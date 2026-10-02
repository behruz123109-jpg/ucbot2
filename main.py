import asyncio
import csv
import hashlib
import hmac
import html
import io
import json
import logging
import os
import random
import re
import time
import traceback
from functools import partial
import urllib.parse
from contextlib import asynccontextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Optional, Union

import aiosqlite
from aiogram import BaseMiddleware, Bot, Dispatcher, F, Router
from aiogram.client.default import DefaultBotProperties
from aiogram.enums import ParseMode
from aiogram.exceptions import TelegramBadRequest, TelegramForbiddenError, TelegramRetryAfter
from aiogram.filters import BaseFilter, Command, CommandObject, CommandStart, StateFilter
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.base import BaseStorage, StorageKey
from aiogram.types import (InputMediaDocument, BufferedInputFile, CallbackQuery, InlineKeyboardMarkup, KeyboardButton,
                            Message, WebAppInfo)
from aiogram.utils.keyboard import InlineKeyboardBuilder, ReplyKeyboardBuilder
from aiohttp import web

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

# ============================================================ SOZLAMALAR
BOT_TOKEN = os.getenv("bott", "")
SUPER_ADMIN_ID = int(os.getenv("id", "0") or 0)
OLD_DB_PATH = "data/pro_uc_bot.db"   # eski joy (kod papkasi yonida) — ma'lumot yo'qolib qolishi mumkin bo'lgan joy


def resolve_db_path() -> str:
    """Baza fayli DOIMIY joyda turishi kerak: kod papkasidan TASHQARIDA.
    Tartib: 1) DB_PATH (qo'lda)  2) hosting volume'i  3) /data  4) uy papkasi (~/.pro_uc_bot)."""
    env = os.getenv("DB_PATH", "").strip()
    if env:
        return env
    for var in ("RAILWAY_VOLUME_MOUNT_PATH", "RENDER_DISK_PATH", "FLY_VOLUME_PATH", "VOLUME_PATH"):
        v = os.getenv(var, "").strip()
        if v:
            return str(Path(v) / "pro_uc_bot.db")
    for d in ("/data", "/var/data"):
        if Path(d).is_dir() and os.access(d, os.W_OK):
            return f"{d}/pro_uc_bot.db"
    try:
        home = Path.home() / ".pro_uc_bot"
        home.mkdir(parents=True, exist_ok=True)
        if os.access(home, os.W_OK):
            return str(home / "pro_uc_bot.db")
    except Exception:
        pass
    return OLD_DB_PATH


def migrate_old_db(new_path: str) -> bool:
    """Eski joydagi bazani (agar bor bo'lsa) yangi doimiy joyga ko'chiradi. Hech narsa o'chirilmaydi."""
    import sqlite3
    old, new = Path(OLD_DB_PATH), Path(new_path)
    try:
        if old.resolve() == new.resolve() or new.exists() or not old.exists() or old.stat().st_size == 0:
            return False
    except OSError:
        return False
    new.parent.mkdir(parents=True, exist_ok=True)
    src_con, dst_con = sqlite3.connect(str(old)), sqlite3.connect(str(new))
    try:
        src_con.backup(dst_con)   # WAL ichidagi ma'lumotlar bilan birga to'liq nusxa
    finally:
        dst_con.close()
        src_con.close()
    return True


DB_PATH = resolve_db_path()
WEBHOOK_URL = os.getenv("WEBHOOK_URL", "").rstrip("/")
WEBHOOK_PATH = os.getenv("WEBHOOK_PATH", "/webhook")
WEBHOOK_SECRET = os.getenv("WEBHOOK_SECRET", "")
WEB_HOST = os.getenv("WEB_HOST", "0.0.0.0")
WEB_PORT = int(os.getenv("WEB_PORT", "8080"))
ORDER_EXPIRE_MIN = int(os.getenv("ORDER_EXPIRE_MIN", "30"))
# --- Telegram bulut zaxirasi: FAQAT BACKUP_CHAT_ID berilsa ishlaydi (sukut bo'yicha O'CHIQ, hech narsa yuborilmaydi) ---
BACKUP_CHAT = int(os.getenv("BACKUP_CHAT_ID", "0") or 0)
BACKUP_EVERY_SEC = max(int(os.getenv("BACKUP_EVERY_SEC", "45") or 45), 15)   # o'zgarish bo'lsa shu oralikda zaxira
BACKUP_CHAT2 = int(os.getenv("BACKUP_CHAT_ID_2", "0") or 0)                    # ixtiyoriy: ikkinchi nusxa (alohida kanal)
# --- Web App (Telegram Mini App) sozlamalari ---
WEBAPP_URL = os.getenv("WEBAPP_URL", "").rstrip("/")   # masalan: https://sizning-domen.com/app
WEBAPP_DIR = Path(os.getenv("WEBAPP_DIR", Path(__file__).parent / "webapp"))
INIT_DATA_MAX_AGE = int(os.getenv("INIT_DATA_MAX_AGE", "86400"))

TZ = timezone(timedelta(hours=5))  # Toshkent vaqti
logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s")
log = logging.getLogger("pro_uc_bot")

DEFAULT_SETTINGS = {
    "card_number": "", "card_owner": "",
    "coin_value": "50",          # 1 tanga = 50 so'm chegirma
    "coin_buy_price": "100",     # 1 tanga sotib olish = 100 so'm (arbitrajga yo'l qo'ymaydi)
    "coin_max_percent": "30",    # narxning maks. necha foizini tanga bilan yopish mumkin
    "welcome_bonus": "5",
    "ref_bonus": "20",
    "ref_purchase_bonus": "20",
    "daily_bonus": "1",
    "cashback_per_100uc": "2",
    "vip_threshold": "600",
    "vip_cashback_mult": "2",
    "topup_min": "5000",
    "support": "@admin",
    "midasbuy_url": "https://www.midasbuy.com",
    "shop_open": "1",
    "work_hours": "09:00 - 23:00",
    "welcome_text": "",
    "backup_msg_id": "0",
    "backup2_msg_id": "0",
    "auto_hours": "0",
    "daily_report": "1",
    "last_report": "",
    "winback_days": "30",
    "lvl_silver": "1500",
    "lvl_gold": "5000",
    "lvl_silver_bonus": "10",
    "lvl_gold_bonus": "25",
    "flash_pct": "0",
    "flash_until": "",
    "price_chat": "0",
    "price_msg_id": "0",
}
SETTING_META = {
    "coin_value": ("1 tanga = necha so'm chegirma", "int"),
    "coin_buy_price": ("1 tanga sotib olish narxi (so'm)", "int"),
    "coin_max_percent": ("Tanga bilan maks. chegirma (%)", "int"),
    "welcome_bonus": ("Yangi foydalanuvchiga tanga", "int"),
    "ref_bonus": ("Do'st qo'shgani uchun tanga", "int"),
    "ref_purchase_bonus": ("Do'st 1-xaridida referalga tanga", "int"),
    "daily_bonus": ("Kunlik bonus (tanga)", "int"),
    "cashback_per_100uc": ("Har 100 UC uchun keshbek (tanga)", "int"),
    "vip_threshold": ("VIP uchun UC chegarasi", "int"),
    "vip_cashback_mult": ("VIP keshbek ko'paytmasi", "int"),
    "topup_min": ("Minimal hisob to'ldirish (so'm)", "int"),
    "support": ("Yordam kontakti (@username)", "str"),
    "midasbuy_url": ("Midasbuy sayt manzili (https://...)", "str"),
    "work_hours": ("Ish vaqti (masalan: 09:00 - 23:00)", "str"),
    "welcome_text": ("Salom xabari ({name} = ism; «-» = asl matn)", "str"),
    "daily_report": ("Kunlik hisobot (1=yoqilgan, 0=o'chiq)", "int"),
    "winback_days": ("Necha kun xarid qilmasa «qaytarish» (kun)", "int"),
    "lvl_silver": ("🥈 Kumush daraja (jami UC)", "int"),
    "lvl_gold": ("🥇 Oltin daraja (jami UC)", "int"),
    "lvl_silver_bonus": ("🥈 Kumush keshbek qo'shimchasi (%)", "int"),
    "lvl_gold_bonus": ("🥇 Oltin keshbek qo'shimchasi (%)", "int"),
}

STATUS_EMOJI = {"awaiting_check": "⏳", "checking": "🔎", "processing": "⚙️", "done": "✅", "cancelled": "❌"}
STATUS_TEXT = {"awaiting_check": "To'lov kutilmoqda", "checking": "Tekshirilmoqda",
               "processing": "UC tashlanmoqda (~5 daqiqa)", "done": "Bajarildi", "cancelled": "Bekor qilindi"}


# ============================================================ MA'LUMOTLAR BAZASI
SCHEMA = """
CREATE TABLE IF NOT EXISTS users(
    id INTEGER PRIMARY KEY, full_name TEXT, username TEXT,
    balance INTEGER NOT NULL DEFAULT 0, coins INTEGER NOT NULL DEFAULT 0,
    total_uc INTEGER NOT NULL DEFAULT 0, is_vip INTEGER NOT NULL DEFAULT 0,
    referrer_id INTEGER, banned INTEGER NOT NULL DEFAULT 0, blocked INTEGER NOT NULL DEFAULT 0,
    last_daily TEXT, ref_paid INTEGER NOT NULL DEFAULT 0, joined_at TEXT DEFAULT (datetime('now','+5 hours'))
);
CREATE TABLE IF NOT EXISTS admins(
    id INTEGER PRIMARY KEY, added_at TEXT DEFAULT (datetime('now','+5 hours'))
);
CREATE TABLE IF NOT EXISTS packages(
    id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, uc INTEGER NOT NULL,
    price INTEGER NOT NULL, cost INTEGER NOT NULL, active INTEGER NOT NULL DEFAULT 1,
    kind TEXT NOT NULL DEFAULT 'uc'
);
CREATE TABLE IF NOT EXISTS orders(
    id INTEGER PRIMARY KEY AUTOINCREMENT, code TEXT NOT NULL, user_id INTEGER NOT NULL,
    pkg_id INTEGER, pkg_name TEXT, uc INTEGER NOT NULL, price INTEGER NOT NULL,
    final INTEGER NOT NULL, cost INTEGER NOT NULL,
    coupon_id INTEGER, coupon_code TEXT, coupon_disc INTEGER NOT NULL DEFAULT 0,
    coin_used INTEGER NOT NULL DEFAULT 0, coin_disc INTEGER NOT NULL DEFAULT 0,
    pubg_id TEXT NOT NULL, pay_method TEXT NOT NULL, status TEXT NOT NULL,
    check_file_id TEXT, admin_id INTEGER,
    created_at TEXT DEFAULT (datetime('now','+5 hours')), done_at TEXT,
    kind TEXT NOT NULL DEFAULT 'uc'
);
CREATE INDEX IF NOT EXISTS idx_orders_user ON orders(user_id);
CREATE INDEX IF NOT EXISTS idx_orders_status ON orders(status);
CREATE TABLE IF NOT EXISTS topups(
    id INTEGER PRIMARY KEY AUTOINCREMENT, code TEXT NOT NULL, user_id INTEGER NOT NULL,
    amount INTEGER NOT NULL, status TEXT NOT NULL, check_file_id TEXT, admin_id INTEGER,
    created_at TEXT DEFAULT (datetime('now','+5 hours')), done_at TEXT
);
CREATE TABLE IF NOT EXISTS coupons(
    id INTEGER PRIMARY KEY AUTOINCREMENT, code TEXT NOT NULL UNIQUE COLLATE NOCASE,
    discount INTEGER NOT NULL, min_uc INTEGER NOT NULL DEFAULT 0,
    limit_count INTEGER NOT NULL, used_count INTEGER NOT NULL DEFAULT 0,
    active INTEGER NOT NULL DEFAULT 1
);
CREATE TABLE IF NOT EXISTS coupon_uses(
    coupon_id INTEGER NOT NULL, user_id INTEGER NOT NULL, PRIMARY KEY(coupon_id, user_id)
);
CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT);
CREATE TABLE IF NOT EXISTS channels(
    id INTEGER PRIMARY KEY AUTOINCREMENT, chat_id INTEGER UNIQUE, title TEXT, link TEXT
);
CREATE TABLE IF NOT EXISTS lottery_log(
    month TEXT PRIMARY KEY, user_id INTEGER, uc INTEGER, cost INTEGER DEFAULT 0,
    created_at TEXT DEFAULT (datetime('now','+5 hours'))
);
CREATE TABLE IF NOT EXISTS fsm(key TEXT PRIMARY KEY, state TEXT, data TEXT);
CREATE TABLE IF NOT EXISTS audit(
    id INTEGER PRIMARY KEY AUTOINCREMENT, admin_id INTEGER, action TEXT NOT NULL, detail TEXT,
    created_at TEXT DEFAULT (datetime('now','+5 hours'))
);
CREATE TABLE IF NOT EXISTS cards(
    id INTEGER PRIMARY KEY AUTOINCREMENT, number TEXT NOT NULL, owner TEXT NOT NULL,
    active INTEGER NOT NULL DEFAULT 1, used INTEGER NOT NULL DEFAULT 0
);
CREATE TABLE IF NOT EXISTS tickets(
    id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL, text TEXT NOT NULL,
    created_at TEXT DEFAULT (datetime('now','+5 hours'))
);
CREATE TABLE IF NOT EXISTS reviews(
    id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL, order_id INTEGER, order_code TEXT,
    uc INTEGER NOT NULL DEFAULT 0, bought_at TEXT, text TEXT, photo_id TEXT,
    status TEXT NOT NULL DEFAULT 'pending',
    created_at TEXT DEFAULT (datetime('now','+5 hours'))
);
CREATE INDEX IF NOT EXISTS idx_reviews_status ON reviews(status);
CREATE TABLE IF NOT EXISTS tournaments(
    id INTEGER PRIMARY KEY AUTOINCREMENT, title TEXT NOT NULL, fee INTEGER NOT NULL DEFAULT 0,
    prize TEXT NOT NULL, start_text TEXT NOT NULL, max_players INTEGER NOT NULL, info TEXT, room_info TEXT,
    status TEXT NOT NULL DEFAULT 'open',
    created_at TEXT DEFAULT (datetime('now','+5 hours'))
);
CREATE TABLE IF NOT EXISTS t_players(
    tid INTEGER NOT NULL, user_id INTEGER NOT NULL, paid INTEGER NOT NULL DEFAULT 0,
    joined_at TEXT DEFAULT (datetime('now','+5 hours')), PRIMARY KEY(tid, user_id)
);
CREATE TABLE IF NOT EXISTS t_comments(
    id INTEGER PRIMARY KEY AUTOINCREMENT, tid INTEGER NOT NULL, user_id INTEGER NOT NULL, text TEXT NOT NULL,
    created_at TEXT DEFAULT (datetime('now','+5 hours'))
);
CREATE INDEX IF NOT EXISTS idx_tcom_tid ON t_comments(tid);
CREATE TABLE IF NOT EXISTS t_reacts(
    tid INTEGER NOT NULL, user_id INTEGER NOT NULL, emoji TEXT NOT NULL, PRIMARY KEY(tid, user_id)
);
"""


class Fail(Exception):
    """Tranzaksiyani bekor qilish uchun."""


DB_MIGRATIONS = [
    ("orders", "claimed_by", "INTEGER"), ("orders", "check_at", "TEXT"), ("orders", "check_uid", "TEXT"),
    ("orders", "check_hash", "TEXT"), ("orders", "card_id", "INTEGER"), ("orders", "flash_pct", "INTEGER NOT NULL DEFAULT 0"),
    ("topups", "check_at", "TEXT"), ("topups", "check_uid", "TEXT"), ("topups", "check_hash", "TEXT"),
    ("topups", "card_id", "INTEGER"),
    ("admins", "role", "TEXT NOT NULL DEFAULT 'full'"),
    ("tournaments", "start_at", "TEXT"), ("tournaments", "reminded", "INTEGER NOT NULL DEFAULT 0"),
    ("t_players", "won", "INTEGER NOT NULL DEFAULT 0"),
    ("users", "last_level", "INTEGER NOT NULL DEFAULT 0"), ("users", "winback_at", "TEXT"),
]


class Database:
    def __init__(self, path: str):
        self.path = path
        self.conn: Optional[aiosqlite.Connection] = None
        self.lock = asyncio.Lock()

    async def connect(self):
        self.conn = await aiosqlite.connect(self.path, isolation_level=None)
        self.conn.row_factory = aiosqlite.Row
        await self.conn.execute("PRAGMA journal_mode=WAL")
        await self.conn.execute("PRAGMA synchronous=NORMAL")
        await self.conn.executescript(SCHEMA)
        ucols = [r["name"] for r in await self.fetchall("PRAGMA table_info(users)")]
        if "ref_paid" not in ucols:   # eski baza: allaqachon xarid qilganlarga referal bonusi berilgan hisoblanadi
            await self.conn.execute("ALTER TABLE users ADD COLUMN ref_paid INTEGER NOT NULL DEFAULT 0")
            await self.conn.execute("UPDATE users SET ref_paid=1 WHERE id IN "
                                    "(SELECT user_id FROM orders WHERE status='done')")
        for tbl, col, ddl in DB_MIGRATIONS:   # yangi ustunlar (eski bazalar uchun xavfsiz)
            cols = [r["name"] for r in await self.fetchall(f"PRAGMA table_info({tbl})")]
            if col not in cols:
                await self.conn.execute(f"ALTER TABLE {tbl} ADD COLUMN {col} {ddl}")
        for tbl in ("packages", "orders"):   # eski bazaga 'kind' ustunini qo'shish
            cols = [r["name"] for r in await self.fetchall(f"PRAGMA table_info({tbl})")]
            if "kind" not in cols:
                await self.conn.execute(f"ALTER TABLE {tbl} ADD COLUMN kind TEXT NOT NULL DEFAULT 'uc'")
        for k, v in DEFAULT_SETTINGS.items():
            await self.conn.execute("INSERT OR IGNORE INTO settings(key,value) VALUES(?,?)", (k, v))
        if SUPER_ADMIN_ID:
            await self.conn.execute("INSERT OR IGNORE INTO admins(id) VALUES(?)", (SUPER_ADMIN_ID,))

    async def close(self):
        if self.conn:
            await self.conn.close()

    async def fetchone(self, q, args=()):
        cur = await self.conn.execute(q, args)
        row = await cur.fetchone()
        await cur.close()
        return row

    async def fetchall(self, q, args=()):
        cur = await self.conn.execute(q, args)
        rows = await cur.fetchall()
        await cur.close()
        return rows

    async def execute(self, q, args=()):
        """(lastrowid, rowcount) qaytaradi."""
        async with self.lock:
            cur = await self.conn.execute(q, args)
            res = (cur.lastrowid, cur.rowcount)
            await cur.close()
            return res

    @asynccontextmanager
    async def tx(self):
        async with self.lock:
            await self.conn.execute("BEGIN IMMEDIATE")
            try:
                yield self.conn
            except BaseException:
                await self.conn.execute("ROLLBACK")
                raise
            else:
                await self.conn.execute("COMMIT")


db = Database(DB_PATH)


async def _one(conn, q, args=()):
    cur = await conn.execute(q, args)
    row = await cur.fetchone()
    await cur.close()
    return row


class SQLiteStorage(BaseStorage):
    """FSM holatlari ham bazada saqlanadi — bot qayta ishga tushsa ham yo'qolmaydi."""

    def __init__(self, database: Database):
        self.db = database

    @staticmethod
    def _k(key: StorageKey) -> str:
        return f"{key.bot_id}:{key.chat_id}:{key.user_id}:{key.thread_id}:{key.destiny}"

    async def set_state(self, key: StorageKey, state=None) -> None:
        s = state.state if isinstance(state, State) else state
        await self.db.execute(
            "INSERT INTO fsm(key,state) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET state=excluded.state",
            (self._k(key), s))

    async def get_state(self, key: StorageKey) -> Optional[str]:
        row = await self.db.fetchone("SELECT state FROM fsm WHERE key=?", (self._k(key),))
        return row["state"] if row else None

    async def set_data(self, key: StorageKey, data) -> None:
        await self.db.execute(
            "INSERT INTO fsm(key,data) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET data=excluded.data",
            (self._k(key), json.dumps(dict(data))))

    async def get_data(self, key: StorageKey) -> dict:
        row = await self.db.fetchone("SELECT data FROM fsm WHERE key=?", (self._k(key),))
        return json.loads(row["data"]) if row and row["data"] else {}

    async def update_data(self, key: StorageKey, data) -> dict:
        cur = await self.get_data(key)
        cur.update(data)
        await self.set_data(key, cur)
        return cur.copy()

    async def close(self) -> None:
        return None



# ============================================================ DOIMIY SAQLASH (Telegram bulut zaxirasi)
BACKUP_TAG = "#PROUC_DB"
BK = {"sig": None, "locked": False, "restored_msg": 0, "warned_big": False}


def _sqlite_ok(path: Path) -> bool:
    import sqlite3
    try:
        con = sqlite3.connect(str(path))
        try:
            r = con.execute("PRAGMA integrity_check").fetchone()
            con.execute("SELECT COUNT(*) FROM users").fetchone()
            return bool(r) and r[0] == "ok"
        finally:
            con.close()
    except Exception:
        return False


async def tg_restore(bot: Bot) -> str:
    """Baza fayli yo'q bo'lsa (yangi server / qayta deploy) — Telegramdagi oxirgi zaxiradan tiklaydi.
    Qaytaradi: 'skip' | 'restored' | 'none' (zaxira topilmadi) | 'error'."""
    p = Path(DB_PATH)
    if (p.exists() and p.stat().st_size > 0) or not BACKUP_CHAT:
        return "skip"
    try:
        chat = await bot.get_chat(BACKUP_CHAT)
        pm = chat.pinned_message
        if not (pm and pm.document and (pm.caption or "").startswith(BACKUP_TAG)):
            log.warning("Bulutda zaxira topilmadi (BACKUP_CHAT=%s) — yangi baza yaratiladi.", BACKUP_CHAT)
            return "none"
        tmp = p.with_suffix(".restore")
        await bot.download(pm.document, destination=str(tmp))
        if not _sqlite_ok(tmp):
            tmp.unlink(missing_ok=True)
            log.error("Bulutdagi zaxira fayli buzuq.")
            return "error"
        for ext in ("-wal", "-shm"):
            Path(str(p) + ext).unlink(missing_ok=True)
        os.replace(tmp, p)
        BK["restored_msg"] = pm.message_id
        log.info("♻️ Baza Telegram zaxirasidan tiklandi.")
        return "restored"
    except Exception:
        log.exception("Zaxirani tiklashda xato")
        return "error"


async def tg_backup(bot: Bot, force: bool = False) -> bool:
    """Bazaning toza nusxasini Telegramga yuboradi. Pin qilingan xabar O'ZGARTIRILADI (yangi xabar/servis xabari chiqmaydi).
    O'zgarish bo'lmasa o'tkazib yuboradi."""
    if not BACKUP_CHAT or not db.conn or BK["locked"]:
        return False
    sig = (await db.fetchone("SELECT total_changes() c"))["c"]
    if not force and sig == BK["sig"]:
        return False
    tmp = Path(DB_PATH).parent / f"tgb_{int(time.time())}.db"
    try:
        async with db.lock:
            await db.conn.execute("VACUUM INTO ?", (str(tmp),))
        data = tmp.read_bytes()
    except Exception:
        log.exception("Zaxira nusxasini olishda xato")
        return False
    finally:
        tmp.unlink(missing_ok=True)
    caption = f"{BACKUP_TAG}\n🕒 {now_tz():%Y-%m-%d %H:%M:%S} • {len(data) // 1024} KB"
    fname = f"pro_uc_bot_{now_tz():%Y%m%d_%H%M}.db"
    try:
        old = int(await S("backup_msg_id") or 0)
        done = False
        if old:   # avval mavjud (pin qilingan) xabarni o'zgartirib ko'ramiz
            try:
                await bot.edit_message_media(
                    chat_id=BACKUP_CHAT, message_id=old,
                    media=InputMediaDocument(media=BufferedInputFile(data, filename=fname), caption=caption))
                done = True
            except TelegramBadRequest as e:
                if "not modified" in str(e).lower():
                    done = True
        if not done:   # xabar yo'q/o'chirilgan — yangisini yuboramiz va pin qilamiz
            msg = await bot.send_document(BACKUP_CHAT, BufferedInputFile(data, filename=fname), caption=caption,
                                          disable_notification=True)
            try:
                await bot.pin_chat_message(BACKUP_CHAT, msg.message_id, disable_notification=True)
            except Exception:
                log.warning("Zaxira xabarini pin qilib bo'lmadi (kanal/guruhda bot admin bo'lishi kerak).")
            await db.execute("UPDATE settings SET value=? WHERE key='backup_msg_id'", (str(msg.message_id),))
            if old and old != msg.message_id:
                try:
                    await bot.delete_message(BACKUP_CHAT, old)
                except Exception:
                    pass
        BK["sig"] = (await db.fetchone("SELECT total_changes() c"))["c"]
        if BACKUP_CHAT2 and (force or time.time() - BK.get("t2", 0) > 1800):   # ikkinchi nusxa: har 30 daqiqada
            BK["t2"] = time.time()
            try:
                m2 = await bot.send_document(BACKUP_CHAT2, BufferedInputFile(data, filename=fname), caption=caption,
                                             disable_notification=True)
                old2 = int(await S("backup2_msg_id") or 0)
                await db.execute("UPDATE settings SET value=? WHERE key='backup2_msg_id'", (str(m2.message_id),))
                if old2:
                    try:
                        await bot.delete_message(BACKUP_CHAT2, old2)
                    except Exception:
                        pass
            except Exception:
                log.warning("Ikkinchi zaxira yuborilmadi (BACKUP_CHAT_ID_2).")
        if len(data) > 19_000_000 and not BK["warned_big"]:
            BK["warned_big"] = True
            for aid in await admin_ids():
                await tell(bot, aid, "⚠️ Baza hajmi 19 MB dan oshdi — Telegram orqali avto-tiklash 20 MB gacha ishlaydi. "
                                     "Eski ma'lumotlarni «🧹 Ma'lumotlarni tozalash» orqali tozalang.")
        return True
    except Exception:
        log.exception("Zaxirani Telegramga yuborishda xato")
        return False


async def backup_loop(bot: Bot):
    """Ma'lumot o'zgargan bo'lsa, har BACKUP_EVERY_SEC soniyada bulutga saqlaydi (yo'qotish oynasi ~1 daqiqa)."""
    await asyncio.sleep(20)
    while True:
        try:
            await tg_backup(bot)
        except Exception:
            log.exception("backup_loop xatosi")
        await asyncio.sleep(BACKUP_EVERY_SEC)


# ============================================================ YORDAMCHI FUNKSIYALAR
def now_tz() -> datetime:
    return datetime.now(TZ)


def fmt(n) -> str:
    return f"{int(n):,}".replace(",", " ")


def esc(s) -> str:
    return html.escape(str(s or ""))


def parse_int(text: str) -> Optional[int]:
    t = re.sub(r"[\s,._]", "", (text or ""))
    if t.lstrip("-").isdigit():
        return int(t)
    return None


def mention(uid: int, name: str) -> str:
    return f'<a href="tg://user?id={uid}">{esc(name) or uid}</a>'


async def S(key: str) -> str:
    row = await db.fetchone("SELECT value FROM settings WHERE key=?", (key,))
    return row["value"] if row else DEFAULT_SETTINGS.get(key, "")


async def Si(key: str) -> int:
    try:
        return int(await S(key))
    except ValueError:
        return int(DEFAULT_SETTINGS.get(key, "0") or 0)


async def _seti(conn, key: str) -> int:
    row = await _one(conn, "SELECT value FROM settings WHERE key=?", (key,))
    try:
        return int(row["value"]) if row else int(DEFAULT_SETTINGS[key])
    except ValueError:
        return int(DEFAULT_SETTINGS[key])


async def is_admin(uid: int) -> bool:
    return bool(await db.fetchone("SELECT 1 FROM admins WHERE id=?", (uid,)))


async def admin_ids() -> list:
    return [r["id"] for r in await db.fetchall("SELECT id FROM admins")]


async def get_user(uid: int):
    return await db.fetchone("SELECT * FROM users WHERE id=?", (uid,))


async def register_user(tg_user, ref_id: Optional[int] = None):
    """(user, yangi_mi, referal_ishladimi)"""
    u = await get_user(tg_user.id)
    name = tg_user.full_name or ""
    if u:
        await db.execute("UPDATE users SET full_name=?, username=?, blocked=0 WHERE id=?",
                         (name, tg_user.username, tg_user.id))
        return await get_user(tg_user.id), False, False
    if ref_id == tg_user.id or (ref_id and not await get_user(ref_id)):
        ref_id = None
    welcome = await Si("welcome_bonus")
    _, rc = await db.execute(
        "INSERT OR IGNORE INTO users(id,full_name,username,coins,referrer_id) VALUES(?,?,?,?,?)",
        (tg_user.id, name, tg_user.username, welcome, ref_id))
    if not rc:
        return await get_user(tg_user.id), False, False
    if ref_id:
        await db.execute("UPDATE users SET coins=coins+? WHERE id=?", (await Si("ref_bonus"), ref_id))
    return await get_user(tg_user.id), True, bool(ref_id)


async def welcome_text(name: str) -> str:
    custom = (await S("welcome_text")).strip()
    return custom.replace("{name}", esc(name)) if custom else WELCOME.format(name=esc(name))


async def gen_code() -> str:
    for _ in range(60):
        code = f"#{random.randint(1000, 9999)}"
        row = await db.fetchone(
            "SELECT 1 FROM orders WHERE code=? AND status IN ('awaiting_check','checking') "
            "UNION SELECT 1 FROM topups WHERE code=? AND status IN ('awaiting_check','checking')",
            (code, code))
        if not row:
            return code
    return f"#{random.randint(100000, 999999)}"


async def safe_edit(msg: Message, text: str, kb: Optional[InlineKeyboardMarkup] = None):
    try:
        await msg.edit_text(text, reply_markup=kb)
    except TelegramBadRequest as e:
        if "not modified" not in str(e).lower():
            await msg.answer(text, reply_markup=kb)


async def send_media(bot: Bot, chat_id: int, fid: Optional[str], text: str, kb=None):
    if not fid:
        return await bot.send_message(chat_id, text, reply_markup=kb)
    if fid.startswith("d:"):
        return await bot.send_document(chat_id, fid[2:], caption=text, reply_markup=kb)
    return await bot.send_photo(chat_id, fid[2:], caption=text, reply_markup=kb)


async def tell(bot: Bot, uid: int, text: str, **kw):
    try:
        await bot.send_message(uid, text, **kw)
    except (TelegramForbiddenError, TelegramBadRequest):
        await db.execute("UPDATE users SET blocked=1 WHERE id=?", (uid,))
    except Exception:
        log.exception("tell xatosi")


# ============================================================ TUGMA MATNLARI VA KLAVIATURALAR
class T:
    BUY = "🛒 UC sotib olish"
    TOPUP = "💰 Hisobni to'ldirish"
    PROFILE = "👤 Profil"
    ORDERS = "📦 Buyurtmalarim"
    COINS = "🪙 Tangalar"
    DAILY = "🎁 Kunlik bonus"
    REF = "👥 Referal"
    RATING = "🏆 Reyting"
    HELP = "☎️ Yordam"
    ADMIN = "⚙️ Admin Panel"
    CANCEL = "❌ Bekor qilish"
    BACK = "🔙 Bosh menyu"
    A_ADMIN = "➕ Admin qo'shish"
    A_CHANNEL = "📢 Majburiy obuna"
    A_BCAST = "📣 Ommaviy xabar"
    A_PKG = "➕ UC paket qo'shish"
    A_PKGS = "📋 Paketlar"
    A_CARD = "💳 Karta ma'lumoti"
    A_STATS = "📊 Statistika"
    A_COUPON = "🎟 Kupon yaratish"
    A_COUPONS = "🎫 Kuponlar"
    A_PENDING = "📥 Kutilayotganlar"
    A_USER = "👤 Foydalanuvchi boshqaruvi"
    A_LOTTERY = "🎰 VIP o'yini"
    A_SETTINGS = "🛠 Tizim sozlamalari"
    A_BACKUP = "💾 Zaxira (Backup)"
    CALC = "🧮 Kalkulyator"
    A_CLEAN = "🧹 Ma'lumotlarni tozalash"
    TOURN = "🏆 Turnirlar"
    REVIEWS = "⭐ Sharxlar"
    A_HOURS = "🕒 Ish rejimi"
    A_REVIEWS = "⭐ Sharx boshqaruvi"
    A_TOURN = "🏆 Turnir boshqaruvi"
    A_AUDIT = "📜 Admin jurnali"
    A_FLASH = "🔥 Aksiya va narx posti"
    A_WINBACK = "💌 Qaytarish kampaniyasi"
    A_midasbuy = "🌐 Midasbuy"


USER_TEXTS = [T.BUY, T.TOPUP, T.PROFILE, T.ORDERS, T.COINS, T.DAILY, T.REF, T.RATING, T.HELP, T.CALC, T.TOURN, T.REVIEWS]
ADMIN_TEXTS = [T.A_ADMIN, T.A_CHANNEL, T.A_BCAST, T.A_PKG, T.A_PKGS, T.A_CARD, T.A_STATS, T.A_COUPON,
               T.A_COUPONS, T.A_PENDING, T.A_USER, T.A_LOTTERY, T.A_SETTINGS, T.A_BACKUP, T.A_midasbuy, T.A_CLEAN,
               T.A_HOURS, T.A_REVIEWS, T.A_TOURN, T.A_AUDIT, T.A_FLASH, T.A_WINBACK]
MENU_TEXTS = set(USER_TEXTS + ADMIN_TEXTS + [T.ADMIN, T.CANCEL, T.BACK])


def menu_kb(admin: bool = False):
    b = ReplyKeyboardBuilder()
    rows = []
    if WEBAPP_URL:
        b.button(text="🎮 Web Do'kon (Mini App)", web_app=WebAppInfo(url=WEBAPP_URL))
        rows.append(1)
    for t in USER_TEXTS:
        b.button(text=t)
    if admin:
        b.button(text=T.ADMIN)
    rows += [2, 2, 2, 2, 2, 2, 1]
    b.adjust(*rows)
    return b.as_markup(resize_keyboard=True)


def admin_kb(operator: bool = False):
    b = ReplyKeyboardBuilder()
    for t in ([T.A_PENDING, T.A_REVIEWS] if operator else ADMIN_TEXTS):
        b.button(text=t)
    b.button(text=T.BACK)
    b.adjust(*([2] * 12))
    return b.as_markup(resize_keyboard=True)


def cancel_kb():
    b = ReplyKeyboardBuilder()
    b.button(text=T.CANCEL)
    return b.as_markup(resize_keyboard=True)


async def main_menu(uid: int):
    return menu_kb(await is_admin(uid))


# ============================================================ FSM HOLATLAR
class Buy(StatesGroup):
    pubg_id = State()
    confirm = State()
    coupon = State()


class Pay(StatesGroup):
    check = State()


class TopUp(StatesGroup):
    amount = State()


class CoinS(StatesGroup):
    exch = State()
    buy = State()


class Calc(StatesGroup):
    uc = State()
    som = State()


class Sup(StatesGroup):
    text = State()


class Rev(StatesGroup):
    text = State()


class TCom(StatesGroup):
    text = State()


class Adm(StatesGroup):
    add_admin = State()
    add_channel = State()
    bcast_msg = State()
    bcast_confirm = State()
    pkg_name = State()
    pkg_uc = State()
    pkg_price = State()
    pkg_cost = State()
    card_number = State()
    card_owner = State()
    cp_code = State()
    cp_disc = State()
    cp_min = State()
    cp_limit = State()
    user_id = State()
    user_bal = State()
    user_coin = State()
    setting_val = State()
    mp_name = State()
    mp_uc = State()
    mp_price = State()
    mp_cost = State()
    t_title = State()
    t_fee = State()
    t_prize = State()
    t_time = State()
    t_max = State()
    t_info = State()
    t_send = State()
    edit_val = State()
    add_channel_name = State()
    bp_pct = State()
    t_win = State()
    support_reply = State()
    fl_pct = State()
    fl_hours = State()
    pp_chat = State()
    wb_amount = State()
    nc_number = State()
    nc_owner = State()


# ============================================================ BIZNES MANTIQ
def flash_price(price: int, pct: int) -> int:
    return max(int(price * (100 - pct) / 100 / 500 + 0.5) * 500, 500) if pct else price


async def flash_pct_now(conn=None) -> int:
    """Faol flash-aksiya foizi (yo'q bo'lsa 0)."""
    c = conn or db.conn
    try:
        pct = int((await _one(c, "SELECT value FROM settings WHERE key='flash_pct'"))["value"] or 0)
        until = (await _one(c, "SELECT value FROM settings WHERE key='flash_until'"))["value"] or ""
    except Exception:
        return 0
    return pct if pct > 0 and until and now_tz().strftime("%Y-%m-%d %H:%M") < until else 0


async def calc_price(conn, uid: int, pkg, coupon_code: Optional[str], use_coins: bool) -> dict:
    fl = await flash_pct_now(conn)
    price = flash_price(pkg["price"], fl)
    res = dict(flash=fl, price=price, coupon_id=None, coupon_code=None, coupon_disc=0, coupon_error=None,
               coin_used=0, coin_disc=0, final=price)
    if coupon_code:
        cp = await _one(conn, "SELECT * FROM coupons WHERE code=? COLLATE NOCASE", (coupon_code,))
        if not cp or not cp["active"]:
            res["coupon_error"] = "Kupon topilmadi yoki faol emas."
        elif cp["used_count"] >= cp["limit_count"]:
            res["coupon_error"] = "Kupon limiti tugagan."
        elif pkg["uc"] < cp["min_uc"]:
            res["coupon_error"] = f"Bu kupon kamida {cp['min_uc']} UC xaridda amal qiladi."
        elif await _one(conn, "SELECT 1 FROM coupon_uses WHERE coupon_id=? AND user_id=?", (cp["id"], uid)):
            res["coupon_error"] = "Siz bu kupondan avval foydalangansiz."
        else:
            cap = max(price - 1000, 0)
            res.update(coupon_id=cp["id"], coupon_code=cp["code"], coupon_disc=min(cp["discount"], cap))
    if use_coins:
        u = await _one(conn, "SELECT coins FROM users WHERE id=?", (uid,))
        coin_value = max(await _seti(conn, "coin_value"), 1)
        pct = min(max(await _seti(conn, "coin_max_percent"), 0), 100)
        remaining = price - res["coupon_disc"]
        max_coins = (remaining * pct // 100) // coin_value
        used = min(u["coins"] if u else 0, max_coins)
        res.update(coin_used=used, coin_disc=used * coin_value)
    res["final"] = max(price - res["coupon_disc"] - res["coin_disc"], 0)
    return res


async def create_order(uid: int, pkg_id: int, pubg_id: str, coupon: Optional[str],
                       use_coins: bool, method: str):
    code = await gen_code()
    try:
        async with db.tx() as c:
            pkg = await _one(c, "SELECT * FROM packages WHERE id=? AND active=1", (pkg_id,))
            if not pkg:
                raise Fail("Paket topilmadi yoki o'chirilgan.")
            pr = await calc_price(c, uid, pkg, coupon, use_coins)
            if pr["coupon_error"]:
                raise Fail(pr["coupon_error"])
            pay_bal = pr["final"] if method == "balance" else 0
            cur = await c.execute(
                "UPDATE users SET balance=balance-?, coins=coins-? WHERE id=? AND balance>=? AND coins>=?",
                (pay_bal, pr["coin_used"], uid, pay_bal, pr["coin_used"]))
            if cur.rowcount == 0:
                raise Fail("Balansingiz yetarli emas. Avval hisobni to'ldiring.")
            if pr["coupon_id"]:
                cur = await c.execute(
                    "UPDATE coupons SET used_count=used_count+1 "
                    "WHERE id=? AND active=1 AND used_count<limit_count", (pr["coupon_id"],))
                if cur.rowcount == 0:
                    raise Fail("Kupon limiti tugagan.")
                try:
                    await c.execute("INSERT INTO coupon_uses(coupon_id,user_id) VALUES(?,?)",
                                    (pr["coupon_id"], uid))
                except aiosqlite.IntegrityError:
                    raise Fail("Siz bu kupondan avval foydalangansiz.")
            status = "checking" if method == "balance" else "awaiting_check"
            cur = await c.execute(
                "INSERT INTO orders(code,user_id,pkg_id,pkg_name,uc,price,final,cost,coupon_id,coupon_code,"
                "coupon_disc,coin_used,coin_disc,pubg_id,pay_method,status,kind) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (code, uid, pkg["id"], pkg["name"], pkg["uc"], pr["price"], pr["final"], pkg["cost"],
                 pr["coupon_id"], pr["coupon_code"], pr["coupon_disc"], pr["coin_used"], pr["coin_disc"],
                 pubg_id, method, status, pkg["kind"]))
            new_id = cur.lastrowid
            if status == "checking":
                await c.execute("UPDATE orders SET check_at=datetime('now','+5 hours') WHERE id=?", (new_id,))
            if pr.get("flash"):
                await c.execute("UPDATE orders SET flash_pct=? WHERE id=?", (pr["flash"], new_id))
            return new_id, code
    except Fail as e:
        return None, str(e)


async def cancel_order(oid: int, statuses=("awaiting_check", "checking")):
    """Buyurtmani bekor qiladi va tanga/kupon/balansni qaytaradi."""
    async with db.tx() as c:
        o = await _one(c, "SELECT * FROM orders WHERE id=?", (oid,))
        if not o or o["status"] not in statuses:
            return None
        await c.execute("UPDATE orders SET status='cancelled', done_at=datetime('now','+5 hours') WHERE id=?",
                        (oid,))
        refund = o["final"] if o["pay_method"] == "balance" else 0
        await c.execute("UPDATE users SET balance=balance+?, coins=coins+? WHERE id=?",
                        (refund, o["coin_used"], o["user_id"]))
        if o["coupon_id"]:
            await c.execute("DELETE FROM coupon_uses WHERE coupon_id=? AND user_id=?",
                            (o["coupon_id"], o["user_id"]))
            await c.execute("UPDATE coupons SET used_count=MAX(used_count-1,0) WHERE id=?", (o["coupon_id"],))
        return o


LEVEL_NAMES = ["🥉 Bronza", "🥈 Kumush", "🥇 Oltin"]


async def level_idx_c(c, total_uc: int) -> int:
    if total_uc >= await _seti(c, "lvl_gold"):
        return 2
    return 1 if total_uc >= await _seti(c, "lvl_silver") else 0


async def level_bonus_c(c, idx: int) -> int:
    return 0 if idx == 0 else max(await _seti(c, "lvl_silver_bonus" if idx == 1 else "lvl_gold_bonus"), 0)


async def level_text(total_uc: int) -> str:
    idx = await level_idx_c(db.conn, total_uc)
    s, g = await Si("lvl_silver"), await Si("lvl_gold")
    bonus = await level_bonus_c(db.conn, idx)
    nxt = ""
    if idx == 0:
        nxt = f" • {LEVEL_NAMES[1]} uchun yana <b>{fmt(max(s - total_uc, 0))} UC</b>"
    elif idx == 1:
        nxt = f" • {LEVEL_NAMES[2]} uchun yana <b>{fmt(max(g - total_uc, 0))} UC</b>"
    return f"{LEVEL_NAMES[idx]}{f' (+{bonus}% keshbek)' if bonus else ''}{nxt}"


async def complete_order(oid: int, admin_id: int, from_status: str = "checking"):
    async with db.tx() as c:
        cur = await c.execute(
            "UPDATE orders SET status='done', done_at=datetime('now','+5 hours'), admin_id=? "
            "WHERE id=? AND status=?", (admin_id, oid, from_status))
        if cur.rowcount == 0:
            return None
        o = await _one(c, "SELECT * FROM orders WHERE id=?", (oid,))
        u = await _one(c, "SELECT * FROM users WHERE id=?", (o["user_id"],))
        per100 = await _seti(c, "cashback_per_100uc")
        mult = max(await _seti(c, "vip_cashback_mult"), 1) if u["is_vip"] else 1
        cashback = o["uc"] * per100 // 100 * mult
        lv_old = await level_idx_c(c, u["total_uc"])
        cashback += cashback * await level_bonus_c(c, lv_old) // 100
        new_total = u["total_uc"] + o["uc"]
        lv_new = await level_idx_c(c, new_total)
        level_up = LEVEL_NAMES[lv_new] if lv_new > u["last_level"] else None
        if lv_new != u["last_level"]:
            await c.execute("UPDATE users SET last_level=? WHERE id=?", (lv_new, u["id"]))
        threshold = await _seti(c, "vip_threshold")
        became_vip = (not u["is_vip"]) and new_total >= threshold
        await c.execute("UPDATE users SET total_uc=?, coins=coins+?, is_vip=? WHERE id=?",
                        (new_total, cashback, 1 if (u["is_vip"] or became_vip) else 0, u["id"]))
        ref_id, ref_bonus = None, 0
        cnt = await _one(c, "SELECT COUNT(*) n FROM orders WHERE user_id=? AND status='done'", (u["id"],))
        if u["referrer_id"] and cnt["n"] == 1 and not u["ref_paid"]:
            ref_id, ref_bonus = u["referrer_id"], await _seti(c, "ref_purchase_bonus")
            await c.execute("UPDATE users SET coins=coins+? WHERE id=?", (ref_bonus, ref_id))
            await c.execute("UPDATE users SET ref_paid=1 WHERE id=?", (u["id"],))
        return dict(order=o, cashback=cashback, became_vip=became_vip, total_uc=new_total,
                    ref_id=ref_id, ref_bonus=ref_bonus, level_up=level_up)


async def midas_accept(oid: int, admin_id: int):
    """Midasbuy: chek tasdiqlandi -> 'processing' (UC tashlanishi kutilmoqda)."""
    _, rc = await db.execute(
        "UPDATE orders SET status='processing', admin_id=? WHERE id=? AND status='checking' AND kind='midas'",
        (admin_id, oid))
    if not rc:
        return None
    return await db.fetchone("SELECT * FROM orders WHERE id=?", (oid,))


async def complete_topup(tid: int, admin_id: int):
    async with db.tx() as c:
        cur = await c.execute(
            "UPDATE topups SET status='done', done_at=datetime('now','+5 hours'), admin_id=? "
            "WHERE id=? AND status='checking'", (admin_id, tid))
        if cur.rowcount == 0:
            return None
        t = await _one(c, "SELECT * FROM topups WHERE id=?", (tid,))
        await c.execute("UPDATE users SET balance=balance+? WHERE id=?", (t["amount"], t["user_id"]))
        return t


async def notify_admins(bot: Bot, kind: str, rid: int, only: Optional[int] = None):
    if kind == "order":
        o = await db.fetchone("SELECT * FROM orders WHERE id=?", (rid,))
        u = await get_user(o["user_id"])
        extra = ""
        if o["coupon_code"]:
            extra += f"\n🎟 Kupon: {esc(o['coupon_code'])} (−{fmt(o['coupon_disc'])})"
        if o["coin_used"]:
            extra += f"\n🪙 Tanga: {o['coin_used']} ta (−{fmt(o['coin_disc'])})"
        extra += await order_extras(o)
        midas = o["kind"] == "midas"
        if not midas:
            tail = f"👉 Pulni tekshiring (izoh: <b>{o['code']}</b>), UC ni PUBG ID ga yuboring, so'ng tasdiqlang."
        elif o["status"] == "processing":
            tail = "⏳ To'lov tasdiqlangan. UC ni Midasbuy orqali yuboring va «UC tashladim» ni bosing."
        else:
            tail = (f"👉 Pulni tekshiring (izoh: <b>{o['code']}</b>) va «To'lovni tasdiqlash» ni bosing — "
                    f"mijozga «5 daqiqada UC tashlaymiz» xabari boradi.")
        text = (f"{'🌐 <b>Midasbuy buyurtma' if midas else '🧾 <b>Yangi buyurtma'} {o['code']}</b>\n"
                f"👤 {mention(u['id'], u['full_name'])} (<code>{u['id']}</code>)"
                f"{' 💎VIP' if u['is_vip'] else ''}\n"
                f"🎮 PUBG ID: <code>{esc(o['pubg_id'])}</code>\n"
                f"💎 Paket: {esc(o['pkg_name'])} — <b>{o['uc']} UC</b>\n"
                f"💵 To'lov: <b>{fmt(o['final'])} so'm</b> (asl narx {fmt(o['price'])}){extra}\n"
                f"💳 Usul: {'Karta (chek)' if o['pay_method'] == 'card' else 'Balansdan (to`langan)'}\n\n"
                f"{tail}")
        multi = len(await admin_ids()) > 1
        claimed = o["claimed_by"]
        mode = "claim" if (multi and claimed is None) else ("locked" if (only and claimed and claimed != only) else "act")
        if mode == "locked":
            cu = await get_user(claimed)
            text += f"\n\n🔒 <b>{esc(cu['full_name'] if cu else claimed)}</b> bajaryapti."
        markup = await order_kb(o, rid, mode, multi)
        fid = o["check_file_id"]
    else:
        t = await db.fetchone("SELECT * FROM topups WHERE id=?", (rid,))
        u = await get_user(t["user_id"])
        text = (f"💰 <b>Hisob to'ldirish {t['code']}</b>\n"
                f"👤 {mention(u['id'], u['full_name'])} (<code>{u['id']}</code>)\n"
                f"💵 Summa: <b>{fmt(t['amount'])} so'm</b>\n\n"
                f"👉 Pul tushganini va izoh <b>{t['code']}</b> ekanini tekshiring.") + await topup_extras(t)
        kb = InlineKeyboardBuilder()
        kb.button(text="✅ Tasdiqlash", callback_data=f"ta:{rid}")
        kb.button(text="❌ Rad etish", callback_data=f"tr:{rid}")
        kb.adjust(1)
        markup = kb.as_markup()
        fid = t["check_file_id"]
    for aid in ([only] if only else await admin_ids()):
        try:
            await send_media(bot, aid, fid, text, markup)
        except Exception:
            log.warning("Adminga xabar yuborilmadi: %s", aid)


async def run_lottery(bot: Bot, month: str) -> str:
    if await db.fetchone("SELECT 1 FROM lottery_log WHERE month=?", (month,)):
        return f"ℹ️ {month} oyi uchun o'yin allaqachon o'tkazilgan."
    row = await db.fetchone(
        "SELECT user_id, SUM(uc) s FROM orders WHERE status='done' AND substr(done_at,1,7)=? "
        "AND user_id IN (SELECT id FROM users WHERE is_vip=1) "
        "GROUP BY user_id ORDER BY s DESC LIMIT 1", (month,))
    if not row:
        return f"ℹ️ {month} oyida VIP xaridorlar topilmadi."
    prize = random.choice([30, 60])
    avg = await db.fetchone("SELECT AVG(cost*1.0/uc) a FROM packages WHERE active=1 AND uc>0")
    cost = int(prize * (avg["a"] or 0))
    await db.execute("INSERT OR IGNORE INTO lottery_log(month,user_id,uc,cost) VALUES(?,?,?,?)",
                     (month, row["user_id"], prize, cost))
    u = await get_user(row["user_id"])
    last = await db.fetchone("SELECT pubg_id FROM orders WHERE user_id=? AND status='done' "
                             "ORDER BY id DESC LIMIT 1", (row["user_id"],))
    pubg = last["pubg_id"] if last else "—"
    await tell(bot, row["user_id"],
               f"🎉 <b>TABRIKLAYMIZ!</b>\n{month} oyining TOP-1 VIP xaridori siz bo'ldingiz "
               f"({row['s']} UC)!\n🎰 O'yin natijasi: <b>{prize} UC</b> yutdingiz!\n"
               f"UC tez orada PUBG ID <code>{esc(pubg)}</code> ga yuboriladi.")
    for aid in await admin_ids():
        await tell(bot, aid,
                   f"🎰 <b>VIP o'yini ({month})</b>\nG'olib: {mention(u['id'], u['full_name'])} "
                   f"(<code>{u['id']}</code>)\nSotib olgan: {row['s']} UC\n🎁 Sovrin: <b>{prize} UC</b>\n"
                   f"🎮 PUBG ID: <code>{esc(pubg)}</code>\n👉 Sovrinni qo'lda yuboring.")
    return f"✅ O'yin o'tkazildi: g'olib {esc(u['full_name'])}, sovrin {prize} UC."


# ============================================================ FILTR VA MIDDLEWARE
class AdminFilter(BaseFilter):
    async def __call__(self, event: Union[Message, CallbackQuery]) -> bool:
        return await is_admin(event.from_user.id)


_sub_cache: dict = {}


async def missing_channels(bot: Bot, uid: int, use_cache: bool = True) -> list:
    now = time.time()
    if use_cache and _sub_cache.get(uid, 0) > now:
        return []
    miss = []
    for ch in await db.fetchall("SELECT * FROM channels"):
        try:
            mb = await bot.get_chat_member(ch["chat_id"], uid)
            if mb.status in ("left", "kicked") or getattr(mb, "is_member", True) is False:
                miss.append(ch)
        except Exception as e:
            log.warning("Kanal tekshirilmadi %s: %s", ch["chat_id"], e)
    if not miss:
        _sub_cache[uid] = now + 60
    return miss


def sub_prompt(miss: list):
    b = InlineKeyboardBuilder()
    for ch in miss:
        b.button(text=f"📢 {ch['title']}", url=ch["link"])
    b.button(text="✅ Obunani tekshirish", callback_data="chk_sub")
    b.adjust(1)
    return "❗️ Botdan foydalanish uchun quyidagi kanallarga obuna bo'ling:", b.as_markup()


class StateInterruptMiddleware(BaseMiddleware):
    """Menyu tugmasi yoki buyruq bosilsa, ochiq FSM holatini tozalaydi."""

    async def __call__(self, handler, event: Message, data):
        text = event.text or ""
        if text in MENU_TEXTS or text.startswith("/"):
            state: Optional[FSMContext] = data.get("state")
            if state:
                await state.clear()
                data["raw_state"] = None
        return await handler(event, data)


class GateMiddleware(BaseMiddleware):
    """Ro'yxatga olish, ban va majburiy obuna tekshiruvi."""

    async def __call__(self, handler, event, data):
        tg = data.get("event_from_user")
        bot: Bot = data["bot"]
        if not tg or tg.is_bot:
            return await handler(event, data)
        is_start = isinstance(event, Message) and (event.text or "").startswith("/start")
        if not is_start:
            await register_user(tg)
        if await is_admin(tg.id):
            return await handler(event, data)
        u = await get_user(tg.id)
        if u and u["banned"]:
            if isinstance(event, CallbackQuery):
                await event.answer("🚫 Siz bloklangansiz.", show_alert=True)
            elif not is_start:
                await event.answer("🚫 Siz bloklangansiz.")
            return
        if is_start or (isinstance(event, CallbackQuery) and event.data == "chk_sub"):
            return await handler(event, data)
        miss = await missing_channels(bot, tg.id)
        if miss:
            text, kb = sub_prompt(miss)
            if isinstance(event, CallbackQuery):
                await event.answer("Avval kanallarga obuna bo'ling!", show_alert=True)
                await event.message.answer(text, reply_markup=kb)
            else:
                await event.answer(text, reply_markup=kb)
            return
        return await handler(event, data)


# ============================================================ ROUTERLAR
common_r = Router()
admin_r = Router()
user_r = Router()
admin_r.message.filter(AdminFilter())
admin_r.callback_query.filter(AdminFilter())


WELCOME = ("👋 Assalomu alaykum, <b>{name}</b>!\n\n💎 <b>Pro UC Bot</b> — PUBG Mobile UC ni ishonchli, "
           "tez va arzon narxlarda sotib oling.\n\n🎁 Har xarid uchun tanga keshbek, VIP maqom, "
           "kuponlar va oylik o'yinlar sizni kutmoqda!")


@common_r.message(F.text == T.CANCEL)
@common_r.message(Command("cancel"))
async def cancel_any(m: Message, state: FSMContext):
    await state.clear()
    await m.answer("❌ Bekor qilindi.", reply_markup=await main_menu(m.from_user.id))


@common_r.message(F.text == T.BACK)
async def back_menu(m: Message, state: FSMContext):
    await state.clear()
    await m.answer("🏠 Bosh menyu", reply_markup=await main_menu(m.from_user.id))


# ------------------------------------------------------------ /start va obuna
@user_r.message(CommandStart())
async def cmd_start(m: Message, state: FSMContext, command: CommandObject, bot: Bot):
    await state.clear()
    ref = None
    if command.args and command.args.startswith("ref_"):
        ref = parse_int(command.args[4:])
    u, new, ref_ok = await register_user(m.from_user, ref)
    if u["banned"]:
        return await m.answer("🚫 Siz bloklangansiz.")
    if new and ref_ok:
        await tell(bot, ref, f"🎉 Sizning havolangiz orqali <b>{esc(m.from_user.full_name)}</b> qo'shildi! "
                             f"+{await Si('ref_bonus')} 🪙 tanga berildi.")
    admin = await is_admin(m.from_user.id)
    if not admin:
        miss = await missing_channels(bot, m.from_user.id, use_cache=False)
        if miss:
            text, kb = sub_prompt(miss)
            return await m.answer(text, reply_markup=kb)
    await m.answer(await welcome_text(m.from_user.full_name) + "\n\n" + await shop_status_text(),
                   reply_markup=menu_kb(admin))


@user_r.callback_query(F.data == "chk_sub")
async def chk_sub(c: CallbackQuery, bot: Bot):
    miss = await missing_channels(bot, c.from_user.id, use_cache=False)
    if miss:
        return await c.answer("❌ Hali barcha kanallarga obuna bo'lmadingiz!", show_alert=True)
    await c.answer("✅ Rahmat!")
    try:
        await c.message.delete()
    except Exception:
        pass
    await c.message.answer(await welcome_text(c.from_user.full_name) + "\n\n" + await shop_status_text(),
                           reply_markup=await main_menu(c.from_user.id))


# ------------------------------------------------------------ Profil, reyting, referal, yordam
@user_r.message(F.text == T.PROFILE)
async def profile(m: Message):
    u = await get_user(m.from_user.id)
    thr = await Si("vip_threshold")
    vip = "💎 <b>VIP</b>" if u["is_vip"] else f"Oddiy (VIP uchun yana <b>{max(thr - u['total_uc'], 0)} UC</b>)"
    refs = await db.fetchone("SELECT COUNT(*) n FROM users WHERE referrer_id=?", (u["id"],))
    await m.answer(
        f"👤 <b>Profil</b>\n\n🆔 ID: <code>{u['id']}</code>\n📛 Ism: {esc(u['full_name'])}\n"
        f"💰 Balans: <b>{fmt(u['balance'])} so'm</b>\n🪙 Tangalar: <b>{fmt(u['coins'])}</b>\n"
        f"💎 Jami xarid: <b>{fmt(u['total_uc'])} UC</b>\n🏅 Maqom: {vip}\n🎖 Daraja: {await level_text(u['total_uc'])}\n"
        f"👥 Takliflar: {refs['n']} ta\n📅 Ro'yxatdan: {u['joined_at'][:10]}")


@user_r.message(F.text == T.RATING)
async def rating(m: Message):
    rows = await db.fetchall("SELECT full_name,total_uc,is_vip FROM users WHERE total_uc>0 "
                             "ORDER BY total_uc DESC LIMIT 10")
    if not rows:
        return await m.answer("🏆 Reyting hozircha bo'sh. Birinchi bo'ling!")
    medals = ["🥇", "🥈", "🥉"] + ["🔹"] * 7
    lines = [f"{medals[i]} {esc((r['full_name'] or '—')[:14])} — <b>{fmt(r['total_uc'])} UC</b>"
             f"{' 💎' if r['is_vip'] else ''}" for i, r in enumerate(rows)]
    await m.answer("🏆 <b>TOP-10 xaridorlar</b>\n\n" + "\n".join(lines) +
                   "\n\n🎰 Har oy VIP'lar orasidan TOP-1 ga <b>30 yoki 60 UC</b> sovg'a o'yini!")


@user_r.message(F.text == T.REF)
async def referral(m: Message, bot: Bot):
    me = await bot.get_me()
    cnt = await db.fetchone("SELECT COUNT(*) n FROM users WHERE referrer_id=?", (m.from_user.id,))
    await m.answer(
        f"👥 <b>Do'stlarni taklif qiling — tanga oling!</b>\n\n"
        f"🔗 Sizning havolangiz:\n<code>https://t.me/{me.username}?start=ref_{m.from_user.id}</code>\n\n"
        f"🎁 Har bir do'st uchun: <b>+{await Si('ref_bonus')} 🪙</b>\n"
        f"🛍 Do'stingiz 1-xaridida: <b>+{await Si('ref_purchase_bonus')} 🪙</b>\n"
        f"👤 Takliflaringiz: <b>{cnt['n']}</b> ta")


@user_r.message(F.text == T.HELP)
async def help_(m: Message):
    kb = InlineKeyboardBuilder()
    kb.button(text="✍️ Savol yozish (bot orqali)", callback_data="sup:new")
    await m.answer(f"☎️ <b>Yordam</b>\n\nSavol va muammolar uchun: {esc(await S('support'))}\n"
                   f"yoki pastdagi tugma orqali shu yerning o'zida yozing — adminlar javob beradi.\n\n"
                   f"ℹ️ To'lov paytida izohga faqat <b>bot bergan kodni</b> (masalan #7492) yozing.",
                   reply_markup=kb.as_markup())


@user_r.message(F.text == T.DAILY)
async def daily(m: Message):
    today = now_tz().strftime("%Y-%m-%d")
    bonus = await Si("daily_bonus")
    _, rc = await db.execute("UPDATE users SET coins=coins+?, last_daily=? WHERE id=? "
                             "AND (last_daily IS NULL OR last_daily<>?)",
                             (bonus, today, m.from_user.id, today))
    if not rc:
        return await m.answer("⏳ Bugungi bonusni olgansiz. Ertaga qaytib keling!")
    await m.answer(f"🎁 Kunlik bonus: <b>+{bonus} 🪙</b> tanga!")


# ------------------------------------------------------------ Tangalar
@user_r.message(F.text == T.COINS)
async def coins_menu(m: Message):
    u = await get_user(m.from_user.id)
    cv, bp = await Si("coin_value"), await Si("coin_buy_price")
    kb = InlineKeyboardBuilder()
    kb.button(text="🔄 Tangani so'mga almashtirish", callback_data="coin:exch")
    kb.button(text="🛍 Tanga sotib olish", callback_data="coin:buy")
    kb.adjust(1)
    await m.answer(
        f"🪙 <b>Tangalar</b>\n\nSizda: <b>{fmt(u['coins'])} 🪙</b>\n\n"
        f"• 1 🪙 = <b>{fmt(cv)} so'm</b> chegirma (xarid paytida ishlatiladi, "
        f"narxning {await Si('coin_max_percent')}% gacha)\n"
        f"• 1 🪙 sotib olish narxi: <b>{fmt(bp)} so'm</b>\n\n"
        f"<b>Tanga qanday olinadi?</b>\n🛒 Har 100 UC uchun +{await Si('cashback_per_100uc')} 🪙 "
        f"(VIP'ga ×{await Si('vip_cashback_mult')})\n👥 Do'st taklif qilish\n🎁 Kunlik bonus", reply_markup=kb.as_markup())


@user_r.callback_query(F.data == "coin:exch")
async def coin_exch(c: CallbackQuery, state: FSMContext):
    await c.answer()
    await state.set_state(CoinS.exch)
    u = await get_user(c.from_user.id)
    await c.message.answer(f"🔄 Nechta tangani so'mga almashtirasiz? (Sizda {u['coins']} 🪙)\n"
                           f"1 🪙 = {fmt(await Si('coin_value'))} so'm", reply_markup=cancel_kb())


@user_r.message(StateFilter(CoinS.exch))
async def coin_exch_amount(m: Message, state: FSMContext):
    n = parse_int(m.text or "")
    if not n or n <= 0:
        return await m.answer("❗️ Musbat son kiriting.")
    val = n * await Si("coin_value")
    _, rc = await db.execute("UPDATE users SET coins=coins-?, balance=balance+? WHERE id=? AND coins>=?",
                             (n, val, m.from_user.id, n))
    if not rc:
        return await m.answer("❌ Tangangiz yetarli emas. Boshqa son kiriting.")
    await state.clear()
    await m.answer(f"✅ {n} 🪙 → <b>{fmt(val)} so'm</b> balansga qo'shildi.",
                   reply_markup=await main_menu(m.from_user.id))


@user_r.callback_query(F.data == "coin:buy")
async def coin_buy(c: CallbackQuery, state: FSMContext):
    await c.answer()
    await state.set_state(CoinS.buy)
    u = await get_user(c.from_user.id)
    await c.message.answer(f"🛍 Nechta tanga sotib olasiz?\n1 🪙 = {fmt(await Si('coin_buy_price'))} so'm\n"
                           f"💰 Balansingiz: {fmt(u['balance'])} so'm", reply_markup=cancel_kb())


@user_r.message(StateFilter(CoinS.buy))
async def coin_buy_amount(m: Message, state: FSMContext):
    n = parse_int(m.text or "")
    if not n or n <= 0:
        return await m.answer("❗️ Musbat son kiriting.")
    cost = n * await Si("coin_buy_price")
    _, rc = await db.execute("UPDATE users SET balance=balance-?, coins=coins+? WHERE id=? AND balance>=?",
                             (cost, n, m.from_user.id, cost))
    if not rc:
        return await m.answer("❌ Balans yetarli emas. Avval hisobni to'ldiring yoki kamroq kiriting.")
    await state.clear()
    await m.answer(f"✅ {n} 🪙 sotib olindi (−{fmt(cost)} so'm).", reply_markup=await main_menu(m.from_user.id))


# ------------------------------------------------------------ 🧮 Kalkulyator
def calc_by_uc(pkgs: list, target: int):
    """Kamida `target` UC ni eng arzon narxda beradigan paketlar to'plami. pkgs: (nom, uc, narx, tannarx)."""
    size = target + max(p[1] for p in pkgs)
    INF = float("inf")
    best, pick = [INF] * (size + 1), [-1] * (size + 1)
    best[0] = 0
    for u in range(1, size + 1):
        for i, (_, uc, pr, _c) in enumerate(pkgs):
            if uc <= u and best[u - uc] + pr < best[u]:
                best[u], pick[u] = best[u - uc] + pr, i
    cand = [u for u in range(target, size + 1) if best[u] < INF]
    if not cand:
        return None
    u = min(cand, key=lambda x: (best[x], x))
    counts, cur = {}, u
    while cur > 0:
        i = pick[cur]
        counts[i] = counts.get(i, 0) + 1
        cur -= pkgs[i][1]
    return counts


def calc_by_budget(pkgs: list, budget: int):
    """Berilgan pulga eng ko'p UC beradigan paketlar to'plami. Juda katta summada None qaytaradi."""
    from functools import reduce
    from math import gcd
    g = reduce(gcd, [p[2] for p in pkgs]) or 1
    B = budget // g
    if B > 150000:
        return None
    dp, ch = [0] * (B + 1), [-1] * (B + 1)
    for b in range(1, B + 1):
        dp[b] = dp[b - 1]
        for i, (_, uc, pr, _c) in enumerate(pkgs):
            w = pr // g
            if w <= b and dp[b - w] + uc > dp[b]:
                dp[b], ch[b] = dp[b - w] + uc, i
    counts, b = {}, B
    while b > 0:
        i = ch[b]
        if i == -1:
            b -= 1
        else:
            counts[i] = counts.get(i, 0) + 1
            b -= pkgs[i][2] // g
    return counts


def calc_report(pkgs: list, counts: dict, admin: bool, head: str) -> str:
    lines, uc_sum, pr_sum, co_sum = [], 0, 0, 0
    for i, c in sorted(counts.items(), key=lambda kv: -pkgs[kv[0]][1]):
        n, uc, pr, co = pkgs[i]
        lines.append(f"• {c} × {esc(n)} — {fmt(pr * c)} so'm")
        uc_sum, pr_sum, co_sum = uc_sum + uc * c, pr_sum + pr * c, co_sum + co * c
    text = (f"🧮 <b>Hisob-kitob</b>\n{head}\n\n📦 <b>Eng qulay to'plam:</b>\n" + "\n".join(lines) +
            f"\n\n💎 Jami: <b>{fmt(uc_sum)} UC</b>\n💵 To'lov: <b>{fmt(pr_sum)} so'm</b>")
    if admin:
        text += f"\n\n🔒 <i>Admin uchun:</i> tannarx {fmt(co_sum)} so'm | foyda <b>{fmt(pr_sum - co_sum)} so'm</b>"
    return text


async def run_blocking(fn, *args):
    return await asyncio.get_running_loop().run_in_executor(None, partial(fn, *args))


async def calc_packages() -> list:
    rows = await db.fetchall("SELECT name,uc,price,cost FROM packages WHERE active=1 AND uc>0 AND price>0 ORDER BY uc")
    fl = await flash_pct_now()
    return [(r["name"], r["uc"], flash_price(r["price"], fl), r["cost"]) for r in rows]


def calc_menu_kb():
    kb = InlineKeyboardBuilder()
    kb.button(text="💎 UC miqdori → narx", callback_data="calc:uc")
    kb.button(text="💵 Pul miqdori → nechta UC", callback_data="calc:som")
    kb.adjust(1)
    return kb.as_markup()


@user_r.message(F.text == T.CALC)
async def calc_menu(m: Message):
    await m.answer("🧮 <b>Kalkulyator</b>\n\nQo'lda son kiriting — bot eng arzon variantni hisoblab beradi.\n"
                   "Nimani hisoblaymiz?", reply_markup=calc_menu_kb())


@user_r.callback_query(F.data.in_({"calc:uc", "calc:som"}))
async def calc_pick(c: CallbackQuery, state: FSMContext):
    await c.answer()
    if not await calc_packages():
        return await c.message.answer("😔 Hozircha paketlar mavjud emas.")
    if c.data == "calc:uc":
        await state.set_state(Calc.uc)
        await c.message.answer("💎 Necha UC kerak? (masalan: 1250)", reply_markup=cancel_kb())
    else:
        await state.set_state(Calc.som)
        await c.message.answer("💵 Qancha pulingiz bor? So'mda yozing (masalan: 150000)", reply_markup=cancel_kb())


@user_r.message(StateFilter(Calc.uc))
async def calc_uc_input(m: Message):
    n = parse_int(m.text or "")
    if not n or not 1 <= n <= 30000:
        return await m.answer("❗️ UC miqdorini 1 dan 30 000 gacha son bilan yozing.")
    pkgs = await calc_packages()
    counts = await run_blocking(calc_by_uc, pkgs, n)
    if not counts:
        return await m.answer("😔 Hisoblab bo'lmadi.")
    extra = sum(pkgs[i][1] * c for i, c in counts.items()) - n
    head = f"🎯 Kerak: <b>{fmt(n)} UC</b>" + (f"\nℹ️ Paketlar bo'yicha {fmt(extra)} UC ortiqcha chiqadi (eng arzon variant)." if extra else "")
    await m.answer(calc_report(pkgs, counts, await is_admin(m.from_user.id), head) +
                   "\n\n🔁 Yana son yuboring yoki ❌ Bekor qilish ni bosing.")


@user_r.message(StateFilter(Calc.som))
async def calc_som_input(m: Message):
    n = parse_int(m.text or "")
    if not n or n <= 0:
        return await m.answer("❗️ Summani so'mda son bilan yozing.")
    pkgs = await calc_packages()
    cheapest = min(p[2] for p in pkgs)
    if n < cheapest:
        return await m.answer(f"😔 Eng arzon paket {fmt(cheapest)} so'm. Kattaroq summa yozing.")
    counts = await run_blocking(calc_by_budget, pkgs, n)
    if counts is None:
        return await m.answer("❗️ Summa juda katta, kichikroq son yozing (masalan 5 000 000 gacha).")
    spent = sum(pkgs[i][2] * c for i, c in counts.items())
    head = f"💰 Pulingiz: <b>{fmt(n)} so'm</b>" + (f"\nℹ️ Qoladi: {fmt(n - spent)} so'm" if n - spent else "")
    await m.answer(calc_report(pkgs, counts, await is_admin(m.from_user.id), head) +
                   "\n\n🔁 Yana son yuboring yoki ❌ Bekor qilish ni bosing.")


# ------------------------------------------------------------ UC sotib olish
@user_r.message(F.text == T.BUY)
async def buy_list(m: Message):
    pk = await db.fetchall("SELECT * FROM packages WHERE active=1 ORDER BY uc")
    if not pk:
        return await m.answer("😔 Hozircha paketlar mavjud emas. Keyinroq urinib ko'ring.")
    b = InlineKeyboardBuilder()
    regular = [p for p in pk if p["kind"] != "midas"]
    midas = [p for p in pk if p["kind"] == "midas"]
    fl = await flash_pct_now()
    last = await db.fetchone("SELECT o.id, o.pubg_id, o.pkg_id, p.name FROM orders o JOIN packages p ON p.id=o.pkg_id "
                             "WHERE o.user_id=? AND o.status='done' AND p.active=1 ORDER BY o.id DESC LIMIT 1",
                             (m.from_user.id,))
    if last:
        b.button(text=f"🔁 Qayta: {last['name']} • ID {last['pubg_id']}", callback_data=f"rb:{last['id']}")
    for p in regular + midas:   # Midasbuy paketlari mijozga oddiy paket ko'rinishida, ro'yxat pastida
        pr = flash_price(p["price"], fl)
        b.button(text=f"{'🔥' if fl else '💎'} {p['name']} • {fmt(pr)} so'm", callback_data=f"buy:{p['id']}")
    b.adjust(1)
    banner = ""
    if fl:
        until = await S("flash_until")
        banner = f"\n\n🔥 <b>AKSIYA −{fl}%!</b> {esc(until[5:] if until else '')} gacha"
    await m.answer("🛒 <b>UC paketini tanlang:</b>\n\n💎 Jami 600 UC dan oshsa — avtomatik <b>VIP</b>!" + banner,
                   reply_markup=b.as_markup())


@user_r.callback_query(F.data.startswith("rb:"))
async def rebuy(c: CallbackQuery, state: FSMContext):
    o = await db.fetchone("SELECT * FROM orders WHERE id=? AND user_id=? AND status='done'",
                          (int(c.data.split(":")[1]), c.from_user.id))
    p = await db.fetchone("SELECT id FROM packages WHERE id=? AND active=1", (o["pkg_id"],)) if o else None
    if not p:
        return await c.answer("Bu paket hozir mavjud emas.", show_alert=True)
    await c.answer()
    await state.clear()
    await state.update_data(pid=o["pkg_id"], coupon=None, use_coins=False, pubg_id=o["pubg_id"])
    await state.set_state(Buy.confirm)
    text, kb = await render_summary(c.from_user.id, await state.get_data())
    if not text:
        await state.clear()
        return await c.message.answer("Paket topilmadi.")
    await c.message.answer(text + await closed_note(), reply_markup=kb)


@user_r.callback_query(F.data.startswith("buy:"))
async def buy_pick(c: CallbackQuery, state: FSMContext):
    pid = int(c.data.split(":")[1])
    p = await db.fetchone("SELECT * FROM packages WHERE id=? AND active=1", (pid,))
    if not p:
        return await c.answer("Paket topilmadi", show_alert=True)
    await c.answer()
    await state.clear()
    await state.update_data(pid=pid, coupon=None, use_coins=False)
    await state.set_state(Buy.pubg_id)
    await c.message.answer(f"✅ Tanlandi: <b>{esc(p['name'])}</b> — {fmt(flash_price(p['price'], await flash_pct_now()))} so'm\n\n"
                           f"🎮 Endi <b>PUBG Mobile ID</b> raqamingizni yuboring:" + await closed_note(),
                           reply_markup=cancel_kb())


async def render_summary(uid: int, data: dict):
    pkg = await db.fetchone("SELECT * FROM packages WHERE id=? AND active=1", (data["pid"],))
    if not pkg:
        return None, None
    pr = await calc_price(db.conn, uid, pkg, data.get("coupon"), data.get("use_coins", False))
    u = await get_user(uid)
    lines = [f"🧾 <b>Buyurtma</b>\n", f"🎮 PUBG ID: <code>{esc(data['pubg_id'])}</code>",
             f"💎 Paket: {esc(pkg['name'])} — {pkg['uc']} UC",
             f"💵 Narx: {fmt(pr['price'])} so'm"]
    if pr["coupon_id"]:
        lines.append(f"🎟 Kupon ({esc(pr['coupon_code'])}): −{fmt(pr['coupon_disc'])} so'm")
    if pr["coin_used"]:
        lines.append(f"🪙 Tanga ({pr['coin_used']} ta): −{fmt(pr['coin_disc'])} so'm")
    lines += ["━━━━━━━━━━", f"✅ <b>To'lov: {fmt(pr['final'])} so'm</b>", f"💰 Balansingiz: {fmt(u['balance'])} so'm"]
    kb = InlineKeyboardBuilder()
    kb.button(text="🎟 Kupon kiritish", callback_data="b:coupon")
    if u["coins"] > 0:
        kb.button(text=f"🪙 Tanga ishlatish: {'✅ Yoqilgan' if data.get('use_coins') else '❌ O`chiq'}",
                  callback_data="b:coins")
    kb.button(text="💳 Karta orqali to'lash", callback_data="b:card")
    kb.button(text="💰 Balansdan to'lash", callback_data="b:bal")
    kb.button(text="✖️ Bekor qilish", callback_data="b:cancel")
    kb.adjust(1)
    return "\n".join(lines), kb.as_markup()


@user_r.message(StateFilter(Buy.pubg_id))
async def buy_pubg(m: Message, state: FSMContext):
    t = (m.text or "").strip()
    if not (t.isdigit() and 7 <= len(t) <= 12):
        return await m.answer("❗️ PUBG ID faqat raqamlardan iborat bo'lishi kerak (7–12 xona). Qayta yuboring:")
    await state.update_data(pubg_id=t)
    await state.set_state(Buy.confirm)
    text, kb = await render_summary(m.from_user.id, await state.get_data())
    if not text:
        await state.clear()
        return await m.answer("Paket topilmadi.", reply_markup=await main_menu(m.from_user.id))
    await m.answer(text, reply_markup=kb)


@user_r.callback_query(StateFilter(Buy.confirm), F.data == "b:coupon")
async def buy_coupon_ask(c: CallbackQuery, state: FSMContext):
    await c.answer()
    await state.set_state(Buy.coupon)
    await c.message.answer("🎟 Kupon (promokod) kodini yuboring:", reply_markup=cancel_kb())


@user_r.message(StateFilter(Buy.coupon))
async def buy_coupon_set(m: Message, state: FSMContext):
    code = (m.text or "").strip()
    data = await state.get_data()
    pkg = await db.fetchone("SELECT * FROM packages WHERE id=?", (data["pid"],))
    pr = await calc_price(db.conn, m.from_user.id, pkg, code, False)
    if pr["coupon_error"]:
        return await m.answer(f"❌ {pr['coupon_error']}\nBoshqa kod yuboring yoki bekor qiling.")
    await state.update_data(coupon=code)
    await state.set_state(Buy.confirm)
    text, kb = await render_summary(m.from_user.id, await state.get_data())
    await m.answer("✅ Kupon qo'llandi!")
    await m.answer(text, reply_markup=kb)


@user_r.callback_query(StateFilter(Buy.confirm), F.data == "b:coins")
async def buy_toggle_coins(c: CallbackQuery, state: FSMContext):
    d = await state.get_data()
    await state.update_data(use_coins=not d.get("use_coins", False))
    text, kb = await render_summary(c.from_user.id, await state.get_data())
    await c.answer()
    await safe_edit(c.message, text, kb)


@user_r.callback_query(F.data == "b:cancel")
async def buy_cancel(c: CallbackQuery, state: FSMContext):
    await state.clear()
    await c.answer("Bekor qilindi")
    await safe_edit(c.message, "❌ Buyurtma bekor qilindi.")
    await c.message.answer("🏠 Bosh menyu", reply_markup=await main_menu(c.from_user.id))


async def pick_card():
    """Eng kam ishlatilgan faol karta; karta yo'q bo'lsa — eski (sozlamalardagi) karta."""
    r = await db.fetchone("SELECT * FROM cards WHERE active=1 ORDER BY used ASC, id ASC LIMIT 1")
    if r:
        return r["number"], r["owner"], r["id"]
    return await S("card_number"), await S("card_owner"), None


async def has_card() -> bool:
    return bool(await S("card_number") or await db.fetchone("SELECT 1 FROM cards WHERE active=1"))


def payment_text(card: str, owner: str, amount: int, code: str) -> str:
    return (f"💳 <b>To'lov ma'lumotlari</b>\n\nKarta: <code>{esc(card)}</code>\nEgasi: <b>{esc(owner)}</b>\n"
            f"Summa: <b>{fmt(amount)} so'm</b>\n\n"
            f"🔑 To'lov izohiga shu kodni yozing: <code>{code}</code>\n"
            f"⚠️ Kod yozilmasa to'lov tasdiqlanmasligi mumkin!\n"
            f"⏳ {ORDER_EXPIRE_MIN} daqiqa ichida to'lang.\n\n📸 To'lagach, <b>chek (skrinshot)</b>ni shu yerga yuboring.")


@user_r.callback_query(StateFilter(Buy.confirm), F.data.in_({"b:card", "b:bal"}))
async def buy_pay(c: CallbackQuery, state: FSMContext, bot: Bot):
    method = "card" if c.data == "b:card" else "balance"
    card, owner, card_id = await pick_card()
    if method == "card" and not card:
        return await c.answer("Karta ma'lumoti kiritilmagan. Admin bilan bog'laning.", show_alert=True)
    d = await state.get_data()
    oid, res = await create_order(c.from_user.id, d["pid"], d["pubg_id"], d.get("coupon"),
                                  d.get("use_coins", False), method)
    if oid is None:
        return await c.answer(res, show_alert=True)
    await c.answer()
    o = await db.fetchone("SELECT * FROM orders WHERE id=?", (oid,))
    if method == "balance":
        await state.clear()
        await safe_edit(c.message, f"✅ <b>Buyurtma {o['code']} qabul qilindi!</b>\nTo'lov balansdan olindi. "
                                   f"UC tez orada PUBG ID <code>{esc(o['pubg_id'])}</code> ga yuboriladi." + await closed_note())
        await c.message.answer("🏠 Bosh menyu", reply_markup=await main_menu(c.from_user.id))
        await notify_admins(bot, "order", oid)
    else:
        if card_id:
            await db.execute("UPDATE cards SET used=used+1 WHERE id=?", (card_id,))
            await db.execute("UPDATE orders SET card_id=? WHERE id=?", (card_id, oid))
        await state.set_state(Pay.check)
        await state.update_data(kind="order", ref_id=oid)
        await safe_edit(c.message, payment_text(card, owner, o["final"], o["code"]))
        await c.message.answer("📸 Chekni yuboring (rasm yoki fayl):", reply_markup=cancel_kb())


@user_r.message(StateFilter(Pay.check), F.photo | F.document)
async def got_check(m: Message, state: FSMContext, bot: Bot):
    d = await state.get_data()
    kind, rid = d.get("kind"), d.get("ref_id")
    fid = ("p:" + m.photo[-1].file_id) if m.photo else ("d:" + m.document.file_id)
    table = "orders" if kind == "order" else "topups"
    obj = m.photo[-1] if m.photo else m.document
    cuid, chash = obj.file_unique_id, None
    try:   # chek faylining barmoq izi (takroriy chekni aniqlash uchun)
        if (obj.file_size or 0) <= 5_000_000:
            bio = io.BytesIO()
            await bot.download(obj, destination=bio)
            chash = hashlib.sha256(bio.getvalue()).hexdigest()
    except Exception:
        log.warning("Chek hash olinmadi")
    _, rc = await db.execute(
        f"UPDATE {table} SET status='checking', check_file_id=?, check_at=datetime('now','+5 hours'), "
        f"check_uid=?, check_hash=? WHERE id=? AND user_id=? AND status='awaiting_check'",
        (fid, cuid, chash, rid, m.from_user.id))
    await state.clear()
    if not rc:
        return await m.answer("⚠️ Bu so'rov muddati tugagan yoki allaqachon yuborilgan.",
                              reply_markup=await main_menu(m.from_user.id))
    await m.answer("✅ Chek qabul qilindi! Admin tekshirib, tez orada tasdiqlaydi." +
                   (await closed_note() if kind == "order" else ""),
                   reply_markup=await main_menu(m.from_user.id))
    await notify_admins(bot, kind, rid)


@user_r.message(StateFilter(Pay.check))
async def check_not_photo(m: Message):
    await m.answer("📸 Iltimos, to'lov chekini <b>rasm</b> yoki <b>fayl</b> ko'rinishida yuboring.")


# ------------------------------------------------------------ Hisobni to'ldirish
@user_r.message(F.text == T.TOPUP)
async def topup_start(m: Message, state: FSMContext):
    if not await has_card():
        return await m.answer("😔 Hozircha to'lov qabul qilinmayapti. Admin bilan bog'laning.")
    await state.set_state(TopUp.amount)
    await m.answer(f"💰 Qancha so'm to'ldirmoqchisiz?\nMinimal: <b>{fmt(await Si('topup_min'))} so'm</b>",
                   reply_markup=cancel_kb())


@user_r.message(StateFilter(TopUp.amount))
async def topup_amount(m: Message, state: FSMContext):
    n = parse_int(m.text or "")
    mn = await Si("topup_min")
    if not n or n < mn or n > 100_000_000:
        return await m.answer(f"❗️ Summa {fmt(mn)} dan 100 000 000 gacha bo'lishi kerak.")
    code = await gen_code()
    tid, _ = await db.execute("INSERT INTO topups(code,user_id,amount,status) VALUES(?,?,?,'awaiting_check')",
                              (code, m.from_user.id, n))
    await state.set_state(Pay.check)
    await state.update_data(kind="topup", ref_id=tid)
    card, owner, card_id = await pick_card()
    if card_id:
        await db.execute("UPDATE cards SET used=used+1 WHERE id=?", (card_id,))
        await db.execute("UPDATE topups SET card_id=? WHERE id=?", (card_id, tid))
    await m.answer(payment_text(card, owner, n, code))


# ------------------------------------------------------------ Buyurtmalarim
@user_r.message(F.text == T.ORDERS)
async def my_orders(m: Message):
    rows = await db.fetchall("SELECT * FROM orders WHERE user_id=? ORDER BY id DESC LIMIT 8", (m.from_user.id,))
    if not rows:
        return await m.answer("📦 Sizda hali buyurtmalar yo'q.")
    lines, kb = [], InlineKeyboardBuilder()
    for o in rows:
        lines.append(f"{STATUS_EMOJI[o['status']]} <b>{o['code']}</b> • {o['uc']} UC • {fmt(o['final'])} so'm • "
                     f"{STATUS_TEXT[o['status']]}\n    🕒 {o['created_at'][:16]}")
        if o["status"] == "awaiting_check":
            kb.button(text=f"📸 Chek yuborish {o['code']}", callback_data=f"oc_send:{o['id']}")
            kb.button(text=f"✖️ Bekor {o['code']}", callback_data=f"oc_cancel:{o['id']}")
    kb.adjust(2)
    await m.answer("📦 <b>So'nggi buyurtmalaringiz</b>\n\n" + "\n".join(lines), reply_markup=kb.as_markup())


@user_r.callback_query(F.data.startswith("oc_send:"))
async def oc_send(c: CallbackQuery, state: FSMContext):
    oid = int(c.data.split(":")[1])
    o = await db.fetchone("SELECT * FROM orders WHERE id=? AND user_id=?", (oid, c.from_user.id))
    if not o or o["status"] != "awaiting_check":
        return await c.answer("Bu buyurtma endi faol emas.", show_alert=True)
    await c.answer()
    await state.set_state(Pay.check)
    await state.update_data(kind="order", ref_id=oid)
    await c.message.answer(payment_text(await S("card_number"), await S("card_owner"), o["final"], o["code"]),
                           reply_markup=cancel_kb())


@user_r.callback_query(F.data.startswith("oc_cancel:"))
async def oc_cancel(c: CallbackQuery):
    oid = int(c.data.split(":")[1])
    o = await db.fetchone("SELECT user_id FROM orders WHERE id=?", (oid,))
    if not o or o["user_id"] != c.from_user.id:
        return await c.answer("Topilmadi", show_alert=True)
    if await cancel_order(oid, ("awaiting_check",)):
        await c.answer("Bekor qilindi")
        await safe_edit(c.message, "✅ Buyurtma bekor qilindi (tanga/kupon qaytarildi).")
    else:
        await c.answer("Bu buyurtmani bekor qilib bo'lmaydi.", show_alert=True)


# ============================================================ ADMIN PANEL
@admin_r.message(F.text == T.ADMIN)
@admin_r.message(Command("admin"))
async def admin_panel(m: Message, state: FSMContext):
    await state.clear()
    await m.answer("⚙️ <b>Admin Panel</b>", reply_markup=admin_kb(await admin_role(m.from_user.id) == "operator"))


# ------------------------------------------------------------ Admin qo'shish
async def admins_view():
    rows = await db.fetchall("SELECT id, role FROM admins ORDER BY added_at")
    b = InlineKeyboardBuilder()
    lines = []
    for r in rows:
        sup = r["id"] == SUPER_ADMIN_ID
        role = "👑 asosiy" if sup else ("🧑‍💼 operator" if r["role"] == "operator" else "🔧 to'liq huquqli")
        lines.append(f"• <code>{r['id']}</code> — {role}")
        if not sup:
            b.button(text=f"🗑 {r['id']}", callback_data=f"ad:{r['id']}")
            b.button(text="🔄 Rolni almashtirish", callback_data=f"adr:{r['id']}")
    b.adjust(2)
    return ("👥 <b>Adminlar:</b>\n" + "\n".join(lines) +
            "\n\n🧑‍💼 <b>Operator</b> faqat chek/buyurtmalarni tasdiqlaydi, sharx va murojaatlarga javob beradi; "
            "narx, sozlama, tozalash kabilarga kira olmaydi.\n\n➕ Yangi admin <b>Telegram ID</b> sini yuboring:"), b.as_markup()


@admin_r.callback_query(F.data.startswith("adr:"))
async def admin_role_toggle(c: CallbackQuery):
    if c.from_user.id != SUPER_ADMIN_ID:
        return await c.answer("Rolni faqat asosiy admin o'zgartira oladi.", show_alert=True)
    aid = int(c.data.split(":")[1])
    if aid == SUPER_ADMIN_ID:
        return await c.answer("Asosiy adminning rolini o'zgartirib bo'lmaydi.", show_alert=True)
    await db.execute("UPDATE admins SET role=CASE role WHEN 'operator' THEN 'full' ELSE 'operator' END WHERE id=?", (aid,))
    await audit(c.from_user.id, "admin_rol", str(aid))
    await c.answer("Rol almashtirildi")
    text, kb = await admins_view()
    await safe_edit(c.message, text, kb)


@admin_r.message(F.text == T.A_ADMIN)
async def add_admin_start(m: Message, state: FSMContext):
    text, kb = await admins_view()
    await state.set_state(Adm.add_admin)
    await m.answer(text, reply_markup=kb)
    await m.answer("ID yuboring yoki bekor qiling 👇", reply_markup=cancel_kb())


@admin_r.callback_query(F.data.startswith("ad:"))
async def del_admin(c: CallbackQuery):
    aid = int(c.data.split(":")[1])
    if aid == SUPER_ADMIN_ID or aid == c.from_user.id:
        return await c.answer("Bu adminni o'chirib bo'lmaydi.", show_alert=True)
    await db.execute("DELETE FROM admins WHERE id=?", (aid,))
    await c.answer("O'chirildi")
    text, kb = await admins_view()
    await safe_edit(c.message, text, kb)


@admin_r.message(StateFilter(Adm.add_admin))
async def add_admin_id(m: Message, state: FSMContext, bot: Bot):
    n = parse_int(m.text or "")
    if not n or n <= 0:
        return await m.answer("❗️ To'g'ri Telegram ID kiriting (faqat raqam).")
    await db.execute("INSERT OR IGNORE INTO admins(id) VALUES(?)", (n,))
    await state.clear()
    await m.answer(f"✅ <code>{n}</code> admin qilindi.", reply_markup=admin_kb())
    await tell(bot, n, "🎉 Siz Pro UC Bot adminisiz! /start ni bosing.")


# ------------------------------------------------------------ Majburiy obuna
async def channels_view():
    rows = await db.fetchall("SELECT * FROM channels")
    b = InlineKeyboardBuilder()
    for r in rows:
        b.button(text=f"✏️ {r['title'][:24]}", callback_data=f"edm:ch:{r['id']}")
        b.button(text="🗑", callback_data=f"chd:{r['id']}")
    b.button(text="➕ Kanal qo'shish", callback_data="cha")
    b.adjust(*([2] * len(rows)), 1)
    text = "📢 <b>Majburiy obuna kanallari</b>\n\n" + (
        "\n".join(f"• {esc(r['title'])} — {esc(r['link'])}" for r in rows) if rows else "Hozircha kanal yo'q.")
    return text + "\n\n⚠️ Bot kanalda <b>admin</b> bo'lishi shart.", b.as_markup()


@admin_r.message(F.text == T.A_CHANNEL)
async def channels_menu(m: Message):
    text, kb = await channels_view()
    await m.answer(text, reply_markup=kb, disable_web_page_preview=True)


@admin_r.callback_query(F.data == "cha")
async def channel_add_ask(c: CallbackQuery, state: FSMContext):
    await c.answer()
    await state.set_state(Adm.add_channel)
    await c.message.answer("📢 Kanal <b>@username</b> yoki ID sini (-100...) yuboring:", reply_markup=cancel_kb())


@admin_r.message(StateFilter(Adm.add_channel))
async def channel_add(m: Message, state: FSMContext, bot: Bot):
    raw = (m.text or "").strip()
    ident = raw if raw.startswith("@") else (parse_int(raw) if parse_int(raw) else "@" + raw.lstrip("@"))
    try:
        chat = await bot.get_chat(ident)
        me = await bot.get_chat_member(chat.id, bot.id)
        if me.status not in ("administrator", "creator"):
            return await m.answer("❌ Bot bu kanalda admin emas. Avval botni admin qiling.")
        link = f"https://t.me/{chat.username}" if chat.username else (
            chat.invite_link or await bot.export_chat_invite_link(chat.id))
    except Exception as e:
        return await m.answer(f"❌ Kanal topilmadi yoki bot ruxsati yo'q.\n<code>{esc(e)}</code>")
    real = chat.title or str(chat.id)
    await state.update_data(ch_id=chat.id, ch_title=real, ch_link=link)
    await state.set_state(Adm.add_channel_name)
    await m.answer(f"✅ Kanal topildi: <b>{esc(real)}</b>\n\n✏️ Obuna tugmasida <b>qanday nom</b> ko'rinsin?\n"
                   f"Masalan: <code>1-kanal</code>, <code>Asosiy kanal</code>, <code>Yangiliklar</code>.\n"
                   f"Asl nomni qoldirish uchun <code>-</code> yozing.", reply_markup=cancel_kb())


@admin_r.message(StateFilter(Adm.add_channel_name))
async def channel_add_name(m: Message, state: FSMContext):
    d = await state.get_data()
    name = (m.text or "").strip()
    if not name:
        return await m.answer("❗️ Nom yozing yoki <code>-</code> yuboring.")
    title = d["ch_title"] if name == "-" else name[:40]
    await db.execute("INSERT OR REPLACE INTO channels(chat_id,title,link) VALUES(?,?,?)",
                     (d["ch_id"], title, d["ch_link"]))
    _sub_cache.clear()
    await state.clear()
    await m.answer(f"✅ Kanal qo'shildi. Tugma nomi: <b>{esc(title)}</b>", reply_markup=admin_kb())


@admin_r.callback_query(F.data.startswith("chd:"))
async def channel_del(c: CallbackQuery):
    await db.execute("DELETE FROM channels WHERE id=?", (int(c.data.split(":")[1]),))
    _sub_cache.clear()
    await c.answer("O'chirildi")
    text, kb = await channels_view()
    await safe_edit(c.message, text, kb)


# ------------------------------------------------------------ Broadcast
@admin_r.message(F.text == T.A_BCAST)
async def bcast_start(m: Message, state: FSMContext):
    await state.set_state(Adm.bcast_msg)
    await m.answer("📣 Yuboriladigan xabarni jo'nating (matn, rasm, video — istalgan format):",
                   reply_markup=cancel_kb())


@admin_r.message(StateFilter(Adm.bcast_msg))
async def bcast_preview(m: Message, state: FSMContext, bot: Bot):
    await state.update_data(src_chat=m.chat.id, src_msg=m.message_id)
    await state.set_state(Adm.bcast_confirm)
    n = await db.fetchone("SELECT COUNT(*) n FROM users WHERE banned=0 AND blocked=0")
    await bot.copy_message(m.chat.id, m.chat.id, m.message_id)
    kb = InlineKeyboardBuilder()
    kb.button(text=f"✅ Yuborish ({n['n']} ta)", callback_data="bc:go")
    kb.button(text="❌ Bekor", callback_data="bc:no")
    kb.adjust(1)
    await m.answer("👆 Xabar shunday ko'rinadi. Yuborilsinmi?", reply_markup=kb.as_markup())


BG_TASKS: set = set()


async def do_broadcast(bot: Bot, admin_id: int, chat_id: int, msg_id: int):
    ids = [r["id"] for r in await db.fetchall("SELECT id FROM users WHERE banned=0 AND blocked=0")]
    ok = fail = 0
    for uid in ids:
        try:
            await bot.copy_message(uid, chat_id, msg_id)
            ok += 1
        except TelegramRetryAfter as e:
            await asyncio.sleep(e.retry_after + 1)
            try:
                await bot.copy_message(uid, chat_id, msg_id)
                ok += 1
            except Exception:
                fail += 1
        except TelegramForbiddenError:
            await db.execute("UPDATE users SET blocked=1 WHERE id=?", (uid,))
            fail += 1
        except Exception:
            fail += 1
        await asyncio.sleep(0.05)
    await tell(bot, admin_id, f"📣 <b>Xabar yuborish yakunlandi</b>\n✅ Yetkazildi: {ok}\n❌ Xato/bloklagan: {fail}")


@admin_r.callback_query(StateFilter(Adm.bcast_confirm), F.data.in_({"bc:go", "bc:no"}))
async def bcast_go(c: CallbackQuery, state: FSMContext, bot: Bot):
    d = await state.get_data()
    await state.clear()
    if c.data == "bc:no":
        await c.answer("Bekor qilindi")
        return await safe_edit(c.message, "❌ Bekor qilindi.")
    await c.answer("Yuborilmoqda...")
    await safe_edit(c.message, "⏳ Xabar yuborilmoqda... Tugagach hisobot keladi.")
    t = asyncio.create_task(do_broadcast(bot, c.from_user.id, d["src_chat"], d["src_msg"]))
    BG_TASKS.add(t)
    t.add_done_callback(BG_TASKS.discard)
    await c.message.answer("Admin panel", reply_markup=admin_kb())


# ------------------------------------------------------------ UC paket qo'shish / boshqarish
@admin_r.message(F.text == T.A_PKG)
async def pkg_start(m: Message, state: FSMContext):
    await state.set_state(Adm.pkg_name)
    await m.answer("➕ Paket <b>nomini</b> yozing (masalan: 60 UC):", reply_markup=cancel_kb())


@admin_r.message(StateFilter(Adm.pkg_name))
async def pkg_name(m: Message, state: FSMContext):
    await state.update_data(name=(m.text or "").strip()[:40])
    await state.set_state(Adm.pkg_uc)
    await m.answer("💎 UC miqdori (raqam):")


@admin_r.message(StateFilter(Adm.pkg_uc))
async def pkg_uc(m: Message, state: FSMContext):
    n = parse_int(m.text or "")
    if not n or n <= 0:
        return await m.answer("❗️ Musbat son kiriting.")
    await state.update_data(uc=n)
    await state.set_state(Adm.pkg_price)
    await m.answer("💵 <b>Sotuv narxi</b> (so'm):")


@admin_r.message(StateFilter(Adm.pkg_price))
async def pkg_price(m: Message, state: FSMContext):
    n = parse_int(m.text or "")
    if not n or n <= 0:
        return await m.answer("❗️ Musbat son kiriting.")
    await state.update_data(price=n)
    await state.set_state(Adm.pkg_cost)
    await m.answer("🏷 <b>Tannarx</b> (siz UC ni qanchaga olasiz, so'm):")


@admin_r.message(StateFilter(Adm.pkg_cost))
async def pkg_cost(m: Message, state: FSMContext):
    n = parse_int(m.text or "")
    if n is None or n < 0:
        return await m.answer("❗️ Son kiriting.")
    d = await state.get_data()
    await db.execute("INSERT INTO packages(name,uc,price,cost) VALUES(?,?,?,?)",
                     (d["name"], d["uc"], d["price"], n))
    await state.clear()
    await m.answer(f"✅ Paket qo'shildi: <b>{esc(d['name'])}</b> — {d['uc']} UC\n"
                   f"Sotuv: {fmt(d['price'])} | Tannarx: {fmt(n)} | Foyda: <b>{fmt(d['price'] - n)} so'm</b>",
                   reply_markup=admin_kb())


async def packages_view():
    rows = await db.fetchall("SELECT * FROM packages WHERE kind='uc' ORDER BY uc")
    b = InlineKeyboardBuilder()
    lines = []
    for p in rows:
        lines.append(f"{'🟢' if p['active'] else '⚪️'} <b>{esc(p['name'])}</b> • {p['uc']} UC • "
                     f"{fmt(p['price'])} / tannarx {fmt(p['cost'])} (foyda {fmt(p['price'] - p['cost'])})")
        b.button(text=f"{'⏸' if p['active'] else '▶️'} {p['name']}", callback_data=f"pt:{p['id']}")
        b.button(text="🗑", callback_data=f"pd:{p['id']}")
        b.button(text="✏️", callback_data=f"edm:pkg:{p['id']}")
    b.button(text="📈 Narxni ommaviy o'zgartirish", callback_data="bp")
    b.adjust(*([3] * len(rows)), 1)
    return "📋 <b>Paketlar</b>\n\n" + ("\n".join(lines) if lines else "Paket yo'q."), b.as_markup()


@admin_r.message(F.text == T.A_PKGS)
async def pkgs_menu(m: Message):
    text, kb = await packages_view()
    await m.answer(text, reply_markup=kb)


@admin_r.callback_query(F.data.startswith(("pt:", "pd:")))
async def pkg_manage(c: CallbackQuery):
    act, pid = c.data.split(":")
    if act == "pt":
        await db.execute("UPDATE packages SET active=1-active WHERE id=?", (int(pid),))
    else:
        await db.execute("DELETE FROM packages WHERE id=?", (int(pid),))
    await c.answer("Bajarildi")
    text, kb = await packages_view()
    await safe_edit(c.message, text, kb)


# ------------------------------------------------------------ Karta
async def cards_view():
    rows = await db.fetchall("SELECT * FROM cards ORDER BY id")
    b = InlineKeyboardBuilder()
    lines = []
    for r in rows:
        lines.append(f"{'🟢' if r['active'] else '⚪️'} <code>{esc(r['number'])}</code> — {esc(r['owner'])} • ishlatilgan: {r['used']}")
        b.button(text=f"{'⏸' if r['active'] else '▶️'} …{r['number'][-4:]}", callback_data=f"cdt:{r['id']}")
        b.button(text="🗑", callback_data=f"cdd:{r['id']}")
    b.button(text="➕ Karta qo'shish", callback_data="cda")
    b.button(text="✏️ Zaxira karta (eski)", callback_data="cdl")
    b.adjust(*([2] * len(rows)), 1, 1)
    legacy = await S("card_number")
    return ("💳 <b>Kartalar</b>\n\n" + ("\n".join(lines) if lines else "Karta qo'shilmagan.") +
            f"\n\nℹ️ Bir nechta faol karta bo'lsa, bot ularni navbat bilan beradi (eng kam ishlatilgani). "
            f"Chek adminga qaysi kartaga to'langani bilan keladi.\n"
            f"🛟 Zaxira karta (faol karta bo'lmasa ishlaydi): <code>{esc(legacy) or '—'}</code>"), b.as_markup()


@admin_r.message(F.text == T.A_CARD)
async def card_start(m: Message, state: FSMContext):
    await state.clear()
    text, kb = await cards_view()
    await m.answer(text, reply_markup=kb)


@admin_r.callback_query(F.data == "cdl")
async def card_legacy(c: CallbackQuery, state: FSMContext):
    await c.answer()
    await state.set_state(Adm.card_number)
    await c.message.answer(f"💳 Hozirgi zaxira karta: <code>{esc(await S('card_number')) or '—'}</code> "
                           f"({esc(await S('card_owner')) or '—'})\n\nYangi <b>karta raqamini</b> yuboring:",
                           reply_markup=cancel_kb())


@admin_r.callback_query(F.data == "cda")
async def card_add(c: CallbackQuery, state: FSMContext):
    await c.answer()
    await state.set_state(Adm.nc_number)
    await c.message.answer("💳 Yangi <b>karta raqamini</b> yuboring:", reply_markup=cancel_kb())


@admin_r.message(StateFilter(Adm.nc_number))
async def card_add_number(m: Message, state: FSMContext):
    digits = re.sub(r"\D", "", m.text or "")
    if not 13 <= len(digits) <= 19:
        return await m.answer("❗️ Karta raqami 13–19 ta raqamdan iborat bo'lishi kerak.")
    await state.update_data(card=" ".join(digits[i:i + 4] for i in range(0, len(digits), 4)))
    await state.set_state(Adm.nc_owner)
    await m.answer("👤 Karta egasining <b>F.I.O.</b> sini yuboring:")


@admin_r.message(StateFilter(Adm.nc_owner))
async def card_add_owner(m: Message, state: FSMContext):
    d = await state.get_data()
    owner = (m.text or "").strip()[:60]
    if not owner:
        return await m.answer("❗️ F.I.O. yozing.")
    await db.execute("INSERT INTO cards(number,owner) VALUES(?,?)", (d["card"], owner))
    await state.clear()
    await audit(m.from_user.id, "karta qo'shildi", "…" + d["card"][-4:])
    await m.answer("✅ Karta qo'shildi.", reply_markup=admin_kb())
    text, kb = await cards_view()
    await m.answer(text, reply_markup=kb)


@admin_r.callback_query(F.data.startswith(("cdt:", "cdd:")))
async def card_manage(c: CallbackQuery):
    act, cid = c.data.split(":")
    if act == "cdt":
        await db.execute("UPDATE cards SET active=1-active WHERE id=?", (int(cid),))
    else:
        await db.execute("DELETE FROM cards WHERE id=?", (int(cid),))
    await audit(c.from_user.id, "karta " + ("almashtirildi" if act == "cdt" else "o'chirildi"), cid)
    await c.answer("Bajarildi")
    text, kb = await cards_view()
    await safe_edit(c.message, text, kb)


@admin_r.message(StateFilter(Adm.card_number))
async def card_number(m: Message, state: FSMContext):
    digits = re.sub(r"\D", "", m.text or "")
    if not 13 <= len(digits) <= 19:
        return await m.answer("❗️ Karta raqami 13–19 ta raqamdan iborat bo'lishi kerak.")
    await state.update_data(card=" ".join(digits[i:i + 4] for i in range(0, len(digits), 4)))
    await state.set_state(Adm.card_owner)
    await m.answer("👤 Karta egasining <b>F.I.O.</b> sini yuboring:")


@admin_r.message(StateFilter(Adm.card_owner))
async def card_owner(m: Message, state: FSMContext):
    d = await state.get_data()
    await db.execute("UPDATE settings SET value=? WHERE key='card_number'", (d["card"],))
    await db.execute("UPDATE settings SET value=? WHERE key='card_owner'", ((m.text or "").strip()[:60],))
    await state.clear()
    await m.answer(f"✅ Karta yangilandi:\n<code>{d['card']}</code>\n{esc(m.text)}", reply_markup=admin_kb())


# ------------------------------------------------------------ Statistika va Avto-Hisobchi
async def stats_text() -> str:
    today = now_tz().strftime("%Y-%m-%d")
    month = now_tz().strftime("%Y-%m")

    async def agg(where="", args=()):
        r = await db.fetchone(
            "SELECT COUNT(*) n, COALESCE(SUM(uc),0) uc, COALESCE(SUM(final),0) rev, COALESCE(SUM(cost),0) cost "
            f"FROM orders WHERE status='done' {where}", args)
        return r

    total, td, mo = await agg(), await agg("AND substr(done_at,1,10)=?", (today,)), \
        await agg("AND substr(done_at,1,7)=?", (month,))
    users = await db.fetchone("SELECT COUNT(*) n, COALESCE(SUM(is_vip),0) v, COALESCE(SUM(balance),0) b, "
                              "COALESCE(SUM(coins),0) c FROM users")
    today_users = await db.fetchone("SELECT COUNT(*) n FROM users WHERE substr(joined_at,1,10)=?", (today,))
    lot = await db.fetchone("SELECT COALESCE(SUM(cost),0) c FROM lottery_log")
    pend = await db.fetchone("SELECT (SELECT COUNT(*) FROM orders WHERE status IN ('checking','processing')) o, "
                             "(SELECT COUNT(*) FROM topups WHERE status='checking') t")
    profit = total["rev"] - total["cost"]
    net = profit - lot["c"]
    return (f"📊 <b>Statistika & Avto-Hisobchi</b>\n\n"
            f"👥 Foydalanuvchilar: <b>{fmt(users['n'])}</b> (bugun +{today_users['n']})\n"
            f"💎 VIP: <b>{users['v']}</b>\n"
            f"💰 Foydalanuvchilar balansi: {fmt(users['b'])} so'm | 🪙 {fmt(users['c'])}\n\n"
            f"<b>📦 Jami savdo</b>\nBuyurtmalar: {fmt(total['n'])} | UC: {fmt(total['uc'])}\n"
            f"Tushum: {fmt(total['rev'])} so'm\nTannarx: {fmt(total['cost'])} so'm\n"
            f"Sovrinlar xarajati: {fmt(lot['c'])} so'm\n"
            f"✅ <b>SOF FOYDA: {fmt(net)} so'm</b>\n\n"
            f"<b>📅 Bugun:</b> {td['n']} ta • {fmt(td['rev'])} so'm • foyda {fmt(td['rev'] - td['cost'])}\n"
            f"<b>🗓 Shu oy:</b> {mo['n']} ta • {fmt(mo['rev'])} so'm • foyda {fmt(mo['rev'] - mo['cost'])}\n\n"
            f"⏳ Kutilayotgan: {pend['o']} buyurtma, {pend['t']} to'ldirish")


@admin_r.message(F.text == T.A_STATS)
async def stats(m: Message):
    kb = InlineKeyboardBuilder()
    kb.button(text="👥 Foydalanuvchilar ro'yxati (CSV)", callback_data="st:users")
    kb.button(text="🧾 Buyurtmalar (CSV)", callback_data="st:orders")
    kb.button(text="📗 Excel hisobot (.xlsx)", callback_data="st:xlsx")
    kb.button(text="⏱ Tezlik hisoboti", callback_data="st:speed")
    kb.adjust(1)
    await m.answer(await stats_text(), reply_markup=kb.as_markup())


@admin_r.callback_query(F.data.in_({"st:users", "st:orders"}))
async def stats_export(c: CallbackQuery):
    await c.answer("Tayyorlanmoqda...")
    buf = io.StringIO()
    w = csv.writer(buf)
    if c.data == "st:users":
        w.writerow(["id", "ism", "username", "balans", "tanga", "jami_uc", "vip", "referal", "ban", "qo'shilgan"])
        for r in await db.fetchall("SELECT * FROM users ORDER BY joined_at"):
            w.writerow([r["id"], r["full_name"], r["username"], r["balance"], r["coins"], r["total_uc"],
                        r["is_vip"], r["referrer_id"], r["banned"], r["joined_at"]])
        name = "users.csv"
    else:
        w.writerow(["id", "kod", "user_id", "pubg_id", "paket", "uc", "narx", "to'langan", "tannarx",
                    "foyda", "holat", "usul", "sana"])
        for r in await db.fetchall("SELECT * FROM orders ORDER BY id"):
            w.writerow([r["id"], r["code"], r["user_id"], r["pubg_id"], r["pkg_name"], r["uc"], r["price"],
                        r["final"], r["cost"], r["final"] - r["cost"], r["status"], r["pay_method"],
                        r["created_at"]])
        name = "orders.csv"
    data = buf.getvalue().encode("utf-8-sig")
    await c.message.answer_document(BufferedInputFile(data, filename=name))


# ------------------------------------------------------------ Kuponlar
@admin_r.message(F.text == T.A_COUPON)
async def cp_start(m: Message, state: FSMContext):
    await state.set_state(Adm.cp_code)
    await m.answer("🎟 Kupon <b>kodini</b> yozing (lotin harf/raqam, masalan: PUBG2026):", reply_markup=cancel_kb())


@admin_r.message(StateFilter(Adm.cp_code))
async def cp_code(m: Message, state: FSMContext):
    code = (m.text or "").strip().upper()
    if not re.fullmatch(r"[A-Z0-9_]{3,20}", code):
        return await m.answer("❗️ Kod 3–20 ta lotin harf/raqamdan iborat bo'lsin.")
    if await db.fetchone("SELECT 1 FROM coupons WHERE code=?", (code,)):
        return await m.answer("❗️ Bu kod mavjud. Boshqa kod yozing.")
    await state.update_data(code=code)
    await state.set_state(Adm.cp_disc)
    await m.answer("💵 Chegirma summasi (so'mda):")


@admin_r.message(StateFilter(Adm.cp_disc))
async def cp_disc(m: Message, state: FSMContext):
    n = parse_int(m.text or "")
    if not n or n <= 0:
        return await m.answer("❗️ Musbat son kiriting.")
    await state.update_data(disc=n)
    await state.set_state(Adm.cp_min)
    await m.answer("💎 Kamida necha UC dan boshlab amal qilsin? (0 — cheklovsiz):")

@admin_r.message(StateFilter(Adm.cp_min))
async def cp_min(m: Message, state: FSMContext):
    n = parse_int(m.text or "")
    if n is None or n < 0:
        return await m.answer("❗️ 0 yoki musbat son kiriting.")
    await state.update_data(min_uc=n)
    await state.set_state(Adm.cp_limit)
    await m.answer("👥 Necha kishiga mo'ljallangan? (limit):")


@admin_r.message(StateFilter(Adm.cp_limit))
async def cp_limit(m: Message, state: FSMContext):
    n = parse_int(m.text or "")
    if not n or n <= 0:
        return await m.answer("❗️ Musbat son kiriting.")
    d = await state.get_data()
    await db.execute("INSERT INTO coupons(code,discount,min_uc,limit_count) VALUES(?,?,?,?)",
                     (d["code"], d["disc"], d["min_uc"], n))
    await state.clear()
    await m.answer(f"✅ Kupon yaratildi!\n🎟 <code>{d['code']}</code>\n💵 Chegirma: {fmt(d['disc'])} so'm\n"
                   f"💎 Minimal: {d['min_uc']} UC\n👥 Limit: {n} kishi", reply_markup=admin_kb())


async def coupons_view():
    rows = await db.fetchall("SELECT * FROM coupons ORDER BY id DESC LIMIT 20")
    b = InlineKeyboardBuilder()
    lines = []
    for r in rows:
        lines.append(f"{'🟢' if r['active'] else '⚪️'} <code>{esc(r['code'])}</code> • −{fmt(r['discount'])} • "
                     f"≥{r['min_uc']} UC • {r['used_count']}/{r['limit_count']}")
        b.button(text=f"{'⏸' if r['active'] else '▶️'} {r['code']}", callback_data=f"cpt:{r['id']}")
        b.button(text="🗑", callback_data=f"cpd:{r['id']}")
        b.button(text="✏️", callback_data=f"edm:cp:{r['id']}")
    b.adjust(3)
    return "🎫 <b>Kuponlar</b>\n\n" + ("\n".join(lines) if lines else "Kupon yo'q."), b.as_markup()


@admin_r.message(F.text == T.A_COUPONS)
async def coupons_menu(m: Message):
    text, kb = await coupons_view()
    await m.answer(text, reply_markup=kb)


@admin_r.callback_query(F.data.startswith(("cpt:", "cpd:")))
async def coupon_manage(c: CallbackQuery):
    act, cid = c.data.split(":")
    if act == "cpt":
        await db.execute("UPDATE coupons SET active=1-active WHERE id=?", (int(cid),))
    else:
        await db.execute("DELETE FROM coupons WHERE id=?", (int(cid),))
    await c.answer("Bajarildi")
    text, kb = await coupons_view()
    await safe_edit(c.message, text, kb)


# ------------------------------------------------------------ Buyurtma / to'ldirishni tasdiqlash
async def mark_msg(c: CallbackQuery, note: str, kb: Optional[InlineKeyboardMarkup] = None):
    try:
        body = c.message.html_text or ""
        if c.message.photo or c.message.document:
            await c.message.edit_caption(caption=f"{body}\n\n{note}", reply_markup=kb)
        else:
            await c.message.edit_text(f"{body}\n\n{note}", reply_markup=kb)
    except TelegramBadRequest:
        try:
            await c.message.edit_reply_markup(reply_markup=kb)
        except TelegramBadRequest:
            pass


@admin_r.callback_query(F.data.startswith("oa:"))
async def order_approve(c: CallbackQuery, bot: Bot):
    oid = int(c.data.split(":")[1])
    if not await claim_guard(c, oid):
        return
    info = await complete_order(oid, c.from_user.id)
    if not info:
        await c.answer("Bu buyurtma allaqachon ishlangan.", show_alert=True)
        return await mark_msg(c, "ℹ️ Allaqachon ishlangan.")
    await c.answer("✅ Tasdiqlandi")
    o = info["order"]
    await audit(c.from_user.id, "UC yuborildi", f"{o['code']} • {o['uc']} UC • {fmt(o['final'])} so'm")
    await mark_msg(c, f"✅ Tasdiqladi: {esc(c.from_user.full_name)}")
    txt = (f"✅ <b>Buyurtma {o['code']} bajarildi!</b>\n💎 {o['uc']} UC PUBG ID <code>{esc(o['pubg_id'])}</code> "
           f"ga yuborildi.\n🪙 Keshbek: <b>+{info['cashback']}</b> tanga\n💎 Jami xaridingiz: {info['total_uc']} UC")
    if info["became_vip"]:
        txt += "\n\n🎉 <b>Tabriklaymiz! Siz endi VIP maqomdasiz!</b> Keshbek ×2 va oylik TOP-1 o'yinida ishtirok."
    if info.get("level_up"):
        txt += f"\n\n🏅 Yangi daraja: <b>{info['level_up']}</b> — keshbekingiz oshdi!"
    await tell(bot, o["user_id"], txt, reply_markup=review_kb(o["id"]))
    if info["ref_id"]:
        await tell(bot, info["ref_id"], f"🎁 Taklif qilgan do'stingiz 1-xaridini qildi! +{info['ref_bonus']} 🪙")


@admin_r.callback_query(F.data.startswith("or:"))
async def order_reject(c: CallbackQuery, bot: Bot):
    oid = int(c.data.split(":")[1])
    if not await claim_guard(c, oid):
        return
    o = await cancel_order(oid)
    if not o:
        await c.answer("Bu buyurtma allaqachon ishlangan.", show_alert=True)
        return await mark_msg(c, "ℹ️ Allaqachon ishlangan.")
    await audit(c.from_user.id, "buyurtma rad", o["code"])
    await c.answer("Rad etildi")
    await mark_msg(c, f"❌ Rad etdi: {esc(c.from_user.full_name)}")
    await tell(bot, o["user_id"], f"❌ <b>Buyurtma {o['code']} rad etildi.</b>\nTo'lov tasdiqlanmadi. "
                                  f"Tanga/kupon/balans qaytarildi. Savol bo'lsa: {esc(await S('support'))}")


@admin_r.callback_query(F.data.startswith("ta:"))
async def topup_approve(c: CallbackQuery, bot: Bot):
    t = await complete_topup(int(c.data.split(":")[1]), c.from_user.id)
    if not t:
        await c.answer("Allaqachon ishlangan.", show_alert=True)
        return await mark_msg(c, "ℹ️ Allaqachon ishlangan.")
    await audit(c.from_user.id, "balans to'ldirildi", f"{t['code']} • {fmt(t['amount'])} so'm • user {t['user_id']}")
    await c.answer("✅ Tasdiqlandi")
    await mark_msg(c, f"✅ Tasdiqladi: {esc(c.from_user.full_name)}")
    await tell(bot, t["user_id"], f"✅ Hisobingiz <b>{fmt(t['amount'])} so'm</b> ga to'ldirildi!")


@admin_r.callback_query(F.data.startswith("tr:"))
async def topup_reject(c: CallbackQuery, bot: Bot):
    tid = int(c.data.split(":")[1])
    _, rc = await db.execute("UPDATE topups SET status='cancelled', done_at=datetime('now','+5 hours') "
                             "WHERE id=? AND status IN ('checking','awaiting_check')", (tid,))
    if not rc:
        await c.answer("Allaqachon ishlangan.", show_alert=True)
        return await mark_msg(c, "ℹ️ Allaqachon ishlangan.")
    t = await db.fetchone("SELECT * FROM topups WHERE id=?", (tid,))
    await audit(c.from_user.id, "to'ldirish rad", t["code"])
    await c.answer("Rad etildi")
    await mark_msg(c, f"❌ Rad etdi: {esc(c.from_user.full_name)}")
    await tell(bot, t["user_id"], f"❌ To'ldirish {t['code']} tasdiqlanmadi. Yordam: {esc(await S('support'))}")


# ------------------------------------------------------------ 🌐 Midasbuy (admin)
async def midas_view():
    rows = await db.fetchall("SELECT * FROM packages WHERE kind='midas' ORDER BY uc")
    st = await db.fetchone(
        "SELECT COALESCE(SUM(status='done'),0) d, COALESCE(SUM(status IN ('checking','processing')),0) p, "
        "COALESCE(SUM(CASE WHEN status='done' THEN final-cost ELSE 0 END),0) f "
        "FROM orders WHERE kind='midas'")
    b = InlineKeyboardBuilder()
    lines = []
    for p in rows:
        lines.append(f"{'🟢' if p['active'] else '⚪️'} <b>{esc(p['name'])}</b> • {p['uc']} UC • "
                     f"{fmt(p['price'])} / tannarx {fmt(p['cost'])} (foyda {fmt(p['price'] - p['cost'])})")
        b.button(text=f"{'⏸' if p['active'] else '▶️'} {p['name']}", callback_data=f"mpt:{p['id']}")
        b.button(text="🗑", callback_data=f"mpd:{p['id']}")
        b.button(text="✏️", callback_data=f"edm:pkg:{p['id']}")
    b.button(text="➕ Midasbuy paket qo'shish", callback_data="mpa")
    b.button(text="📈 Narxni ommaviy o'zgartirish", callback_data="bp")
    b.adjust(*([3] * len(rows)), 1, 1)
    text = ("🌐 <b>Midasbuy paketlari</b>\n\n" + ("\n".join(lines) if lines else "Paket yo'q.") +
            f"\n\n✅ Bajarilgan: {st['d']} ta | ⏳ Jarayonda: {st['p']} ta | 💰 Foyda: {fmt(st['f'])} so'm\n"
            "🔒 Mijozlar Midasbuy haqida bilmaydi: bu paketlar ularga oddiy UC paket sifatida, ro'yxat pastida ko'rinadi.")
    return text, b.as_markup()


async def midas_home():
    url = await S("midasbuy_url")
    b = InlineKeyboardBuilder()
    if url.startswith("https://"):
        b.button(text="🌐 Midasbuy'ga kirish", web_app=WebAppInfo(url=url))
    if url.startswith(("http://", "https://")):
        b.button(text="🔗 Brauzerda ochish", url=url)
    b.button(text="📦 Midasbuy paketlari", callback_data="mpl")
    b.button(text="✏️ Sayt manzilini o'zgartirish", callback_data="set:midasbuy_url")
    b.adjust(1)
    return (f"🌐 <b>Midasbuy</b>\n\nSaytga kirib, mijoz PUBG ID siga UC ni tashlang.\n"
            f"🔗 Manzil: {esc(url)}\n\n🔒 Bu bo'lim faqat adminlarga ko'rinadi."), b.as_markup()


@admin_r.message(F.text == T.A_midasbuy)
async def midas_menu(m: Message):
    text, kb = await midas_home()
    await m.answer(text, reply_markup=kb, disable_web_page_preview=True)


@admin_r.callback_query(F.data == "mpl")
async def midas_pkgs(c: CallbackQuery):
    await c.answer()
    text, kb = await midas_view()
    await c.message.answer(text, reply_markup=kb)


@admin_r.callback_query(F.data.startswith(("mpt:", "mpd:")))
async def midas_pkg_manage(c: CallbackQuery):
    act, pid = c.data.split(":")
    if act == "mpt":
        await db.execute("UPDATE packages SET active=1-active WHERE id=? AND kind='midas'", (int(pid),))
    else:
        await db.execute("DELETE FROM packages WHERE id=? AND kind='midas'", (int(pid),))
    await c.answer("Bajarildi")
    text, kb = await midas_view()
    await safe_edit(c.message, text, kb)


@admin_r.callback_query(F.data == "mpa")
async def midas_pkg_add(c: CallbackQuery, state: FSMContext):
    await c.answer()
    await state.set_state(Adm.mp_name)
    await c.message.answer("🌐 Midasbuy paket <b>nomini</b> yozing (masalan: 60 UC):", reply_markup=cancel_kb())


@admin_r.message(StateFilter(Adm.mp_name))
async def midas_pkg_name(m: Message, state: FSMContext):
    await state.update_data(name=(m.text or "").strip()[:40])
    await state.set_state(Adm.mp_uc)
    await m.answer("💎 UC miqdori (raqam):")


@admin_r.message(StateFilter(Adm.mp_uc))
async def midas_pkg_uc(m: Message, state: FSMContext):
    n = parse_int(m.text or "")
    if not n or n <= 0:
        return await m.answer("❗️ Musbat son kiriting.")
    await state.update_data(uc=n)
    await state.set_state(Adm.mp_price)
    await m.answer("💵 <b>Sotuv narxi</b> (so'm):")


@admin_r.message(StateFilter(Adm.mp_price))
async def midas_pkg_price(m: Message, state: FSMContext):
    n = parse_int(m.text or "")
    if not n or n <= 0:
        return await m.answer("❗️ Musbat son kiriting.")
    await state.update_data(price=n)
    await state.set_state(Adm.mp_cost)
    await m.answer("🏷 <b>Tannarx</b> (siz UC ni qanchaga olasiz, so'm):")


@admin_r.message(StateFilter(Adm.mp_cost))
async def midas_pkg_cost(m: Message, state: FSMContext):
    n = parse_int(m.text or "")
    if n is None or n < 0:
        return await m.answer("❗️ Son kiriting.")
    d = await state.get_data()
    await db.execute("INSERT INTO packages(name,uc,price,cost,kind) VALUES(?,?,?,?,'midas')",
                     (d["name"], d["uc"], d["price"], n))
    await state.clear()
    await m.answer(f"✅ Midasbuy paket qo'shildi: <b>{esc(d['name'])}</b> — {d['uc']} UC\n"
                   f"Sotuv: {fmt(d['price'])} | Tannarx: {fmt(n)} | Foyda: <b>{fmt(d['price'] - n)} so'm</b>",
                   reply_markup=admin_kb())


# ------------------------------------------------------------ 🌐 Midasbuy buyurtma oqimi
async def _midas_processing_kb(oid: int) -> InlineKeyboardMarkup:
    kb = InlineKeyboardBuilder()
    kb.button(text="✅ UC tashladim", callback_data=f"md:{oid}")
    kb.button(text="❌ Bekor qilish", callback_data=f"mr:{oid}")
    murl = await S("midasbuy_url")
    if murl.startswith(("http://", "https://")):
        kb.button(text="🌐 Midasbuy'ga kirish", url=murl)
    kb.adjust(1)
    return kb.as_markup()


@admin_r.callback_query(F.data.startswith("ma:"))
async def midas_pay_ok(c: CallbackQuery, bot: Bot):
    oid = int(c.data.split(":")[1])
    if not await claim_guard(c, oid):
        return
    o = await midas_accept(oid, c.from_user.id)
    if not o:
        cur = await db.fetchone("SELECT status FROM orders WHERE id=?", (oid,))
        if cur and cur["status"] == "processing":   # boshqa admin allaqachon tasdiqlagan
            await c.answer("To'lov allaqachon tasdiqlangan. UC ni tashlab, tugmani bosing.", show_alert=True)
            return await mark_msg(c, "⚙️ To'lov tasdiqlangan.", await _midas_processing_kb(oid))
        await c.answer("Bu buyurtma allaqachon ishlangan.", show_alert=True)
        return await mark_msg(c, "ℹ️ Allaqachon ishlangan.")
    await audit(c.from_user.id, "midas to'lov ok", o["code"])
    await c.answer("✅ To'lov tasdiqlandi")
    await mark_msg(c, f"✅ To'lovni tasdiqladi: {esc(c.from_user.full_name)}\n"
                      f"⏳ Endi UC ni tashlab, «UC tashladim» ni bosing.", await _midas_processing_kb(oid))
    await tell(bot, o["user_id"],
               f"✅ <b>To'lovingiz tasdiqlandi!</b>\n🧾 Buyurtma <b>{o['code']}</b>\n"
               f"💎 {o['uc']} UC — PUBG ID <code>{esc(o['pubg_id'])}</code>\n\n"
               f"⏳ <b>5 daqiqa ichida</b> UC ni tashlab beramiz. Iltimos, biroz kutib turing.")


@admin_r.callback_query(F.data.startswith("md:"))
async def midas_uc_sent(c: CallbackQuery, bot: Bot):
    oid = int(c.data.split(":")[1])
    if not await claim_guard(c, oid):
        return
    info = await complete_order(oid, c.from_user.id, from_status="processing")
    if not info:
        await c.answer("Bu buyurtma allaqachon ishlangan.", show_alert=True)
        return await mark_msg(c, "ℹ️ Allaqachon ishlangan.")
    await c.answer("✅ Bajarildi")
    o = info["order"]
    await audit(c.from_user.id, "midas UC tashlandi", f"{o['code']} • {o['uc']} UC • {fmt(o['final'])} so'm")
    await mark_msg(c, f"✅ UC tashlandi (tasdiqladi: {esc(c.from_user.full_name)})")
    txt = (f"✅ <b>UC tushdi! Buyurtma {o['code']} bajarildi.</b>\n"
           f"💎 {o['uc']} UC PUBG ID <code>{esc(o['pubg_id'])}</code> ga tushdi.\n\n"
           f"🙏 <b>Xaridingiz uchun rahmat!</b>\n"
           f"🎁 Sizga bonus: <b>+{info['cashback']}</b> 🪙 tanga\n💎 Jami xaridingiz: {info['total_uc']} UC")
    if info["became_vip"]:
        txt += "\n\n🎉 <b>Tabriklaymiz! Siz endi VIP maqomdasiz!</b> Keshbek ×2 va oylik TOP-1 o'yinida ishtirok."
    if info.get("level_up"):
        txt += f"\n\n🏅 Yangi daraja: <b>{info['level_up']}</b> — keshbekingiz oshdi!"
    await tell(bot, o["user_id"], txt, reply_markup=review_kb(o["id"]))
    if info["ref_id"]:
        await tell(bot, info["ref_id"], f"🎁 Taklif qilgan do'stingiz 1-xaridini qildi! +{info['ref_bonus']} 🪙")


@admin_r.callback_query(F.data.startswith("mr:"))
async def midas_reject(c: CallbackQuery, bot: Bot):
    oid = int(c.data.split(":")[1])
    if not await claim_guard(c, oid):
        return
    o = await cancel_order(oid, ("awaiting_check", "checking", "processing"))
    if not o:
        await c.answer("Bu buyurtma allaqachon ishlangan.", show_alert=True)
        return await mark_msg(c, "ℹ️ Allaqachon ishlangan.")
    await audit(c.from_user.id, "midas bekor", o["code"])
    await c.answer("Bekor qilindi")
    await mark_msg(c, f"❌ Bekor qildi: {esc(c.from_user.full_name)}")
    await tell(bot, o["user_id"], f"❌ <b>Buyurtma {o['code']} bekor qilindi.</b>\n"
                                  f"Tanga/kupon{'/balans' if o['pay_method'] == 'balance' else ''} qaytarildi. "
                                  f"Savol bo'lsa: {esc(await S('support'))}")


@admin_r.message(F.text == T.A_PENDING)
async def pending(m: Message, bot: Bot):
    os_ = await db.fetchall("SELECT id FROM orders WHERE status IN ('checking','processing') ORDER BY id LIMIT 15")
    ts = await db.fetchall("SELECT id FROM topups WHERE status='checking' ORDER BY id LIMIT 15")
    if not os_ and not ts:
        return await m.answer("✅ Kutilayotgan so'rovlar yo'q.")
    for r in os_:
        o = await db.fetchone("SELECT check_file_id FROM orders WHERE id=?", (r["id"],))
        await _resend(bot, m.from_user.id, "order", r["id"])
    for r in ts:
        await _resend(bot, m.from_user.id, "topup", r["id"])


async def _resend(bot: Bot, admin_id: int, kind: str, rid: int):
    await notify_admins(bot, kind, rid, only=admin_id)


# ------------------------------------------------------------ Foydalanuvchi boshqaruvi
def user_card(u) -> tuple:
    text = (f"👤 <b>{esc(u['full_name'])}</b> (<code>{u['id']}</code>)\n@{esc(u['username'] or '—')}\n"
            f"💰 Balans: {fmt(u['balance'])} so'm | 🪙 {fmt(u['coins'])}\n"
            f"💎 Jami: {fmt(u['total_uc'])} UC | {'VIP' if u['is_vip'] else 'Oddiy'}\n"
            f"🚫 Ban: {'ha' if u['banned'] else 'yo`q'} | 📅 {u['joined_at'][:10]}")
    kb = InlineKeyboardBuilder()
    kb.button(text="💰 Balans ±", callback_data=f"um:bal:{u['id']}")
    kb.button(text="🪙 Tanga ±", callback_data=f"um:coin:{u['id']}")
    kb.button(text="✅ Unban" if u["banned"] else "🚫 Ban", callback_data=f"um:ban:{u['id']}")
    kb.button(text="💎 VIP almashtirish", callback_data=f"um:vip:{u['id']}")
    kb.adjust(2)
    return text, kb.as_markup()


@admin_r.message(F.text == T.A_USER)
async def user_start(m: Message, state: FSMContext):
    await state.set_state(Adm.user_id)
    await m.answer("👤 Foydalanuvchi <b>ID</b> sini yuboring:", reply_markup=cancel_kb())


@admin_r.message(StateFilter(Adm.user_id))
async def user_find(m: Message, state: FSMContext):
    n = parse_int(m.text or "")
    u = await get_user(n) if n else None
    if not u:
        return await m.answer("❌ Foydalanuvchi topilmadi. ID ni tekshiring.")
    await state.clear()
    text, kb = user_card(u)
    await m.answer(text, reply_markup=kb)
    await m.answer("Admin panel", reply_markup=admin_kb())


@admin_r.callback_query(F.data.startswith("um:"))
async def user_action(c: CallbackQuery, state: FSMContext):
    _, act, uid = c.data.split(":")
    uid = int(uid)
    if act in ("bal", "coin"):
        await c.answer()
        await state.set_state(Adm.user_bal if act == "bal" else Adm.user_coin)
        await state.update_data(uid=uid)
        return await c.message.answer("Qo'shish uchun musbat, ayirish uchun manfiy son yuboring (masalan 5000 yoki -2000):",
                                      reply_markup=cancel_kb())
    if act == "ban":
        await db.execute("UPDATE users SET banned=1-banned WHERE id=?", (uid,))
    elif act == "vip":
        await db.execute("UPDATE users SET is_vip=1-is_vip WHERE id=?", (uid,))
    await c.answer("Bajarildi")
    text, kb = user_card(await get_user(uid))
    await safe_edit(c.message, text, kb)


@admin_r.message(StateFilter(Adm.user_bal))
@admin_r.message(StateFilter(Adm.user_coin))
async def user_adjust(m: Message, state: FSMContext, bot: Bot):
    n = parse_int(m.text or "")
    if n is None:
        return await m.answer("❗️ Son kiriting.")
    d = await state.get_data()
    is_bal = (await state.get_state()) == Adm.user_bal.state
    col = "balance" if is_bal else "coins"
    await db.execute(f"UPDATE users SET {col}=MAX({col}+?,0) WHERE id=?", (n, d["uid"]))
    await state.clear()
    u = await get_user(d["uid"])
    text, kb = user_card(u)
    await m.answer("✅ Yangilandi.\n\n" + text, reply_markup=kb)
    await m.answer("Admin panel", reply_markup=admin_kb())
    await tell(bot, d["uid"], f"ℹ️ Hisobingiz o'zgartirildi: {'balans' if is_bal else 'tanga'} {n:+d}")


# ------------------------------------------------------------ VIP o'yini
@admin_r.message(F.text == T.A_LOTTERY)
async def lottery_menu(m: Message):
    cur = now_tz().strftime("%Y-%m")
    prev = (now_tz().replace(day=1) - timedelta(days=1)).strftime("%Y-%m")
    kb = InlineKeyboardBuilder()
    kb.button(text=f"🎰 Joriy oy ({cur})", callback_data=f"lot:{cur}")
    kb.button(text=f"🎰 O'tgan oy ({prev})", callback_data=f"lot:{prev}")
    kb.adjust(1)
    await m.answer("🎰 <b>VIP TOP-1 o'yini</b>\n\nTanlangan oyda eng ko'p UC olgan VIP g'olib bo'ladi va "
                   "tasodifiy <b>30 yoki 60 UC</b> yutadi.\n(Har oyning 1-kuni avtomatik ham o'tkaziladi.)",
                   reply_markup=kb.as_markup())


@admin_r.callback_query(F.data.startswith("lot:"))
async def lottery_run(c: CallbackQuery, bot: Bot):
    await c.answer("O'yin boshlandi...")
    await c.message.answer(await run_lottery(bot, c.data.split(":")[1]))


# ------------------------------------------------------------ Tizim sozlamalari
async def settings_view():
    b = InlineKeyboardBuilder()
    lines = []
    for k, (label, _) in SETTING_META.items():
        lines.append(f"• {label}: <b>{esc((await S(k))[:60]) or '—'}</b>")
        b.button(text=label[:30], callback_data=f"set:{k}")
    b.adjust(1)
    return "🛠 <b>Tizim sozlamalari</b>\n\n" + "\n".join(lines) + "\n\nO'zgartirish uchun tanlang:", b.as_markup()


@admin_r.message(F.text == T.A_SETTINGS)
async def settings_menu(m: Message):
    text, kb = await settings_view()
    await m.answer(text, reply_markup=kb)


@admin_r.callback_query(F.data.startswith("set:"))
async def setting_pick(c: CallbackQuery, state: FSMContext):
    key = c.data.split(":")[1]
    if key not in SETTING_META:
        return await c.answer("Noma'lum")
    await c.answer()
    await state.set_state(Adm.setting_val)
    await state.update_data(key=key)
    await c.message.answer(f"✏️ <b>{SETTING_META[key][0]}</b>\nHozirgi: <b>{esc((await S(key))[:400]) or '—'}</b>\nYangi qiymat:",
                           reply_markup=cancel_kb())


@admin_r.message(StateFilter(Adm.setting_val))
async def setting_save(m: Message, state: FSMContext):
    d = await state.get_data()
    key = d["key"]
    val = (m.text or "").strip()
    if key == "midasbuy_url" and not re.match(r"https?://\S+$", val):
        return await m.answer("❗️ Manzil http:// yoki https:// bilan boshlanishi kerak.")
    if key == "welcome_text":
        if val == "-":
            val = ""
        elif len(val) > 1500:
            return await m.answer("❗️ Matn 1500 belgidan oshmasin.")
        else:
            try:   # HTML xato bo'lsa /start hammada buzilmasligi uchun avval ko'rib chiqamiz
                await m.answer("👀 <b>Ko'rinishi:</b>\n\n" + val.replace("{name}", esc(m.from_user.full_name)))
            except TelegramBadRequest as e:
                return await m.answer(f"❌ Matnda HTML xatosi (teg yopilmagan bo'lishi mumkin):\n<code>{esc(e)}</code>")
    if SETTING_META[key][1] == "int":
        n = parse_int(val)
        if n is None or n < 0:
            return await m.answer("❗️ 0 yoki musbat son kiriting.")
        val = str(n)
    await db.execute("UPDATE settings SET value=? WHERE key=?", (val, key))
    await state.clear()
    await m.answer("✅ Saqlandi.", reply_markup=admin_kb())


# ------------------------------------------------------------ 🧹 Ma'lumotlarni tozalash
CLEAN_META = {"d": ("Kunlik", 10, "%Y-%m-%d", "bugungi"),
              "m": ("Oylik", 7, "%Y-%m", "shu oydagi"),
              "y": ("Yillik", 4, "%Y", "shu yildagi")}


def clean_cond(p: str):
    _, n, f, _ = CLEAN_META[p]
    return (f"status IN ('done','cancelled') AND substr(COALESCE(done_at,created_at),1,{n})=?",
            (now_tz().strftime(f),))


@admin_r.message(F.text == T.A_CLEAN)
async def clean_menu(m: Message):
    kb = InlineKeyboardBuilder()
    kb.button(text="🗓 Kunlik (bugun)", callback_data="cl:d")
    kb.button(text="📆 Oylik (shu oy)", callback_data="cl:m")
    kb.button(text="🗃 Yillik (shu yil)", callback_data="cl:y")
    kb.adjust(1)
    await m.answer("🧹 <b>Ma'lumotlarni tozalash</b>\n\nTanlangan davrdagi <b>yakunlangan</b> (bajarilgan/bekor qilingan) "
                   "buyurtma va to'ldirishlar o'chiriladi — statistika va hisobot nolga tushadi.\n\n"
                   "✅ Tegilmaydi: foydalanuvchilar, balans, tanga, VIP, kutilayotgan buyurtmalar.\n"
                   "💾 O'chirishdan oldin o'chiriladigan ma'lumot CSV fayl qilib sizga yuboriladi.",
                   reply_markup=kb.as_markup())


@admin_r.callback_query(F.data.startswith("cl:"))
async def clean_ask(c: CallbackQuery):
    p = c.data.split(":")[1]
    if p not in CLEAN_META:
        return await c.answer("Noma'lum")
    cond, args = clean_cond(p)
    no = (await db.fetchone(f"SELECT COUNT(*) n FROM orders WHERE {cond}", args))["n"]
    nt = (await db.fetchone(f"SELECT COUNT(*) n FROM topups WHERE {cond}", args))["n"]
    await c.answer()
    if not no and not nt:
        return await c.message.answer("ℹ️ Bu davr uchun tozalanadigan ma'lumot yo'q.")
    label, _, _, word = CLEAN_META[p]
    warn = ("\n⚠️ VIP o'yini hali o'tkazilmagan oyning ma'lumotini o'chirsangiz, o'yin g'olibi aniqlanmaydi."
            if p in ("m", "y") else "")
    kb = InlineKeyboardBuilder()
    kb.button(text="✅ Ha, tozalash", callback_data=f"clok:{p}")
    kb.button(text="❌ Yo'q", callback_data="clno")
    kb.adjust(2)
    await c.message.answer(f"🧹 <b>{label} tozalash</b>\n\nO'chiriladi ({word}): 🧾 {no} ta buyurtma, 💰 {nt} ta to'ldirish."
                           f"{warn}\n\nDavom etamizmi?", reply_markup=kb.as_markup())


@admin_r.callback_query(F.data == "clno")
async def clean_no(c: CallbackQuery):
    await c.answer("Bekor qilindi")
    await safe_edit(c.message, "❌ Tozalash bekor qilindi.")


@admin_r.callback_query(F.data.startswith("clok:"))
async def clean_do(c: CallbackQuery):
    p = c.data.split(":")[1]
    if p not in CLEAN_META:
        return await c.answer("Noma'lum")
    await c.answer("Tozalanmoqda...")
    cond, args = clean_cond(p)
    label = CLEAN_META[p][0]
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["turi", "id", "kod", "user_id", "pubg_id", "paket", "uc", "narx", "to'langan", "tannarx",
                "holat", "usul", "sana"])
    for r in await db.fetchall(f"SELECT * FROM orders WHERE {cond} ORDER BY id", args):
        w.writerow(["buyurtma", r["id"], r["code"], r["user_id"], r["pubg_id"], r["pkg_name"], r["uc"], r["price"],
                    r["final"], r["cost"], r["status"], r["pay_method"], r["created_at"]])
    for r in await db.fetchall(f"SELECT * FROM topups WHERE {cond} ORDER BY id", args):
        w.writerow(["to'ldirish", r["id"], r["code"], r["user_id"], "", "", "", "", r["amount"], "",
                    r["status"], "", r["created_at"]])
    stamp = now_tz().strftime("%Y%m%d_%H%M")
    await c.message.answer_document(BufferedInputFile(buf.getvalue().encode("utf-8-sig"),
                                                      filename=f"tozalangan_{p}_{stamp}.csv"),
                                    caption="💾 O'chirilgan ma'lumotlar nusxasi")
    _, no = await db.execute(f"DELETE FROM orders WHERE {cond}", args)
    _, nt = await db.execute(f"DELETE FROM topups WHERE {cond}", args)
    await safe_edit(c.message, f"✅ <b>{label} tozalash bajarildi.</b>\n🧾 {no} ta buyurtma, 💰 {nt} ta to'ldirish o'chirildi.")


# ------------------------------------------------------------ Backup
@admin_r.message(F.text == T.A_BACKUP)
async def backup(m: Message, bot: Bot):
    cloud = await tg_backup(bot, force=True)
    if cloud:
        await m.answer("☁️ Telegram bulutiga ham saqlandi (pin qilingan xabar). Server almashsa bot o'zi tiklaydi.")
    elif BK["locked"]:
        await m.answer("⚠️ Avto-zaxira o'chirilgan (tiklash xatosi). Quyidagi faylni qo'lda saqlab qo'ying.")
    tmp = Path(DB_PATH).parent / f"backup_{int(time.time())}.db"
    try:
        async with db.lock:
            await db.conn.execute("VACUUM INTO ?", (str(tmp),))
        data = tmp.read_bytes()
        await m.answer_document(BufferedInputFile(data, filename=f"pro_uc_bot_{now_tz():%Y%m%d_%H%M}.db"),
                                caption=f"💾 Baza zaxira nusxasi.\n📍 Hozirgi joy: {DB_PATH}\nYangi serverda shu faylni DB_PATH ga qo'ying "
                                        f"yoki botga «/restore» izohi bilan yuboring — hamma narsa saqlanadi.")
    except Exception as e:
        await m.answer(f"❌ Zaxira xatosi: {esc(e)}")
    finally:
        tmp.unlink(missing_ok=True)


# ============================================================ 🛡 ROLLAR, JURNAL, BAND QILISH, XATO HISOBOTI
async def admin_role(uid: int) -> Optional[str]:
    if uid == SUPER_ADMIN_ID:
        return "super"
    r = await db.fetchone("SELECT role FROM admins WHERE id=?", (uid,))
    return r["role"] if r else None


OP_TEXTS = {T.A_PENDING, T.A_REVIEWS, T.ADMIN}
OP_CB = ("oa:", "or:", "ma:", "md:", "mr:", "ocl:", "orl:", "ta:", "tr:", "rvp:", "rvr:", "sr:")


class RoleMiddleware(BaseMiddleware):
    """Operator roli: faqat chek/buyurtma, sharx va murojaatlar bilan ishlaydi."""

    async def __call__(self, handler, event, data):
        uid = event.from_user.id
        if await admin_role(uid) != "operator":
            return await handler(event, data)
        if isinstance(event, Message):
            txt = event.text or ""
            st = data.get("state")
            cur = (await st.get_state()) if st else None
            if txt in OP_TEXTS or txt.startswith("/admin") or (cur and str(cur).endswith("support_reply")):
                return await handler(event, data)
            return await event.answer("🚫 Bu bo'lim faqat to'liq huquqli adminlar uchun.")
        if (event.data or "").startswith(OP_CB):
            return await handler(event, data)
        return await event.answer("🚫 Bu amal faqat to'liq huquqli adminlar uchun.", show_alert=True)


async def audit(admin_id: int, action: str, detail: str = ""):
    try:
        await db.execute("INSERT INTO audit(admin_id,action,detail) VALUES(?,?,?)",
                         (admin_id, action[:40], (detail or "")[:300]))
    except Exception:
        log.exception("audit xatosi")


@admin_r.message(F.text == T.A_AUDIT)
async def audit_view(m: Message):
    rows = await db.fetchall("SELECT a.*, u.full_name FROM audit a LEFT JOIN users u ON u.id=a.admin_id "
                             "ORDER BY a.id DESC LIMIT 40")
    if not rows:
        return await m.answer("📜 Jurnal hozircha bo'sh.")
    out, total = [], 0
    for r in rows:
        line = (f"<code>{r['created_at'][5:16]}</code> • {esc((r['full_name'] or str(r['admin_id']))[:14])} • "
                f"<b>{esc(r['action'])}</b> {esc(r['detail'] or '')}")
        total += len(line) + 1
        if total > 3600:
            break
        out.append(line)
    await m.answer("📜 <b>Admin harakatlari jurnali</b> (oxirgi)\n\n" + "\n".join(out))


async def claim_order(oid: int, admin_id: int):
    """Buyurtmani adminga band qiladi. (True, id) yoki (False, band qilgan admin id)."""
    _, rc = await db.execute(
        "UPDATE orders SET claimed_by=? WHERE id=? AND status IN ('checking','processing') "
        "AND (claimed_by IS NULL OR claimed_by=?)", (admin_id, oid, admin_id))
    if rc:
        return True, admin_id
    o = await db.fetchone("SELECT status, claimed_by FROM orders WHERE id=?", (oid,))
    if not o or o["status"] not in ("checking", "processing"):
        return True, None   # allaqachon yakunlangan — keyingi tekshiruvni asosiy handler qiladi
    return False, o["claimed_by"]


async def claim_guard(c: CallbackQuery, oid: int) -> bool:
    ok, who = await claim_order(oid, c.from_user.id)
    if ok:
        return True
    u = await get_user(who) if who else None
    await c.answer(f"🔒 Bu buyurtmani {u['full_name'] if u else who} bajaryapti.", show_alert=True)
    return False


async def order_kb(o, rid: int, mode: str, multi: bool):
    if mode == "locked":
        return None
    midas = o["kind"] == "midas"
    rej = f"mr:{rid}" if midas else f"or:{rid}"
    kb = InlineKeyboardBuilder()
    if mode == "claim":
        kb.button(text="🙋 Men bajaraman", callback_data=f"ocl:{rid}")
        kb.button(text="❌ Rad etish", callback_data=rej)
    elif not midas:
        kb.button(text="✅ UC yuborildi — Tasdiqlash", callback_data=f"oa:{rid}")
        kb.button(text="❌ Rad etish", callback_data=f"or:{rid}")
    elif o["status"] == "processing":
        kb.button(text="✅ UC tashladim", callback_data=f"md:{rid}")
        kb.button(text="❌ Bekor qilish", callback_data=f"mr:{rid}")
    else:
        kb.button(text="✅ To'lovni tasdiqlash", callback_data=f"ma:{rid}")
        kb.button(text="❌ Rad etish", callback_data=f"mr:{rid}")
    if midas:
        murl = await S("midasbuy_url")
        if murl.startswith(("http://", "https://")):
            kb.button(text="🌐 Midasbuy'ga kirish", url=murl)
    if mode == "act" and multi:
        kb.button(text="🔓 Bo'shatish", callback_data=f"orl:{rid}")
    kb.adjust(1)
    return kb.as_markup()


@admin_r.callback_query(F.data.startswith("ocl:"))
async def order_claim(c: CallbackQuery):
    oid = int(c.data.split(":")[1])
    o = await db.fetchone("SELECT * FROM orders WHERE id=?", (oid,))
    if not o or o["status"] not in ("checking", "processing"):
        await c.answer("Bu buyurtma allaqachon ishlangan.", show_alert=True)
        return await mark_msg(c, "ℹ️ Allaqachon ishlangan.")
    ok, who = await claim_order(oid, c.from_user.id)
    if not ok:
        u = await get_user(who) if who else None
        await c.answer(f"🔒 Bu buyurtmani {u['full_name'] if u else who} bajaryapti.", show_alert=True)
        return await mark_msg(c, f"🔒 {esc(u['full_name'] if u else who)} bajaryapti.")
    await audit(c.from_user.id, "buyurtma band", o["code"])
    await c.answer("🙋 Buyurtma sizga biriktirildi")
    o = await db.fetchone("SELECT * FROM orders WHERE id=?", (oid,))
    await mark_msg(c, f"🙋 Bajaruvchi: {esc(c.from_user.full_name)}", await order_kb(o, oid, "act", True))


@admin_r.callback_query(F.data.startswith("orl:"))
async def order_release(c: CallbackQuery):
    oid = int(c.data.split(":")[1])
    _, rc = await db.execute("UPDATE orders SET claimed_by=NULL WHERE id=? AND claimed_by=? "
                             "AND status IN ('checking','processing')", (oid, c.from_user.id))
    if not rc:
        return await c.answer("Bu buyurtma sizga biriktirilmagan.", show_alert=True)
    o = await db.fetchone("SELECT * FROM orders WHERE id=?", (oid,))
    await audit(c.from_user.id, "buyurtma bo'shatildi", o["code"])
    await c.answer("🔓 Bo'shatildi")
    await mark_msg(c, f"🔓 {esc(c.from_user.full_name)} bo'shatdi.", await order_kb(o, oid, "claim", True))


async def dup_text(table: str, rid: int, cuid, chash) -> str:
    if not (cuid or chash):
        return ""
    for tb in ("orders", "topups"):
        r = await db.fetchone(
            f"SELECT code, user_id FROM {tb} WHERE id!=? AND ((check_uid IS NOT NULL AND check_uid=?) "
            f"OR (check_hash IS NOT NULL AND check_hash=?)) LIMIT 1", (rid if tb == table else -1, cuid, chash))
        if r:
            return f"⚠️ <b>Bu chek avval ham yuborilgan:</b> {esc(r['code'])} (ID <code>{r['user_id']}</code>) — ehtiyot bo'ling!"
    return ""


async def risk_lines(uid: int, pubg: str) -> str:
    out = []
    r = await db.fetchone("SELECT COUNT(*) n FROM orders WHERE user_id=? AND status='cancelled' AND check_at IS NOT NULL "
                          "AND done_at > datetime('now','+5 hours','-14 days')", (uid,))
    if r["n"] >= 2:
        out.append(f"⚠️ Oxirgi 14 kunda {r['n']} ta buyurtmasi rad etilgan")
    r = await db.fetchone("SELECT COUNT(DISTINCT user_id) n FROM orders WHERE pubg_id=? "
                          "AND created_at > datetime('now','+5 hours','-7 days')", (pubg,))
    if r["n"] >= 3:
        out.append(f"⚠️ Bu PUBG ID ni {r['n']} ta turli akkaunt ishlatgan (7 kun)")
    return "".join("\n" + x for x in out)


async def card_line(card_id) -> str:
    if not card_id:
        return ""
    c = await db.fetchone("SELECT * FROM cards WHERE id=?", (card_id,))
    return f"\n💳 Qabul qiladigan karta: …{c['number'][-4:]} ({esc(c['owner'])})" if c else ""


async def order_extras(o) -> str:
    s = await card_line(o["card_id"]) + await risk_lines(o["user_id"], o["pubg_id"])
    d = await dup_text("orders", o["id"], o["check_uid"], o["check_hash"])
    return s + (("\n" + d) if d else "")


async def topup_extras(t) -> str:
    s = await card_line(t["card_id"])
    d = await dup_text("topups", t["id"], t["check_uid"], t["check_hash"])
    return s + (("\n" + d) if d else "")


_err_last: dict = {}


async def report_error(bot: Bot, where: str, exc) -> None:
    """Xatoni super adminga yuboradi (bir xil xato 5 daqiqada bir marta)."""
    try:
        key = f"{where}:{type(exc).__name__}:{str(exc)[:40]}"
        if time.time() - _err_last.get(key, 0) < 300:
            return
        _err_last[key] = time.time()
        tb = "".join(traceback.format_exception(type(exc), exc, exc.__traceback__))[-900:]
        for aid in ([SUPER_ADMIN_ID] if SUPER_ADMIN_ID else await admin_ids()):
            await bot.send_message(aid, f"⚠️ <b>Bot xatosi</b> ({esc(where)})\n<pre>{esc(tb)}</pre>")
    except Exception:
        log.exception("report_error xatosi")


async def on_error(event, bot: Bot):
    exc = getattr(event, "exception", None)
    log.error("Handler xatosi", exc_info=exc)
    if exc is not None:
        await report_error(bot, "handler", exc)
    try:   # foydalanuvchi jim qolmasin
        upd = getattr(event, "update", None)
        if upd is not None and getattr(upd, "message", None):
            await bot.send_message(upd.message.chat.id, "⚠️ Vaqtincha xatolik yuz berdi. Iltimos, qayta urinib ko'ring.")
        elif upd is not None and getattr(upd, "callback_query", None):
            await upd.callback_query.answer("⚠️ Vaqtincha xatolik. Qayta urinib ko'ring.", show_alert=True)
    except Exception:
        pass
    return True


@common_r.message(Command("ping"))
async def ping(m: Message):
    """Bot ishlayotganini tekshirish: /ping"""
    n = await db.fetchone("SELECT COUNT(*) c FROM users")
    await m.answer(f"🏓 pong — bot ishlayapti\n👥 Foydalanuvchilar: {n['c']}\n💾 Baza: <code>{esc(DB_PATH)}</code>")


# ============================================================ 🗓 TURNIR: vaqt, eslatma, g'olib mukofoti
def parse_when(text: str) -> Optional[datetime]:
    t = (text or "").strip().lower()
    now = now_tz().replace(tzinfo=None)
    try:
        m = re.fullmatch(r"(bugun|ertaga)?\s*(\d{1,2})[:.](\d{2})", t)
        if m:
            base = now.replace(hour=int(m.group(2)), minute=int(m.group(3)), second=0, microsecond=0)
            if m.group(1) == "ertaga":
                base += timedelta(days=1)
            elif not m.group(1) and base < now:
                base += timedelta(days=1)
            return base
        m = re.fullmatch(r"(\d{1,2})[./-](\d{1,2})(?:[./-](\d{2,4}))?\s+(\d{1,2})[:.](\d{2})", t)
        if m:
            y = int(m.group(3)) if m.group(3) else now.year
            if y < 100:
                y += 2000
            d = datetime(y, int(m.group(2)), int(m.group(1)), int(m.group(4)), int(m.group(5)))
            if not m.group(3) and d < now - timedelta(days=1):
                d = d.replace(year=y + 1)
            return d
    except ValueError:
        return None
    return None


async def tournament_reminders(bot: Bot):
    now = now_tz().replace(tzinfo=None)
    for t in await db.fetchall("SELECT * FROM tournaments WHERE status IN ('open','started') "
                               "AND start_at IS NOT NULL AND reminded=0"):
        try:
            st = datetime.strptime(t["start_at"], "%Y-%m-%d %H:%M")
        except ValueError:
            continue
        if not (-timedelta(minutes=5) <= st - now <= timedelta(minutes=15)):
            continue
        _, rc = await db.execute("UPDATE tournaments SET reminded=1 WHERE id=? AND reminded=0", (t["id"],))
        if not rc:
            continue
        room = f"🔑 <b>Xona ma'lumoti:</b>\n{esc(t['room_info'])}" if t["room_info"] else \
            "🔑 Xona kodi va nomi tez orada yuboriladi."
        for p in await db.fetchall("SELECT user_id FROM t_players WHERE tid=?", (t["id"],)):
            await tell(bot, p["user_id"], f"⏰ <b>{esc(t['title'])}</b> turniri tez orada boshlanadi ({esc(t['start_text'])})!\n\n{room}")
        if not t["room_info"]:
            for aid in await admin_ids():
                await tell(bot, aid, f"⏰ «{esc(t['title'])}» turniri boshlanishiga 15 daqiqa qoldi, "
                                     f"lekin <b>xona kodi yuborilmagan</b>. Turnir boshqaruvidan yuboring.")


@admin_r.callback_query(F.data.startswith("atnw:"))
async def atn_winner_pick(c: CallbackQuery):
    tid = int(c.data.split(":")[1])
    rows = await db.fetchall("SELECT p.user_id, p.won, u.full_name FROM t_players p LEFT JOIN users u ON u.id=p.user_id "
                             "WHERE p.tid=? ORDER BY p.joined_at LIMIT 40", (tid,))
    if not rows:
        return await c.answer("Hali ishtirokchi yo'q.", show_alert=True)
    await c.answer()
    b = InlineKeyboardBuilder()
    for r in rows:
        b.button(text=f"{(r['full_name'] or str(r['user_id']))[:22]}{' 🏆' if r['won'] else ''}",
                 callback_data=f"atnwp:{tid}:{r['user_id']}")
    b.adjust(2)
    await c.message.answer("🏆 G'olibni tanlang:", reply_markup=b.as_markup())


@admin_r.callback_query(F.data.startswith("atnwp:"))
async def atn_winner_sel(c: CallbackQuery, state: FSMContext):
    _, tid, uid = c.data.split(":")
    await c.answer()
    await state.set_state(Adm.t_win)
    await state.update_data(tid=int(tid), uid=int(uid))
    await c.message.answer("💰 Mukofot summasini yozing (so'm) — g'olib <b>balansiga</b> qo'shiladi.\n"
                           "Faqat xabar yuborish kerak bo'lsa (UC ni o'zingiz tashlasangiz) <code>0</code> yozing:",
                           reply_markup=cancel_kb())


@admin_r.message(StateFilter(Adm.t_win))
async def atn_winner_pay(m: Message, state: FSMContext, bot: Bot):
    d = await state.get_data()
    tid, uid = d["tid"], d["uid"]
    n = parse_int(m.text or "")
    if n is None or n < 0 or n > 100_000_000:
        return await m.answer("❗️ 0 yoki musbat son kiriting.")
    t = await db.fetchone("SELECT * FROM tournaments WHERE id=?", (tid,))
    p = await db.fetchone("SELECT 1 FROM t_players WHERE tid=? AND user_id=?", (tid, uid))
    if not (t and p):
        await state.clear()
        return await m.answer("Turnir yoki ishtirokchi topilmadi.", reply_markup=admin_kb())
    if n > 0:
        async with db.tx() as c:
            await c.execute("UPDATE users SET balance=balance+? WHERE id=?", (n, uid))
            await c.execute("UPDATE t_players SET won=won+? WHERE tid=? AND user_id=?", (n, tid, uid))
    await state.clear()
    await audit(m.from_user.id, "turnir mukofoti", f"#{tid} → {uid}: {n}")
    if n > 0:
        txt = (f"🏆 <b>Tabriklaymiz!</b> «{esc(t['title'])}» turnirida g'olib bo'ldingiz!\n"
               f"💰 Mukofot <b>{fmt(n)} so'm</b> balansingizga qo'shildi.")
    else:
        txt = (f"🏆 <b>Tabriklaymiz!</b> «{esc(t['title'])}» turnirida g'olib bo'ldingiz!\n"
               f"🎁 Mukofot ({esc(t['prize'])}) tez orada beriladi.")
    await tell(bot, uid, txt)
    await m.answer(f"✅ G'olibga yuborildi{f' va {fmt(n)} so`m balansga qo`shildi' if n else ''}.", reply_markup=admin_kb())


# ============================================================ 📈 NARXNI OMMAVIY O'ZGARTIRISH
def round_price(v: float, step: int = 500) -> int:
    return max(int(v / step + 0.5) * step, step)


BP_SCOPE = {"uc": ("kind='uc'", "Oddiy paketlar"), "midas": ("kind='midas'", "🌐 Midasbuy paketlari"),
            "all": ("1=1", "Hamma paketlar")}


@admin_r.callback_query(F.data == "bp")
async def bp_start(c: CallbackQuery):
    await c.answer()
    kb = InlineKeyboardBuilder()
    for k, (_, label) in BP_SCOPE.items():
        kb.button(text=label, callback_data=f"bps:{k}")
    kb.adjust(1)
    await c.message.answer("📈 <b>Narxni ommaviy o'zgartirish</b>\n\nQaysi paketlar uchun?", reply_markup=kb.as_markup())


@admin_r.callback_query(F.data.startswith("bps:"))
async def bp_scope(c: CallbackQuery, state: FSMContext):
    k = c.data.split(":")[1]
    if k not in BP_SCOPE:
        return await c.answer("Noma'lum")
    await c.answer()
    await state.set_state(Adm.bp_pct)
    await state.update_data(scope=k)
    await c.message.answer("📊 Necha <b>foizga</b> o'zgartiramiz?\nOshirish: <code>5</code> yoki <code>+5</code>, "
                           "kamaytirish: <code>-3</code> (−50 dan +100 gacha). Narx 500 so'mga yaxlitlanadi.",
                           reply_markup=cancel_kb())


async def bp_rows(scope: str):
    return await db.fetchall(f"SELECT * FROM packages WHERE {BP_SCOPE[scope][0]} ORDER BY kind, uc")


@admin_r.message(StateFilter(Adm.bp_pct))
async def bp_pct(m: Message, state: FSMContext):
    d = await state.get_data()
    raw = (m.text or "").replace(",", ".").replace("%", "").replace("+", "").strip()
    try:
        pct = float(raw)
    except ValueError:
        return await m.answer("❗️ Foizni son bilan yozing (masalan 5 yoki -3).")
    tenths = int(round(pct * 10))
    if tenths == 0 or not -500 <= tenths <= 1000:
        return await m.answer("❗️ Foiz −50 dan +100 gacha va 0 dan farqli bo'lsin.")
    rows = await bp_rows(d["scope"])
    if not rows:
        await state.clear()
        return await m.answer("Bu guruhda paket yo'q.", reply_markup=admin_kb())
    await state.clear()
    k = 1 + tenths / 1000
    lines = [f"• {esc(r['name'])}: {fmt(r['price'])} → <b>{fmt(round_price(r['price'] * k))}</b>" for r in rows[:15]]
    kb = InlineKeyboardBuilder()
    kb.button(text="✅ Faqat sotuv narxi", callback_data=f"bpo:{d['scope']}:{tenths}:0")
    kb.button(text="✅ Sotuv narxi + tannarx", callback_data=f"bpo:{d['scope']}:{tenths}:1")
    kb.button(text="❌ Bekor qilish", callback_data="bpn")
    kb.adjust(1)
    await m.answer(f"📈 <b>{BP_SCOPE[d['scope']][1]}: {tenths / 10:+g}%</b> ({len(rows)} ta paket)\n\n" + "\n".join(lines) +
                   ("\n…" if len(rows) > 15 else "") + "\n\nTasdiqlaysizmi?", reply_markup=kb.as_markup())
    await m.answer("👆 Tugmani tanlang.", reply_markup=admin_kb())


@admin_r.callback_query(F.data == "bpn")
async def bp_cancel(c: CallbackQuery):
    await c.answer("Bekor qilindi")
    await safe_edit(c.message, "❌ Narx o'zgartirilmadi.")


@admin_r.callback_query(F.data.startswith("bpo:"))
async def bp_apply(c: CallbackQuery):
    _, scope, tenths, with_cost = c.data.split(":")
    if scope not in BP_SCOPE:
        return await c.answer("Noma'lum")
    k = 1 + int(tenths) / 1000
    n = 0
    async with db.tx() as cx:
        cur = await cx.execute(f"SELECT * FROM packages WHERE {BP_SCOPE[scope][0]}")
        rows = await cur.fetchall()
        await cur.close()
        for r in rows:
            cost = round_price(r["cost"] * k) if with_cost == "1" and r["cost"] > 0 else r["cost"]
            await cx.execute("UPDATE packages SET price=?, cost=? WHERE id=?", (round_price(r["price"] * k), cost, r["id"]))
            n += 1
    await audit(c.from_user.id, "narx ommaviy", f"{BP_SCOPE[scope][1]} {int(tenths) / 10:+g}% ({n} ta)"
                                                 f"{' +tannarx' if with_cost == '1' else ''}")
    await c.answer("Bajarildi")
    await safe_edit(c.message, f"✅ <b>{n} ta paket narxi yangilandi</b> ({int(tenths) / 10:+g}%).")


_sup_last: dict = {}


@user_r.callback_query(F.data == "sup:new")
async def sup_new(c: CallbackQuery, state: FSMContext):
    await c.answer()
    await state.set_state(Sup.text)
    await c.message.answer("✍️ Savol yoki muammoingizni yozing (buyurtma kodi bo'lsa, ko'rsating):", reply_markup=cancel_kb())


@user_r.message(StateFilter(Sup.text), F.text)
async def sup_save(m: Message, state: FSMContext, bot: Bot):
    text = (m.text or "").strip()
    if not text or len(text) > 800:
        return await m.answer("❗️ Xabar 1–800 belgi bo'lsin.")
    if time.time() - _sup_last.get(m.from_user.id, 0) < 20:
        return await m.answer("⏳ Biroz kuting va qayta yuboring (20 soniya).")
    _sup_last[m.from_user.id] = time.time()
    tid, _ = await db.execute("INSERT INTO tickets(user_id,text) VALUES(?,?)", (m.from_user.id, text))
    await state.clear()
    await m.answer("✅ Xabaringiz adminga yuborildi. Javobni shu yerda olasiz.", reply_markup=await main_menu(m.from_user.id))
    u = await get_user(m.from_user.id)
    kb = InlineKeyboardBuilder()
    kb.button(text="↩️ Javob berish", callback_data=f"sr:{m.from_user.id}")
    for aid in await admin_ids():
        await tell(bot, aid, f"✉️ <b>Murojaat #{tid}</b>\n👤 {mention(u['id'], u['full_name'])} (<code>{u['id']}</code>) "
                             f"• 💰 {fmt(u['balance'])} • 💎 {fmt(u['total_uc'])} UC\n\n{esc(text)}", reply_markup=kb.as_markup())


@user_r.message(StateFilter(Sup.text))
async def sup_wrong(m: Message):
    await m.answer("✍️ Iltimos, savolni matn ko'rinishida yozing.")


@admin_r.callback_query(F.data.startswith("sr:"))
async def sup_reply_ask(c: CallbackQuery, state: FSMContext):
    await c.answer()
    await state.set_state(Adm.support_reply)
    await state.update_data(uid=int(c.data.split(":")[1]))
    await c.message.answer("↩️ Mijozga javobingizni yozing:", reply_markup=cancel_kb())


@admin_r.message(StateFilter(Adm.support_reply), F.text)
async def sup_reply_send(m: Message, state: FSMContext, bot: Bot):
    d = await state.get_data()
    text = (m.text or "").strip()
    if not text:
        return await m.answer("❗️ Javob matnini yozing.")
    kb = InlineKeyboardBuilder()
    kb.button(text="✍️ Yana yozish", callback_data="sup:new")
    await tell(bot, d["uid"], f"💬 <b>Admin javobi:</b>\n\n{esc(text[:1500])}", reply_markup=kb.as_markup())
    await state.clear()
    await audit(m.from_user.id, "murojaatga javob", f"user {d['uid']}")
    await m.answer("✅ Javob yuborildi.", reply_markup=admin_kb(await admin_role(m.from_user.id) == "operator"))


# ================= 🔥 AKSIYA VA NARX POSTI (admin)
async def price_post_body(bot: Bot) -> str:
    fl = await flash_pct_now()
    rows = await db.fetchall("SELECT name,uc,price FROM packages WHERE active=1 AND uc>0 ORDER BY uc")
    lines = []
    for r in rows:
        np = flash_price(r["price"], fl)
        lines.append(f"💎 <b>{esc(r['name'])}</b> — {fmt(np)} so'm" + (f"  <s>{fmt(r['price'])}</s>" if fl else ""))
    head = f"🔥 <b>AKSIYA −{fl}%!</b>\n\n" if fl else ""
    try:
        me = await bot.get_me()
        link = f"\n\n🤖 Buyurtma: @{me.username}"
    except Exception:
        link = ""
    return head + "💎 <b>PUBG UC narxlari</b>\n\n" + "\n".join(lines or ["Hozircha paket yo'q."]) + link


async def refresh_price_post(bot: Bot) -> bool:
    chat = int(await S("price_chat") or 0)
    if not chat:
        return False
    body = await price_post_body(bot)
    text = body + f"\n\n🕒 Yangilandi: {now_tz():%d.%m %H:%M}"
    msg_id = int(await S("price_msg_id") or 0)
    try:
        if msg_id:
            try:
                await bot.edit_message_text(text, chat_id=chat, message_id=msg_id)
                done = True
            except TelegramBadRequest as e:
                done = "not modified" in str(e).lower()
        else:
            done = False
        if not done:
            m = await bot.send_message(chat, text)
            await db.execute("UPDATE settings SET value=? WHERE key='price_msg_id'", (str(m.message_id),))
        await db.execute("INSERT OR REPLACE INTO settings(key,value) VALUES('price_hash',?)",
                         (hashlib.sha1(body.encode()).hexdigest(),))
        return True
    except Exception:
        log.exception("Narx posti yangilanmadi")
        return False


async def price_post_tick(bot: Bot):
    if not int(await S("price_chat") or 0):
        return
    h = hashlib.sha1((await price_post_body(bot)).encode()).hexdigest()
    cur = await db.fetchone("SELECT value FROM settings WHERE key='price_hash'")
    if not cur or cur["value"] != h:
        await refresh_price_post(bot)


async def flash_tick(bot: Bot):
    """Aksiya vaqti tugasa — o'chiradi."""
    if int(await S("flash_pct") or 0) > 0 and not await flash_pct_now():
        await db.execute("UPDATE settings SET value='0' WHERE key='flash_pct'")
        await db.execute("UPDATE settings SET value='' WHERE key='flash_until'")
        for aid in await admin_ids():
            await tell(bot, aid, "⏹ Flash-aksiya vaqti tugadi, narxlar asl holiga qaytdi.")


async def flash_view():
    fl = await flash_pct_now()
    chat = int(await S("price_chat") or 0)
    b = InlineKeyboardBuilder()
    if fl:
        b.button(text="⏹ Aksiyani to'xtatish", callback_data="fl:stop")
        b.button(text="📣 Mijozlarga e'lon qilish", callback_data="fl:ann")
    else:
        b.button(text="▶️ Aksiya boshlash", callback_data="fl:start")
    b.button(text="📌 Narxlar posti kanali" + (" ✅" if chat else ""), callback_data="pp:set")
    if chat:
        b.button(text="🔄 Postni hozir yangilash", callback_data="pp:now")
        b.button(text="🗑 Post kanalini o'chirish", callback_data="pp:off")
    b.adjust(1)
    st = f"🔥 Faol: <b>−{fl}%</b>, {esc((await S('flash_until'))[5:])} gacha" if fl else "Hozir aksiya yo'q."
    return (f"🔥 <b>Aksiya va narx posti</b>\n\n{st}\n📌 Narxlar posti: "
            f"{'<code>' + str(chat) + '</code> (narx o`zgarsa avtomatik yangilanadi)' if chat else 'sozlanmagan'}\n\n"
            "ℹ️ Aksiya hamma paketga qo'llanadi; kupon va tanga bilan birga ishlaydi."), b.as_markup()


@admin_r.message(F.text == T.A_FLASH)
async def flash_menu(m: Message):
    text, kb = await flash_view()
    await m.answer(text, reply_markup=kb)


@admin_r.callback_query(F.data == "fl:start")
async def flash_start(c: CallbackQuery, state: FSMContext):
    await c.answer()
    await state.set_state(Adm.fl_pct)
    await c.message.answer("🔥 Necha <b>foiz</b> chegirma? (1 dan 50 gacha):", reply_markup=cancel_kb())


@admin_r.message(StateFilter(Adm.fl_pct))
async def flash_pct_in(m: Message, state: FSMContext):
    n = parse_int(m.text or "")
    if not n or not 1 <= n <= 50:
        return await m.answer("❗️ 1 dan 50 gacha son kiriting.")
    await state.update_data(pct=n)
    await state.set_state(Adm.fl_hours)
    await m.answer("⏳ Necha <b>soat</b> davom etsin? (1 dan 168 gacha):")


@admin_r.message(StateFilter(Adm.fl_hours))
async def flash_hours_in(m: Message, state: FSMContext):
    h = parse_int(m.text or "")
    if not h or not 1 <= h <= 168:
        return await m.answer("❗️ 1 dan 168 gacha son kiriting.")
    d = await state.get_data()
    until = (now_tz().replace(tzinfo=None) + timedelta(hours=h)).strftime("%Y-%m-%d %H:%M")
    await db.execute("UPDATE settings SET value=? WHERE key='flash_pct'", (str(d["pct"]),))
    await db.execute("UPDATE settings SET value=? WHERE key='flash_until'", (until,))
    await state.clear()
    await audit(m.from_user.id, "flash aksiya", f"−{d['pct']}% {h} soat")
    await m.answer(f"✅ Aksiya boshlandi: −{d['pct']}% ({until} gacha).", reply_markup=admin_kb())
    text, kb = await flash_view()
    await m.answer(text, reply_markup=kb)


@admin_r.callback_query(F.data == "fl:stop")
async def flash_stop(c: CallbackQuery):
    await db.execute("UPDATE settings SET value='0' WHERE key='flash_pct'")
    await db.execute("UPDATE settings SET value='' WHERE key='flash_until'")
    await audit(c.from_user.id, "flash to'xtatildi", "")
    await c.answer("To'xtatildi")
    text, kb = await flash_view()
    await safe_edit(c.message, text, kb)


@admin_r.callback_query(F.data == "fl:ann")
async def flash_announce(c: CallbackQuery, bot: Bot):
    fl = await flash_pct_now()
    if not fl:
        return await c.answer("Aksiya faol emas.", show_alert=True)
    await c.answer("Yuborilmoqda...")
    ids = [r["id"] for r in await db.fetchall("SELECT id FROM users WHERE banned=0 AND blocked=0")]
    kb = InlineKeyboardBuilder()
    kb.button(text="🛒 Hozir sotib olish", callback_data="goto:buy")
    text = f"🔥 <b>AKSIYA −{fl}%!</b>\n⏳ {esc((await S('flash_until'))[5:])} gacha barcha UC paketlarga chegirma!\nUlgurib qoling 👇"
    task = asyncio.create_task(mass_send(bot, c.from_user.id, ids, text, kb.as_markup(), "Aksiya e'loni"))
    BG_TASKS.add(task)
    task.add_done_callback(BG_TASKS.discard)
    await c.message.answer(f"⏳ E'lon {len(ids)} ta foydalanuvchiga yuborilmoqda. Tugagach hisobot keladi.")


@user_r.callback_query(F.data == "goto:buy")
async def goto_buy(c: CallbackQuery):
    await c.answer()
    await buy_list(c.message)


@admin_r.callback_query(F.data == "pp:set")
async def pp_set(c: CallbackQuery, state: FSMContext):
    await c.answer()
    await state.set_state(Adm.pp_chat)
    await c.message.answer("📌 Narxlar posti chiqadigan <b>kanal</b>ni yuboring: <code>@username</code> yoki ID "
                           "(<code>-100...</code>). Bot o'sha kanalda <b>admin</b> bo'lishi kerak.", reply_markup=cancel_kb())


@admin_r.message(StateFilter(Adm.pp_chat))
async def pp_chat_in(m: Message, state: FSMContext, bot: Bot):
    raw = (m.text or "").strip()
    try:
        chat = await bot.get_chat(raw if raw.startswith("@") else int(raw))
    except Exception:
        return await m.answer("❌ Kanal topilmadi. @username yoki ID ni tekshiring va botni kanalga admin qiling.")
    await db.execute("UPDATE settings SET value=? WHERE key='price_chat'", (str(chat.id),))
    await db.execute("UPDATE settings SET value='0' WHERE key='price_msg_id'")
    await state.clear()
    ok = await refresh_price_post(bot)
    await audit(m.from_user.id, "narx posti kanali", str(chat.id))
    await m.answer("✅ Kanal saqlandi va narxlar posti yuborildi." if ok else
                   "⚠️ Kanal saqlandi, lekin post yuborilmadi — botni kanalga admin qiling (post yuborish huquqi bilan).",
                   reply_markup=admin_kb())


@admin_r.callback_query(F.data == "pp:now")
async def pp_now(c: CallbackQuery, bot: Bot):
    ok = await refresh_price_post(bot)
    await c.answer("🔄 Yangilandi" if ok else "Xato: kanal/ruxsatni tekshiring", show_alert=not ok)


@admin_r.callback_query(F.data == "pp:off")
async def pp_off(c: CallbackQuery):
    await db.execute("UPDATE settings SET value='0' WHERE key IN ('price_chat','price_msg_id')")
    await c.answer("O'chirildi")
    text, kb = await flash_view()
    await safe_edit(c.message, text, kb)


# ================= 💌 QAYTARISH KAMPANIYASI
async def winback_ids() -> list:
    days = max(await Si("winback_days"), 1)
    rows = await db.fetchall(
        "SELECT u.id FROM users u WHERE u.banned=0 AND u.blocked=0 AND u.id IN "
        "(SELECT user_id FROM orders WHERE status='done' GROUP BY user_id HAVING MAX(done_at) < datetime('now','+5 hours',?)) "
        "AND (u.winback_at IS NULL OR u.winback_at < datetime('now','+5 hours','-30 days'))", (f"-{days} days",))
    return [r["id"] for r in rows]


@admin_r.message(F.text == T.A_WINBACK)
async def winback_start(m: Message, state: FSMContext):
    n = len(await winback_ids())
    if not n:
        return await m.answer(f"💌 Hozir {await Si('winback_days')} kundan beri xarid qilmagan mijoz yo'q.")
    await state.set_state(Adm.wb_amount)
    await m.answer(f"💌 <b>Qaytarish kampaniyasi</b>\n\n{await Si('winback_days')} kundan beri xarid qilmagan mijozlar: <b>{n}</b> ta.\n"
                   f"Ularga bir martalik kupon yuboriladi. Kupon <b>chegirma summasi</b>ni yozing (so'm):", reply_markup=cancel_kb())


@admin_r.message(StateFilter(Adm.wb_amount))
async def winback_amount(m: Message, state: FSMContext):
    n = parse_int(m.text or "")
    if not n or not 1000 <= n <= 1_000_000:
        return await m.answer("❗️ 1 000 dan 1 000 000 gacha summa kiriting.")
    ids = await winback_ids()
    await state.clear()
    kb = InlineKeyboardBuilder()
    kb.button(text=f"✅ Ha, {len(ids)} ta mijozga yuborish", callback_data=f"wbok:{n}")
    kb.button(text="❌ Bekor qilish", callback_data="bpn")
    kb.adjust(1)
    await m.answer(f"💌 {len(ids)} ta mijozga <b>{fmt(n)} so'm</b> chegirmali kupon yuboriladi. Tasdiqlaysizmi?", reply_markup=kb.as_markup())
    await m.answer("👆 Tugmani tanlang.", reply_markup=admin_kb())


@admin_r.callback_query(F.data.startswith("wbok:"))
async def winback_send(c: CallbackQuery, bot: Bot):
    amount = int(c.data.split(":")[1])
    ids = await winback_ids()
    if not ids:
        return await c.answer("Mijoz topilmadi.", show_alert=True)
    code = None
    for _ in range(10):
        cand = f"BACK{random.randint(1000, 9999)}"
        try:
            await db.execute("INSERT INTO coupons(code,discount,min_uc,limit_count) VALUES(?,?,0,?)", (cand, amount, len(ids) + 5))
            code = cand
            break
        except aiosqlite.IntegrityError:
            continue
    if not code:
        return await c.answer("Kupon kodi yaratilmadi, qayta urinib ko'ring.", show_alert=True)
    for i in ids:
        await db.execute("UPDATE users SET winback_at=datetime('now','+5 hours') WHERE id=?", (i,))
    await audit(c.from_user.id, "qaytarish kampaniyasi", f"{code} • {fmt(amount)} • {len(ids)} ta")
    await c.answer("Yuborilmoqda...")
    kb = InlineKeyboardBuilder()
    kb.button(text="🛒 Sotib olish", callback_data="goto:buy")
    text = (f"💌 <b>Sizni sog'indik!</b>\n\nSizga shaxsiy kupon: <code>{code}</code> — <b>{fmt(amount)} so'm</b> chegirma.\n"
            f"Buyurtma berishda «🎟 Kupon kiritish» ni bosib shu kodni yozing. Kupon bir martalik.")
    task = asyncio.create_task(mass_send(bot, c.from_user.id, ids, text, kb.as_markup(), "Qaytarish kuponi"))
    BG_TASKS.add(task)
    task.add_done_callback(BG_TASKS.discard)
    await safe_edit(c.message, f"⏳ Kupon <code>{code}</code> {len(ids)} ta mijozga yuborilmoqda. Tugagach hisobot keladi.")


# ================= ⏱ TEZLIK HISOBOTI, 📗 EXCEL, 📆 KUNLIK HISOBOT
async def speed_report(days: int = 30) -> str:
    rows = await db.fetchall(
        "SELECT o.code, o.kind, o.admin_id, (julianday(o.done_at)-julianday(o.check_at))*1440 AS mins, u.full_name "
        "FROM orders o LEFT JOIN users u ON u.id=o.admin_id "
        "WHERE o.status='done' AND o.check_at IS NOT NULL AND o.done_at > datetime('now','+5 hours',?)", (f"-{days} days",))
    vals = sorted(r["mins"] for r in rows if r["mins"] is not None and r["mins"] >= 0)
    if not vals:
        return f"⏱ <b>Tezlik ({days} kun)</b>\n\nHali ma'lumot yo'q (chek→UC vaqti yangi buyurtmalardan boshlab hisoblanadi)."
    n = len(vals)
    med = vals[n // 2] if n % 2 else (vals[n // 2 - 1] + vals[n // 2]) / 2
    p90 = vals[min(int(n * 0.9), n - 1)]
    fast = sum(1 for v in vals if v <= 5) * 100 // n
    per: dict = {}
    for r in rows:
        if r["mins"] is not None and r["mins"] >= 0:
            per.setdefault(r["full_name"] or str(r["admin_id"]), []).append(r["mins"])
    adm = "\n".join(f"• {esc(k[:16])}: {sum(v) / len(v):.1f} daq ({len(v)} ta)" for k, v in sorted(per.items(), key=lambda kv: sum(kv[1]) / len(kv[1])))
    slow = sorted((r for r in rows if r["mins"] is not None), key=lambda r: -r["mins"])[:3]
    return (f"⏱ <b>Tezlik hisoboti ({days} kun)</b>\nChek kelgandan UC tushguncha:\n\n"
            f"📦 Buyurtmalar: <b>{n}</b>\n⚡️ O'rtacha: <b>{sum(vals) / n:.1f} daq</b> | median: {med:.1f} | 90%: {p90:.1f}\n"
            f"✅ 5 daqiqada bajarilgan: <b>{fast}%</b>\n\n👥 Adminlar bo'yicha:\n{adm}\n\n"
            f"🐢 Eng sekin: " + ", ".join(f"{esc(r['code'])} ({r['mins']:.0f} daq)" for r in slow))


async def daily_report(bot: Bot):
    if (await S("daily_report")) != "1":
        return
    n = now_tz()
    today = n.strftime("%Y-%m-%d")
    if (n.hour, n.minute) < (23, 55) or (await S("last_report")) == today:
        return
    await db.execute("UPDATE settings SET value=? WHERE key='last_report'", (today,))
    r = await db.fetchone("SELECT COUNT(*) n, COALESCE(SUM(uc),0) uc, COALESCE(SUM(final),0) rev, COALESCE(SUM(cost),0) cost, "
                          "COALESCE(SUM(kind='midas'),0) mid FROM orders WHERE status='done' AND substr(done_at,1,10)=?", (today,))
    tp = await db.fetchone("SELECT COUNT(*) n, COALESCE(SUM(amount),0) s FROM topups WHERE status='done' AND substr(done_at,1,10)=?", (today,))
    nu = await db.fetchone("SELECT COUNT(*) n FROM users WHERE substr(joined_at,1,10)=?", (today,))
    pend = await db.fetchone("SELECT (SELECT COUNT(*) FROM orders WHERE status IN ('checking','processing')) o, "
                             "(SELECT COUNT(*) FROM topups WHERE status='checking') t")
    text = (f"📆 <b>Kunlik hisobot — {n:%d.%m.%Y}</b>\n\n🧾 Bajarilgan buyurtma: <b>{r['n']}</b> ta (🌐 {r['mid']})\n"
            f"💎 UC: <b>{fmt(r['uc'])}</b>\n💵 Tushum: <b>{fmt(r['rev'])} so'm</b>\n"
            f"💰 Foyda: <b>{fmt(r['rev'] - r['cost'])} so'm</b>\n"
            f"➕ Balans to'ldirish: {tp['n']} ta • {fmt(tp['s'])} so'm\n👥 Yangi foydalanuvchi: {nu['n']}\n"
            f"⏳ Kutilayotgan: {pend['o']} buyurtma, {pend['t']} to'ldirish")
    for a in await db.fetchall("SELECT id FROM admins WHERE role!='operator'"):
        await tell(bot, a["id"], text)


def _xl_safe(v):
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


def build_xlsx(summary, orders, topups) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Xulosa"
    ws.append(["Oy", "Buyurtma", "UC", "Tushum", "Tannarx", "Foyda"])
    for r in summary:
        ws.append([r["m"], r["n"], r["uc"], r["rev"], r["cost"], r["rev"] - r["cost"]])
    ws2 = wb.create_sheet("Buyurtmalar")
    ws2.append(["Kod", "Foydalanuvchi", "PUBG ID", "Paket", "UC", "Narx", "To'langan", "Tannarx", "Usul", "Holat", "Sana"])
    for r in orders:
        ws2.append([_xl_safe(x) for x in (r["code"], r["user_id"], r["pubg_id"], r["pkg_name"], r["uc"], r["price"], r["final"],
                                           r["cost"], r["pay_method"], r["status"], r["created_at"])])
    ws3 = wb.create_sheet("To'ldirishlar")
    ws3.append(["Kod", "Foydalanuvchi", "Summa", "Holat", "Sana"])
    for r in topups:
        ws3.append([_xl_safe(x) for x in (r["code"], r["user_id"], r["amount"], r["status"], r["created_at"])])
    for w in (ws, ws2, ws3):
        for cell in w[1]:
            cell.font = Font(bold=True)
        for col in w.columns:
            w.column_dimensions[col[0].column_letter].width = 16
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


@admin_r.callback_query(F.data == "st:xlsx")
async def stats_xlsx(c: CallbackQuery):
    await c.answer("Tayyorlanmoqda...")
    summary = await db.fetchall("SELECT substr(done_at,1,7) m, COUNT(*) n, COALESCE(SUM(uc),0) uc, COALESCE(SUM(final),0) rev, "
                                "COALESCE(SUM(cost),0) cost FROM orders WHERE status='done' GROUP BY m ORDER BY m DESC LIMIT 24")
    orders = await db.fetchall("SELECT * FROM orders ORDER BY id DESC LIMIT 5000")
    topups = await db.fetchall("SELECT * FROM topups ORDER BY id DESC LIMIT 5000")
    try:
        data = await run_blocking(build_xlsx, summary, orders, topups)
    except ImportError:
        return await c.message.answer("❌ Excel uchun serverda <code>pip install openpyxl</code> kerak. CSV tugmalari ishlaydi.")
    await c.message.answer_document(BufferedInputFile(data, filename=f"hisobot_{now_tz():%Y%m%d}.xlsx"),
                                    caption="📗 Excel hisobot: Xulosa (oylar) • Buyurtmalar • To'ldirishlar")


@admin_r.callback_query(F.data == "st:speed")
async def stats_speed(c: CallbackQuery):
    await c.answer()
    await c.message.answer(await speed_report())


# ============================================================ ✏️ UMUMIY TAHRIRLASH (paket, kupon, kanal, turnir)
# maydon: (nomi, turi, maks.uzunlik)  turlar: str | opt (bo'sh bo'lishi mumkin, «-») | pos | nonneg | code | url
EDIT = {
    "pkg": ("packages", "📦 Paket", {"name": ("Nomi", "str", 40), "uc": ("UC miqdori", "pos", 0),
                                      "price": ("Sotuv narxi (so'm)", "pos", 0), "cost": ("Tannarx (so'm)", "nonneg", 0)}),
    "cp": ("coupons", "🎟 Kupon", {"code": ("Kod", "code", 20), "discount": ("Chegirma (so'm)", "pos", 0),
                                    "min_uc": ("Minimal UC", "nonneg", 0), "limit_count": ("Limit (necha kishi)", "pos", 0)}),
    "ch": ("channels", "📢 Kanal", {"title": ("Tugma nomi", "str", 40), "link": ("Havola", "url", 300)}),
    "tn": ("tournaments", "🏆 Turnir", {"title": ("Nomi", "str", 60), "fee": ("Kirish narxi (so'm)", "nonneg", 0),
                                         "prize": ("Mukofot", "str", 200), "start_text": ("Vaqti", "str", 80),
                                         "max_players": ("Nechta kishi", "pos", 0), "info": ("Qo'shimcha ma'lumot", "opt", 800),
                                         "room_info": ("Xona kodi/nomi", "opt", 500)}),
}


def _ev(v, kind):
    if v is None or v == "":
        return "—"
    return fmt(v) if kind in ("pos", "nonneg") else esc(str(v)[:120])


async def edit_menu_view(ent: str, rid: int):
    table, title, fields = EDIT[ent]
    row = await db.fetchone(f"SELECT * FROM {table} WHERE id=?", (rid,))
    if not row:
        return None
    lines = [f"• {lbl}: <b>{_ev(row[f], kind)}</b>" for f, (lbl, kind, _) in fields.items()]
    b = InlineKeyboardBuilder()
    for f, (lbl, _, _) in fields.items():
        b.button(text=f"✏️ {lbl}", callback_data=f"edf:{ent}:{rid}:{f}")
    b.button(text="🔙 Orqaga", callback_data=f"edb:{ent}:{rid}")
    b.adjust(*([2] * ((len(fields) + 1) // 2)), 1)
    return f"{title} <b>#{rid}</b> — tahrirlash\n\n" + "\n".join(lines) + "\n\nQaysi maydonni o'zgartirasiz?", b.as_markup()


@admin_r.callback_query(F.data.startswith("edm:"))
async def edit_menu(c: CallbackQuery):
    _, ent, rid = c.data.split(":")
    if ent not in EDIT:
        return await c.answer("Noma'lum")
    v = await edit_menu_view(ent, int(rid))
    if not v:
        return await c.answer("Topilmadi (o'chirilgan bo'lishi mumkin).", show_alert=True)
    await c.answer()
    await safe_edit(c.message, *v)


@admin_r.callback_query(F.data.startswith("edb:"))
async def edit_back(c: CallbackQuery):
    _, ent, rid = c.data.split(":")
    rid = int(rid)
    await c.answer()
    if ent == "pkg":
        p = await db.fetchone("SELECT kind FROM packages WHERE id=?", (rid,))
        v = await (midas_view() if p and p["kind"] == "midas" else packages_view())
    elif ent == "cp":
        v = await coupons_view()
    elif ent == "ch":
        v = await channels_view()
    else:
        v = await atn_view(rid)
    await safe_edit(c.message, *v)


@admin_r.callback_query(F.data.startswith("edf:"))
async def edit_field(c: CallbackQuery, state: FSMContext):
    _, ent, rid, f = c.data.split(":")
    if ent not in EDIT or f not in EDIT[ent][2]:
        return await c.answer("Noma'lum")
    table, _, fields = EDIT[ent]
    row = await db.fetchone(f"SELECT * FROM {table} WHERE id=?", (int(rid),))
    if not row:
        return await c.answer("Topilmadi.", show_alert=True)
    lbl, kind, _ = fields[f]
    await c.answer()
    await state.set_state(Adm.edit_val)
    await state.update_data(ent=ent, rid=int(rid), f=f)
    hint = {"opt": "\nTozalash uchun <code>-</code> yozing.", "url": "\n(http:// yoki https:// bilan)",
            "pos": "\n(musbat son)", "nonneg": "\n(0 yoki musbat son)"}.get(kind, "")
    await c.message.answer(f"✏️ <b>{lbl}</b>\nHozirgi: <b>{_ev(row[f], kind)}</b>\n\nYangi qiymatni yozing:{hint}",
                           reply_markup=cancel_kb())


@admin_r.message(StateFilter(Adm.edit_val))
async def edit_save(m: Message, state: FSMContext):
    d = await state.get_data()
    ent, rid, f = d["ent"], d["rid"], d["f"]
    table, _, fields = EDIT[ent]
    lbl, kind, maxlen = fields[f]
    raw = (m.text or "").strip()
    if kind in ("pos", "nonneg"):
        val = parse_int(raw)
        if val is None or val < (1 if kind == "pos" else 0):
            return await m.answer("❗️ " + ("Musbat" if kind == "pos" else "0 yoki musbat") + " son kiriting.")
    elif kind == "code":
        val = raw.upper()
        if not re.fullmatch(r"[A-Z0-9_]{3,20}", val):
            return await m.answer("❗️ Kod 3–20 ta lotin harf/raqamdan iborat bo'lsin.")
    elif kind == "url":
        val = raw
        if not re.match(r"https?://\S+$", val):
            return await m.answer("❗️ Havola http:// yoki https:// bilan boshlansin.")
    elif kind == "opt":
        val = "" if raw == "-" else raw[:maxlen]
    else:
        if not raw:
            return await m.answer("❗️ Bo'sh bo'lmasin.")
        val = raw[:maxlen]
    if ent == "tn" and f == "start_text":
        dt = parse_when(val)
        await db.execute("UPDATE tournaments SET start_at=?, reminded=0 WHERE id=?",
                         (dt.strftime("%Y-%m-%d %H:%M") if dt else None, rid))
        if dt:
            val = f"{dt:%d.%m %H:%M}"
    if ent == "tn" and f == "max_players" and val < await tn_count(rid):
        return await m.answer(f"❗️ Hozir {await tn_count(rid)} kishi ro'yxatda — undan kam bo'lishi mumkin emas.")
    old = await db.fetchone(f"SELECT {f} v FROM {table} WHERE id=?", (rid,))
    await audit(m.from_user.id, f"tahrir {ent}#{rid}", f"{f}: {str(old['v'] if old else '')[:40]} → {str(val)[:40]}")
    try:
        await db.execute(f"UPDATE {table} SET {f}=? WHERE id=?", (val, rid))
    except aiosqlite.IntegrityError:
        return await m.answer("❗️ Bu kod allaqachon mavjud. Boshqasini yozing.")
    if ent == "ch":
        _sub_cache.clear()
    await state.clear()
    await m.answer("✅ Yangilandi.", reply_markup=admin_kb())
    v = await edit_menu_view(ent, rid)
    if v:
        await m.answer(v[0], reply_markup=v[1])


# ============================================================ 🕒 ISH REJIMI
def parse_hours(s: str):
    m = re.match(r"\s*(\d{1,2})[:.](\d{2})\s*[-–—]\s*(\d{1,2})[:.](\d{2})\s*$", s or "")
    if not m:
        return None
    a, b = int(m.group(1)) * 60 + int(m.group(2)), int(m.group(3)) * 60 + int(m.group(4))
    return (a, b) if a < 1440 and b < 1440 else None


async def shop_is_open() -> bool:
    if (await S("auto_hours")) == "1":
        rng = parse_hours(await S("work_hours"))
        if rng:
            n = now_tz()
            cur, (a, b) = n.hour * 60 + n.minute, rng
            if a == b:
                return True
            return (a <= cur < b) if a < b else (cur >= a or cur < b)   # tundan o'tadigan vaqt ham ishlaydi
    return (await S("shop_open")) == "1"


async def shop_status_text() -> str:
    hours = esc(await S("work_hours"))
    if await shop_is_open():
        return f"🟢 <b>Do'kon hozir OCHIQ</b>\n🕒 Ish vaqti: {hours}"
    return (f"🔴 <b>Do'kon hozir YOPIQ</b>\n🕒 Ish vaqti: {hours}\n"
            f"ℹ️ Buyurtma berishingiz mumkin — do'kon ochilganda UC tashlab beramiz.")


async def closed_note() -> str:
    if await shop_is_open():
        return ""
    return ("\n\n🔴 <b>Do'kon hozirda yopiq.</b> Buyurtma berishda davom etsangiz, do'kon ochilganda UC tashlab beramiz.\n"
            f"🕒 Ish vaqti: {esc(await S('work_hours'))}")


async def hours_view():
    auto = (await S("auto_hours")) == "1"
    b = InlineKeyboardBuilder()
    if not auto:
        manual_open = (await S("shop_open")) == "1"
        b.button(text="🔴 Do'konni YOPISH" if manual_open else "🟢 Do'konni OCHISH", callback_data="shop:toggle")
    b.button(text=f"🤖 Avto rejim: {'YOQILGAN ✅' if auto else 'o`chiq'}", callback_data="shop:auto")
    b.button(text="✏️ Ish vaqtini o'zgartirish", callback_data="set:work_hours")
    b.adjust(1)
    note = ("🤖 Avto rejim: do'kon ish vaqtiga qarab o'zi ochiladi va yopiladi." if auto
            else "✋ Qo'lda rejim: do'konni o'zingiz ochasiz/yopasiz.")
    return ("🕒 <b>Ish rejimi</b>\n\n" + await shop_status_text() + f"\n\n{note}\n"
            "Ish vaqti formati: <code>09:00 - 23:00</code>"), b.as_markup()


@admin_r.message(F.text == T.A_HOURS)
async def hours_menu(m: Message):
    text, kb = await hours_view()
    await m.answer(text, reply_markup=kb)


@admin_r.callback_query(F.data == "shop:toggle")
async def shop_toggle(c: CallbackQuery):
    new = "0" if await shop_is_open() else "1"
    await db.execute("UPDATE settings SET value=? WHERE key='shop_open'", (new,))
    await audit(c.from_user.id, "do'kon", "ochildi" if new == "1" else "yopildi")
    await c.answer("🟢 Do'kon ochildi" if new == "1" else "🔴 Do'kon yopildi")
    text, kb = await hours_view()
    await safe_edit(c.message, text, kb)


@admin_r.callback_query(F.data == "shop:auto")
async def shop_auto(c: CallbackQuery):
    new = "0" if (await S("auto_hours")) == "1" else "1"
    if new == "1" and not parse_hours(await S("work_hours")):
        return await c.answer("Avval ish vaqtini 09:00 - 23:00 formatida kiriting.", show_alert=True)
    await db.execute("UPDATE settings SET value=? WHERE key='auto_hours'", (new,))
    await audit(c.from_user.id, "avto rejim", "yoqildi" if new == "1" else "o'chirildi")
    await c.answer("🤖 Avto rejim " + ("yoqildi" if new == "1" else "o'chirildi"))
    text, kb = await hours_view()
    await safe_edit(c.message, text, kb)


# ============================================================ ⭐ SHARXLAR (isbot tasmasi)
def review_kb(oid: int) -> InlineKeyboardMarkup:
    kb = InlineKeyboardBuilder()
    kb.button(text="✍️ Sharx yozish (ixtiyoriy)", callback_data=f"rv:{oid}")
    return kb.as_markup()


async def start_review(c: CallbackQuery, state: FSMContext, oid: int):
    o = await db.fetchone("SELECT * FROM orders WHERE id=? AND user_id=? AND status='done'", (oid, c.from_user.id))
    if not o:
        return await c.answer("Buyurtma topilmadi.", show_alert=True)
    if await db.fetchone("SELECT 1 FROM reviews WHERE order_id=?", (oid,)):
        return await c.answer("Bu buyurtma uchun sharx yuborgansiz. Rahmat! 🙏", show_alert=True)
    await c.answer()
    await state.set_state(Rev.text)
    await state.update_data(oid=oid)
    await c.message.answer("✍️ <b>Sharxingizni yozing</b> (xohlasangiz rasm/skrinshot ham qo'shing — bu isbot bo'ladi).\n"
                           "Majburiy emas — bekor qilish uchun pastdagi tugmani bosing.", reply_markup=cancel_kb())


@user_r.callback_query(F.data.startswith("rv:"))
async def review_ask(c: CallbackQuery, state: FSMContext):
    await start_review(c, state, int(c.data.split(":")[1]))


@user_r.callback_query(F.data == "rvn")
async def review_new(c: CallbackQuery, state: FSMContext):
    o = await db.fetchone(
        "SELECT id FROM orders WHERE user_id=? AND status='done' AND id NOT IN "
        "(SELECT order_id FROM reviews WHERE order_id IS NOT NULL) ORDER BY id DESC LIMIT 1", (c.from_user.id,))
    if not o:
        return await c.answer("Sharx qoldirish uchun avval xarid qiling (yoki barcha buyurtmalaringizga sharx yuborgansiz).",
                              show_alert=True)
    await start_review(c, state, o["id"])


@user_r.message(StateFilter(Rev.text), F.text | F.photo)
async def review_save(m: Message, state: FSMContext, bot: Bot):
    d = await state.get_data()
    oid = d.get("oid")
    text = (m.text or m.caption or "").strip()
    photo = m.photo[-1].file_id if m.photo else None
    if not text and not photo:
        return await m.answer("✍️ Matn yoki rasm yuboring.")
    if len(text) > 500:
        return await m.answer("❗️ Sharx 500 belgidan oshmasin. Qisqaroq yozing.")
    o = await db.fetchone("SELECT * FROM orders WHERE id=? AND user_id=? AND status='done'", (oid, m.from_user.id))
    if not o or await db.fetchone("SELECT 1 FROM reviews WHERE order_id=?", (oid,)):
        await state.clear()
        return await m.answer("ℹ️ Bu buyurtma uchun sharx allaqachon yuborilgan.", reply_markup=await main_menu(m.from_user.id))
    rid, _ = await db.execute(
        "INSERT INTO reviews(user_id,order_id,order_code,uc,bought_at,text,photo_id) VALUES(?,?,?,?,?,?,?)",
        (m.from_user.id, oid, o["code"], o["uc"], o["done_at"], text, photo))
    await state.clear()
    await m.answer("🙏 <b>Rahmat!</b> Sharxingiz adminga yuborildi. Tasdiqlansa, «⭐ Sharxlar» tasmasida chiqadi.",
                   reply_markup=await main_menu(m.from_user.id))
    await notify_review(bot, rid)


@user_r.message(StateFilter(Rev.text))
async def review_wrong(m: Message):
    await m.answer("✍️ Iltimos, sharxni <b>matn</b> yoki <b>rasm</b> ko'rinishida yuboring.")


async def notify_review(bot: Bot, rid: int, only: Optional[int] = None):
    r = await db.fetchone("SELECT * FROM reviews WHERE id=?", (rid,))
    u = await get_user(r["user_id"])
    text = (f"⭐ <b>Yangi sharx #{rid}</b>\n👤 {mention(u['id'], u['full_name'])} (<code>{u['id']}</code>)\n"
            f"🧾 Buyurtma {esc(r['order_code'])} • {r['uc']} UC • {esc((r['bought_at'] or '')[:16])}\n\n"
            f"💬 {esc((r['text'] or '—')[:700])}")
    kb = InlineKeyboardBuilder()
    kb.button(text="📢 Tasmaga qo'shish", callback_data=f"rvp:{rid}")
    kb.button(text="🗑 Rad etish", callback_data=f"rvr:{rid}")
    kb.adjust(1)
    fid = ("p:" + r["photo_id"]) if r["photo_id"] else None
    for aid in ([only] if only else await admin_ids()):
        try:
            await send_media(bot, aid, fid, text, kb.as_markup())
        except Exception:
            log.warning("Sharx adminga yuborilmadi: %s", aid)


@admin_r.callback_query(F.data.startswith(("rvp:", "rvr:")))
async def review_moderate(c: CallbackQuery, bot: Bot):
    act, rid = c.data.split(":")
    rid = int(rid)
    new = "published" if act == "rvp" else "rejected"
    _, rc = await db.execute("UPDATE reviews SET status=? WHERE id=? AND status='pending'", (new, rid))
    if not rc:
        await c.answer("Bu sharx allaqachon ko'rib chiqilgan.", show_alert=True)
        return await mark_msg(c, "ℹ️ Allaqachon ko'rib chiqilgan.")
    await c.answer("📢 Tasmaga qo'shildi" if new == "published" else "Rad etildi")
    await mark_msg(c, f"📢 Tasmaga qo'shdi: {esc(c.from_user.full_name)}" if new == "published"
                   else f"🗑 Rad etdi: {esc(c.from_user.full_name)}")
    if new == "published":
        r = await db.fetchone("SELECT user_id FROM reviews WHERE id=?", (rid,))
        await tell(bot, r["user_id"], "📢 Sharxingiz «⭐ Sharxlar» tasmasiga qo'shildi. Rahmat! 🙏")


@admin_r.message(F.text == T.A_REVIEWS)
async def reviews_admin(m: Message, bot: Bot):
    st = await db.fetchone("SELECT COALESCE(SUM(status='published'),0) p, COALESCE(SUM(status='pending'),0) w FROM reviews")
    await m.answer(f"⭐ <b>Sharx boshqaruvi</b>\n\n📢 Tasmada: <b>{st['p']}</b> ta\n⏳ Kutilmoqda: <b>{st['w']}</b> ta\n\n"
                   f"ℹ️ Tasmadagi sharxni o'chirish: «⭐ Sharxlar» ni oching — har kartada 🗑 tugmasi bor.")
    for r in await db.fetchall("SELECT id FROM reviews WHERE status='pending' ORDER BY id LIMIT 15"):
        await notify_review(bot, r["id"], only=m.from_user.id)


async def review_page(idx: int, admin: bool):
    rows = await db.fetchall("SELECT r.*, u.full_name FROM reviews r LEFT JOIN users u ON u.id=r.user_id "
                             "WHERE r.status='published' ORDER BY r.id DESC LIMIT 200")
    if not rows:
        return None
    n = len(rows)
    idx %= n
    r = rows[idx]
    text = (f"⭐ <b>Mijozlar sharxlari</b> ({idx + 1}/{n})\n\n"
            + (f"«{esc((r['text'] or '')[:600])}»\n\n" if r["text"] else "")
            + f"👤 {esc((r['full_name'] or 'Mijoz')[:12])} • ✅ {r['uc']} UC xarid • {(r['bought_at'] or r['created_at'] or '')[:10]}")
    kb = InlineKeyboardBuilder()
    sizes = []
    if n > 1:
        kb.button(text="◀️", callback_data=f"rt:{(idx - 1) % n}")
        kb.button(text="▶️", callback_data=f"rt:{(idx + 1) % n}")
        sizes.append(2)
    kb.button(text="✍️ Sharx qoldirish", callback_data="rvn")
    sizes.append(1)
    if admin:
        kb.button(text="🗑 Tasmadan o'chirish", callback_data=f"rvd:{r['id']}")
        sizes.append(1)
    kb.adjust(*sizes)
    return r["photo_id"], text, kb.as_markup()


async def send_review_page(msg: Message, idx: int, admin: bool):
    page = await review_page(idx, admin)
    if not page:
        return await msg.answer("⭐ Hozircha sharxlar yo'q. Xariddan keyin birinchi bo'lib sharx qoldiring!")
    photo, text, kb = page
    if photo:
        await msg.answer_photo(photo, caption=text, reply_markup=kb)
    else:
        await msg.answer(text, reply_markup=kb)


@user_r.message(F.text == T.REVIEWS)
async def reviews_menu(m: Message):
    await send_review_page(m, 0, await is_admin(m.from_user.id))


@user_r.callback_query(F.data.startswith("rt:"))
async def review_nav(c: CallbackQuery):
    await c.answer()
    try:
        await c.message.delete()
    except Exception:
        pass
    await send_review_page(c.message, int(c.data.split(":")[1]), await is_admin(c.from_user.id))


@admin_r.callback_query(F.data.startswith("rvd:"))
async def review_delete(c: CallbackQuery):
    await db.execute("UPDATE reviews SET status='rejected' WHERE id=?", (int(c.data.split(":")[1]),))
    await c.answer("🗑 Tasmadan o'chirildi")
    try:
        await c.message.delete()
    except Exception:
        pass
    await send_review_page(c.message, 0, True)


# ============================================================ 🏆 TURNIRLAR
TN_REACTS = ["🔥", "👍", "❤️", "😮"]
TN_STATUS = {"open": "🟢 Ro'yxat ochiq", "started": "🔒 Ro'yxat yopilgan", "finished": "🏁 Yakunlangan",
             "cancelled": "❌ Bekor qilingan"}
_com_last: dict = {}


def fee_text(f: int) -> str:
    return "Bepul" if f == 0 else f"{fmt(f)} so'm"


async def tn_count(tid: int) -> int:
    return (await db.fetchone("SELECT COUNT(*) n FROM t_players WHERE tid=?", (tid,)))["n"]


async def tn_text(t, uid: int = 0) -> str:
    n = await tn_count(t["id"])
    text = (f"🏆 <b>{esc(t['title'])}</b>\n\n💰 Kirish narxi: <b>{fee_text(t['fee'])}</b>\n🎁 Mukofot: <b>{esc(t['prize'])}</b>\n"
            f"🕒 Vaqti: <b>{esc(t['start_text'])}</b>\n👥 O'yinchilar: <b>{n}/{t['max_players']}</b>\n"
            f"📌 Holat: {TN_STATUS.get(t['status'], t['status'])}")
    if t["info"]:
        text += f"\n\n📝 {esc(t['info'])}"
    if uid and await db.fetchone("SELECT 1 FROM t_players WHERE tid=? AND user_id=?", (t["id"], uid)):
        text += "\n\n✅ <b>Siz ro'yxatdasiz.</b>"
        text += (f"\n🔑 <b>Xona ma'lumoti:</b>\n{esc(t['room_info'])}" if t["room_info"]
                 else "\n⏳ Xona kodi va nomi boshlanishdan oldin shu yerga va chatga yuboriladi.")
    return text


async def tn_kb(t, uid: int):
    tid = t["id"]
    counts = {r["emoji"]: r["n"] for r in await db.fetchall(
        "SELECT emoji, COUNT(*) n FROM t_reacts WHERE tid=? GROUP BY emoji", (tid,))}
    mine = await db.fetchone("SELECT emoji FROM t_reacts WHERE tid=? AND user_id=?", (tid, uid))
    kb = InlineKeyboardBuilder()
    for i, e in enumerate(TN_REACTS):
        kb.button(text=f"{'✅' if mine and mine['emoji'] == e else ''}{e} {counts.get(e, 0)}", callback_data=f"tnx:{tid}:{i}")
    sizes = [len(TN_REACTS)]
    joined = await db.fetchone("SELECT 1 FROM t_players WHERE tid=? AND user_id=?", (tid, uid))
    if not joined and t["status"] == "open":
        if await tn_count(tid) >= t["max_players"]:
            kb.button(text="😔 Joylar tugagan", callback_data=f"tnv:{tid}")
        else:
            kb.button(text=f"✅ Qatnashish — {fee_text(t['fee'])}", callback_data=f"tnj:{tid}")
        sizes.append(1)
    cc = (await db.fetchone("SELECT COUNT(*) n FROM t_comments WHERE tid=?", (tid,)))["n"]
    kb.button(text=f"💬 Izohlar ({cc})", callback_data=f"tnc:{tid}")
    kb.button(text="🔙 Turnirlar", callback_data="tnl")
    sizes += [1, 1]
    kb.adjust(*sizes)
    return kb.as_markup()


async def tn_list_view():
    rows = await db.fetchall("SELECT * FROM tournaments WHERE status IN ('open','started') ORDER BY id DESC LIMIT 10")
    if not rows:
        return "🏆 Hozircha faol turnir yo'q. Yangi turnir e'lon qilinganda xabar beramiz!", None
    b = InlineKeyboardBuilder()
    for t in rows:
        b.button(text=f"🏆 {t['title'][:28]} • {fee_text(t['fee'])} • {await tn_count(t['id'])}/{t['max_players']}",
                 callback_data=f"tnv:{t['id']}")
    b.adjust(1)
    return "🏆 <b>Turnirlar</b>\n\nQatnashish uchun turnirni tanlang:", b.as_markup()


@user_r.message(F.text == T.TOURN)
async def tn_menu(m: Message):
    text, kb = await tn_list_view()
    await m.answer(text, reply_markup=kb)


@user_r.callback_query(F.data == "tnl")
async def tn_list_cb(c: CallbackQuery):
    await c.answer()
    text, kb = await tn_list_view()
    await safe_edit(c.message, text, kb)


async def tn_show(c: CallbackQuery, tid: int):
    t = await db.fetchone("SELECT * FROM tournaments WHERE id=?", (tid,))
    if not t:
        return await c.answer("Turnir topilmadi", show_alert=True)
    await safe_edit(c.message, await tn_text(t, c.from_user.id), await tn_kb(t, c.from_user.id))


@user_r.callback_query(F.data.startswith("tnv:"))
async def tn_view_cb(c: CallbackQuery):
    await c.answer()
    await tn_show(c, int(c.data.split(":")[1]))


@user_r.callback_query(F.data.startswith("tnx:"))
async def tn_react(c: CallbackQuery):
    _, tid, i = c.data.split(":")
    tid, i = int(tid), int(i)
    if not await db.fetchone("SELECT 1 FROM tournaments WHERE id=?", (tid,)) or not 0 <= i < len(TN_REACTS):
        return await c.answer("Topilmadi")
    e = TN_REACTS[i]
    cur = await db.fetchone("SELECT emoji FROM t_reacts WHERE tid=? AND user_id=?", (tid, c.from_user.id))
    if cur and cur["emoji"] == e:
        await db.execute("DELETE FROM t_reacts WHERE tid=? AND user_id=?", (tid, c.from_user.id))
    else:
        await db.execute("INSERT INTO t_reacts(tid,user_id,emoji) VALUES(?,?,?) "
                         "ON CONFLICT(tid,user_id) DO UPDATE SET emoji=excluded.emoji", (tid, c.from_user.id, e))
    await c.answer(e)
    await tn_show(c, tid)


async def tn_join(tid: int, uid: int):
    """(turnir, xabar) — muvaffaqiyatli bo'lsa xabar None."""
    try:
        async with db.tx() as c:
            t = await _one(c, "SELECT * FROM tournaments WHERE id=?", (tid,))
            if not t or t["status"] != "open":
                raise Fail("Bu turnirga ro'yxat yopilgan.")
            if await _one(c, "SELECT 1 FROM t_players WHERE tid=? AND user_id=?", (tid, uid)):
                raise Fail("Siz allaqachon ro'yxatdasiz.")
            if (await _one(c, "SELECT COUNT(*) n FROM t_players WHERE tid=?", (tid,)))["n"] >= t["max_players"]:
                raise Fail("Afsus, joylar tugagan.")
            cur = await c.execute("UPDATE users SET balance=balance-? WHERE id=? AND balance>=?", (t["fee"], uid, t["fee"]))
            if cur.rowcount == 0:
                raise Fail(f"Balansingiz yetarli emas ({fee_text(t['fee'])}). «{T.TOPUP}» orqali to'ldiring.")
            await c.execute("INSERT INTO t_players(tid,user_id,paid) VALUES(?,?,?)", (tid, uid, t["fee"]))
            return t, None
    except Fail as e:
        return None, str(e)


@user_r.callback_query(F.data.startswith("tnj:"))
async def tn_join_ask(c: CallbackQuery, bot: Bot):
    tid = int(c.data.split(":")[1])
    t = await db.fetchone("SELECT * FROM tournaments WHERE id=?", (tid,))
    if not t or t["status"] != "open":
        return await c.answer("Bu turnirga ro'yxat yopilgan.", show_alert=True)
    if t["fee"] == 0:
        return await tn_join_do(c, bot, tid)
    u = await get_user(c.from_user.id)
    await c.answer()
    kb = InlineKeyboardBuilder()
    kb.button(text="✅ Ha, qatnashaman", callback_data=f"tnjy:{tid}")
    kb.button(text="🔙 Orqaga", callback_data=f"tnv:{tid}")
    kb.adjust(1)
    await safe_edit(c.message, f"🏆 <b>{esc(t['title'])}</b>\n\nKirish narxi: <b>{fmt(t['fee'])} so'm</b> balansingizdan yechiladi.\n"
                               f"💰 Balansingiz: <b>{fmt(u['balance'])} so'm</b>\n\nTasdiqlaysizmi?", kb.as_markup())


@user_r.callback_query(F.data.startswith("tnjy:"))
async def tn_join_yes(c: CallbackQuery, bot: Bot):
    await tn_join_do(c, bot, int(c.data.split(":")[1]))


async def tn_join_do(c: CallbackQuery, bot: Bot, tid: int):
    t, err = await tn_join(tid, c.from_user.id)
    if err:
        await c.answer(err, show_alert=True)
        return await tn_show(c, tid)
    await c.answer("✅ Ro'yxatga olindingiz!", show_alert=True)
    await tn_show(c, tid)
    n = await tn_count(tid)
    kb = InlineKeyboardBuilder()
    kb.button(text="✉️ Shu odamga kod/nom yuborish", callback_data=f"atns:{tid}:{c.from_user.id}")
    kb.adjust(1)
    txt = (f"🏆 <b>Yangi qatnashchi!</b>\n{esc(t['title'])}\n👤 {mention(c.from_user.id, c.from_user.full_name)} "
           f"(<code>{c.from_user.id}</code>)\n💰 To'ladi: {fee_text(t['fee'])}\n👥 {n}/{t['max_players']}"
           + ("\n🔔 <b>Joylar to'ldi!</b>" if n >= t["max_players"] else ""))
    for aid in await admin_ids():
        await tell(bot, aid, txt, reply_markup=kb.as_markup())


async def tn_comments_view(tid: int, admin: bool):
    rows = await db.fetchall("SELECT c.*, u.full_name FROM t_comments c LEFT JOIN users u ON u.id=c.user_id "
                             "WHERE c.tid=? ORDER BY c.id DESC LIMIT 10", (tid,))
    t = await db.fetchone("SELECT title FROM tournaments WHERE id=?", (tid,))
    lines = [f"<b>{esc((r['full_name'] or 'Mijoz')[:14])}</b> <i>({r['created_at'][11:16]})</i>: {esc(r['text'])}"
             for r in reversed(rows)]
    text = f"💬 <b>{esc(t['title'] if t else '')} — izohlar</b>\n\n" + ("\n".join(lines) if lines else "Hali izoh yo'q. Birinchi bo'ling!")
    kb = InlineKeyboardBuilder()
    kb.button(text="✍️ Izoh yozish", callback_data=f"tncw:{tid}")
    kb.button(text="🔙 Turnirga qaytish", callback_data=f"tnv:{tid}")
    sizes = [1, 1]
    if admin and rows:
        for r in rows:
            kb.button(text=f"🗑 #{r['id']}", callback_data=f"tncd:{r['id']}:{tid}")
        sizes += [5] * ((len(rows) + 4) // 5)
    kb.adjust(*sizes)
    return text, kb.as_markup()


@user_r.callback_query(F.data.startswith("tnc:"))
async def tn_comments_cb(c: CallbackQuery):
    await c.answer()
    text, kb = await tn_comments_view(int(c.data.split(":")[1]), await is_admin(c.from_user.id))
    await safe_edit(c.message, text, kb)


@user_r.callback_query(F.data.startswith("tncw:"))
async def tn_comment_ask(c: CallbackQuery, state: FSMContext):
    await c.answer()
    await state.set_state(TCom.text)
    await state.update_data(tid=int(c.data.split(":")[1]))
    await c.message.answer("✍️ Izohingizni yozing (200 belgigacha, havolalar taqiqlangan):", reply_markup=cancel_kb())


@user_r.message(StateFilter(TCom.text), F.text)
async def tn_comment_save(m: Message, state: FSMContext):
    d = await state.get_data()
    tid = d.get("tid")
    text = (m.text or "").strip()
    if not text or len(text) > 200:
        return await m.answer("❗️ Izoh 1–200 belgi bo'lsin.")
    if re.search(r"(https?://|t\.me|www\.|\.com|\.uz)", text, re.I) and not await is_admin(m.from_user.id):
        return await m.answer("🚫 Izohda havola yozish mumkin emas.")
    now = time.time()
    if now - _com_last.get(m.from_user.id, 0) < 5:
        return await m.answer("⏳ Biroz sekinroq — 5 soniyadan keyin yozing.")
    if not await db.fetchone("SELECT 1 FROM tournaments WHERE id=?", (tid,)):
        await state.clear()
        return await m.answer("Turnir topilmadi.", reply_markup=await main_menu(m.from_user.id))
    _com_last[m.from_user.id] = now
    await db.execute("INSERT INTO t_comments(tid,user_id,text) VALUES(?,?,?)", (tid, m.from_user.id, text))
    await state.clear()
    await m.answer("✅ Izoh qo'shildi!", reply_markup=await main_menu(m.from_user.id))
    ctext, ckb = await tn_comments_view(tid, await is_admin(m.from_user.id))
    await m.answer(ctext, reply_markup=ckb)


@user_r.message(StateFilter(TCom.text))
async def tn_comment_wrong(m: Message):
    await m.answer("✍️ Izohni matn ko'rinishida yozing.")


@admin_r.callback_query(F.data.startswith("tncd:"))
async def tn_comment_delete(c: CallbackQuery):
    _, cid, tid = c.data.split(":")
    await db.execute("DELETE FROM t_comments WHERE id=?", (int(cid),))
    await c.answer("🗑 O'chirildi")
    text, kb = await tn_comments_view(int(tid), True)
    await safe_edit(c.message, text, kb)


# ------------------------------------------------------------ 🏆 Turnir boshqaruvi (admin)
async def atn_menu_view():
    rows = await db.fetchall("SELECT * FROM tournaments WHERE status IN ('open','started') ORDER BY id DESC LIMIT 10")
    b = InlineKeyboardBuilder()
    b.button(text="➕ Turnir ochish", callback_data="atna")
    for t in rows:
        b.button(text=f"⚙️ #{t['id']} {t['title'][:24]} ({await tn_count(t['id'])}/{t['max_players']})",
                 callback_data=f"atnv:{t['id']}")
    b.adjust(1)
    return "🏆 <b>Turnir boshqaruvi</b>\n\n" + ("Faol turnirlar quyida." if rows else "Faol turnir yo'q."), b.as_markup()


@admin_r.message(F.text == T.A_TOURN)
async def atn_menu(m: Message):
    text, kb = await atn_menu_view()
    await m.answer(text, reply_markup=kb)


@admin_r.callback_query(F.data == "atn")
async def atn_menu_cb(c: CallbackQuery):
    await c.answer()
    text, kb = await atn_menu_view()
    await safe_edit(c.message, text, kb)


@admin_r.callback_query(F.data == "atna")
async def atn_create(c: CallbackQuery, state: FSMContext):
    await c.answer()
    await state.set_state(Adm.t_title)
    await c.message.answer("🏆 <b>Yangi turnir</b>\n\n1/6 — Turnir <b>nomini</b> yozing:", reply_markup=cancel_kb())


@admin_r.message(StateFilter(Adm.t_title))
async def atn_title(m: Message, state: FSMContext):
    if not (m.text or "").strip():
        return await m.answer("❗️ Nom yozing.")
    await state.update_data(title=m.text.strip()[:60])
    await state.set_state(Adm.t_fee)
    await m.answer("2/6 — <b>Kirish narxi</b> (so'm). Bepul bo'lsa 0 yozing:")


@admin_r.message(StateFilter(Adm.t_fee))
async def atn_fee(m: Message, state: FSMContext):
    n = parse_int(m.text or "")
    if n is None or n < 0:
        return await m.answer("❗️ 0 yoki musbat son kiriting.")
    await state.update_data(fee=n)
    await state.set_state(Adm.t_prize)
    await m.answer("3/6 — <b>Mukofot</b> (matn, masalan: 1-o'rin 660 UC, 2-o'rin 325 UC):")


@admin_r.message(StateFilter(Adm.t_prize))
async def atn_prize(m: Message, state: FSMContext):
    if not (m.text or "").strip():
        return await m.answer("❗️ Mukofotni yozing.")
    await state.update_data(prize=m.text.strip()[:200])
    await state.set_state(Adm.t_time)
    await m.answer("4/6 — <b>Boshlanish vaqti</b>: <code>21:00</code>, <code>ertaga 20:30</code> yoki <code>25.10 21:00</code> "
                   "(shunda 15 daqiqa oldin ishtirokchilarga eslatma avtomatik ketadi):")


@admin_r.message(StateFilter(Adm.t_time))
async def atn_time(m: Message, state: FSMContext):
    if not (m.text or "").strip():
        return await m.answer("❗️ Vaqtni yozing.")
    dt = parse_when(m.text)
    await state.update_data(start_text=(f"{dt:%d.%m %H:%M}" if dt else m.text.strip()[:80]),
                            start_at=(dt.strftime("%Y-%m-%d %H:%M") if dt else None))
    if not dt:
        await m.answer("ℹ️ Vaqt formati tushunilmadi — matn sifatida saqlandi, <b>avto-eslatma ishlamaydi</b>.")
    await state.set_state(Adm.t_max)
    await m.answer("5/6 — <b>Nechta kishi</b> qatnasha oladi? (son):")


@admin_r.message(StateFilter(Adm.t_max))
async def atn_max(m: Message, state: FSMContext):
    n = parse_int(m.text or "")
    if not n or not 2 <= n <= 10000:
        return await m.answer("❗️ 2 dan 10 000 gacha son kiriting.")
    await state.update_data(max_players=n)
    await state.set_state(Adm.t_info)
    await m.answer("6/6 — <b>Qo'shimcha ma'lumot</b>: xarita (Livik/Erangel...), rejim (solo/duo/squad), qoidalar. "
                   "Kerak bo'lmasa <code>-</code> yozing:")


@admin_r.message(StateFilter(Adm.t_info))
async def atn_info(m: Message, state: FSMContext):
    d = await state.get_data()
    info = (m.text or "").strip()
    info = "" if info == "-" else info[:800]
    tid, _ = await db.execute(
        "INSERT INTO tournaments(title,fee,prize,start_text,max_players,info,start_at) VALUES(?,?,?,?,?,?,?)",
        (d["title"], d["fee"], d["prize"], d["start_text"], d["max_players"], info, d.get("start_at")))
    await state.clear()
    await m.answer("✅ <b>Turnir ochildi!</b> Mijozlarga e'lon qilish uchun 📣 tugmasini bosing.", reply_markup=admin_kb())
    text, kb = await atn_view(tid)
    await m.answer(text, reply_markup=kb)


async def atn_view(tid: int):
    t = await db.fetchone("SELECT * FROM tournaments WHERE id=?", (tid,))
    text = "⚙️ <b>Turnir #%d</b>\n\n" % tid + await tn_text(t)
    if t["room_info"]:
        text += f"\n\n🔑 Yuborilgan kod/nom:\n{esc(t['room_info'])}"
    b = InlineKeyboardBuilder()
    sizes = []
    if t["status"] in ("open", "started"):
        b.button(text="📣 Barchaga e'lon qilish", callback_data=f"atnn:{tid}")
        b.button(text="📨 Hammaga kod/nom yuborish", callback_data=f"atns:{tid}:0")
        b.button(text="🔒 Ro'yxatni yopish" if t["status"] == "open" else "🔓 Ro'yxatni ochish", callback_data=f"atnt:{tid}")
        sizes += [1, 1, 1]
    b.button(text="👥 Ishtirokchilar", callback_data=f"atnp:{tid}")
    b.button(text="💬 Izohlar", callback_data=f"tnc:{tid}")
    sizes.append(2)
    if t["status"] != "cancelled":
        b.button(text="🏆 G'olibga mukofot berish", callback_data=f"atnw:{tid}")
        sizes.append(1)
    if t["status"] in ("open", "started"):
        b.button(text="🏁 Yakunlash", callback_data=f"atnf:{tid}")
        b.button(text="❌ Bekor qilish (pul qaytadi)", callback_data=f"atnx:{tid}")
        sizes += [1, 1]
    b.button(text="✏️ Tahrirlash", callback_data=f"edm:tn:{tid}")
    b.button(text="🔙 Turnirlar", callback_data="atn")
    sizes += [1, 1]
    b.adjust(*sizes)
    return text, b.as_markup()


@admin_r.callback_query(F.data.startswith("atnv:"))
async def atn_view_cb(c: CallbackQuery):
    await c.answer()
    text, kb = await atn_view(int(c.data.split(":")[1]))
    await safe_edit(c.message, text, kb)


@admin_r.callback_query(F.data.startswith("atnp:"))
async def atn_players(c: CallbackQuery):
    tid = int(c.data.split(":")[1])
    rows = await db.fetchall("SELECT p.*, u.full_name FROM t_players p LEFT JOIN users u ON u.id=p.user_id "
                             "WHERE p.tid=? ORDER BY p.joined_at LIMIT 80", (tid,))
    await c.answer()
    lines = [f"{i}. {mention(r['user_id'], r['full_name'])} (<code>{r['user_id']}</code>) • {fee_text(r['paid'])}"
             f"{' 🏆 +' + fmt(r['won']) if r['won'] else ''}"
             for i, r in enumerate(rows, 1)]
    await c.message.answer(f"👥 <b>Turnir #{tid} ishtirokchilari ({len(rows)})</b>\n\n" + ("\n".join(lines) or "Hali yo'q."))


@admin_r.callback_query(F.data.startswith("atnt:"))
async def atn_toggle(c: CallbackQuery):
    tid = int(c.data.split(":")[1])
    await db.execute("UPDATE tournaments SET status=CASE status WHEN 'open' THEN 'started' ELSE 'open' END "
                     "WHERE id=? AND status IN ('open','started')", (tid,))
    await c.answer("Bajarildi")
    text, kb = await atn_view(tid)
    await safe_edit(c.message, text, kb)


async def mass_send(bot: Bot, admin_id: int, ids: list, text: str, kb=None, label: str = "Xabar"):
    ok = fail = 0
    for uid in ids:
        try:
            await bot.send_message(uid, text, reply_markup=kb)
            ok += 1
        except TelegramRetryAfter as e:
            await asyncio.sleep(e.retry_after + 1)
            try:
                await bot.send_message(uid, text, reply_markup=kb)
                ok += 1
            except Exception:
                fail += 1
        except TelegramForbiddenError:
            await db.execute("UPDATE users SET blocked=1 WHERE id=?", (uid,))
            fail += 1
        except Exception:
            fail += 1
        await asyncio.sleep(0.05)
    await tell(bot, admin_id, f"📣 <b>{label} yuborildi</b>\n✅ Yetkazildi: {ok}\n❌ Xato/bloklagan: {fail}")


@admin_r.callback_query(F.data.startswith("atnn:"))
async def atn_announce(c: CallbackQuery, bot: Bot):
    tid = int(c.data.split(":")[1])
    t = await db.fetchone("SELECT * FROM tournaments WHERE id=?", (tid,))
    if not t or t["status"] != "open":
        return await c.answer("Faqat ro'yxati ochiq turnirni e'lon qilish mumkin.", show_alert=True)
    await c.answer("Yuborilmoqda...")
    ids = [r["id"] for r in await db.fetchall("SELECT id FROM users WHERE banned=0 AND blocked=0")]
    kb = InlineKeyboardBuilder()
    kb.button(text="🏆 Ko'rish va qatnashish", callback_data=f"tnv:{tid}")
    task = asyncio.create_task(mass_send(bot, c.from_user.id, ids, "🔥 <b>YANGI TURNIR!</b>\n\n" + await tn_text(t),
                                         kb.as_markup(), "Turnir e'loni"))
    BG_TASKS.add(task)
    task.add_done_callback(BG_TASKS.discard)
    await c.message.answer(f"⏳ E'lon {len(ids)} ta foydalanuvchiga yuborilmoqda. Tugagach hisobot keladi.")


@admin_r.callback_query(F.data.startswith("atns:"))
async def atn_send_ask(c: CallbackQuery, state: FSMContext):
    _, tid, uid = c.data.split(":")
    await c.answer()
    await state.set_state(Adm.t_send)
    await state.update_data(tid=int(tid), uid=int(uid))
    who = "barcha ishtirokchilarga" if int(uid) == 0 else f"<code>{uid}</code> ga"
    await c.message.answer(f"📨 {who} yuboriladigan <b>kod va nom</b>ni yozing (masalan: Xona ID: 123456 | Parol: 789 | "
                           f"Nom: PROUC):", reply_markup=cancel_kb())


@admin_r.message(StateFilter(Adm.t_send))
async def atn_send(m: Message, state: FSMContext, bot: Bot):
    d = await state.get_data()
    tid, uid = d["tid"], d["uid"]
    text = (m.text or "").strip()
    if not text:
        return await m.answer("❗️ Matn yozing.")
    t = await db.fetchone("SELECT * FROM tournaments WHERE id=?", (tid,))
    if not t:
        await state.clear()
        return await m.answer("Turnir topilmadi.", reply_markup=admin_kb())
    ids = [uid] if uid else [r["user_id"] for r in await db.fetchall("SELECT user_id FROM t_players WHERE tid=?", (tid,))]
    if uid == 0:
        await db.execute("UPDATE tournaments SET room_info=? WHERE id=?", (text[:500], tid))
    await state.clear()
    ok = 0
    for x in ids:
        try:
            await bot.send_message(x, f"🏆 <b>{esc(t['title'])}</b>\n🔑 <b>Xona ma'lumoti:</b>\n{esc(text)}\n\n🕒 {esc(t['start_text'])}\nOmad! 🍀")
            ok += 1
        except Exception:
            pass
        await asyncio.sleep(0.05)
    await m.answer(f"✅ Yuborildi: {ok}/{len(ids)}", reply_markup=admin_kb())


@admin_r.callback_query(F.data.startswith("atnf:"))
async def atn_finish(c: CallbackQuery, bot: Bot):
    tid = int(c.data.split(":")[1])
    _, rc = await db.execute("UPDATE tournaments SET status='finished' WHERE id=? AND status IN ('open','started')", (tid,))
    if not rc:
        return await c.answer("Allaqachon yakunlangan/bekor qilingan.", show_alert=True)
    await c.answer("🏁 Yakunlandi")
    t = await db.fetchone("SELECT title FROM tournaments WHERE id=?", (tid,))
    for r in await db.fetchall("SELECT user_id FROM t_players WHERE tid=?", (tid,)):
        await tell(bot, r["user_id"], f"🏁 <b>{esc(t['title'])}</b> turniri yakunlandi. Qatnashganingiz uchun rahmat! 🙏")
    text, kb = await atn_view(tid)
    await safe_edit(c.message, text, kb)


@admin_r.callback_query(F.data.startswith("atnx:"))
async def atn_cancel_ask(c: CallbackQuery):
    tid = int(c.data.split(":")[1])
    await c.answer()
    kb = InlineKeyboardBuilder()
    kb.button(text="✅ Ha, bekor qilish", callback_data=f"atny:{tid}")
    kb.button(text="🔙 Yo'q", callback_data=f"atnv:{tid}")
    kb.adjust(2)
    await safe_edit(c.message, "❌ Turnir bekor qilinsa, barcha ishtirokchilarga kirish narxi <b>balansga qaytariladi</b>. "
                               "Davom etamizmi?", kb.as_markup())


@admin_r.callback_query(F.data.startswith("atny:"))
async def atn_cancel_do(c: CallbackQuery, bot: Bot):
    tid = int(c.data.split(":")[1])
    players = []
    async with db.tx() as cx:
        t = await _one(cx, "SELECT * FROM tournaments WHERE id=?", (tid,))
        if not t or t["status"] not in ("open", "started"):
            return await c.answer("Allaqachon yakunlangan/bekor qilingan.", show_alert=True)
        cur = await cx.execute("SELECT * FROM t_players WHERE tid=?", (tid,))
        players = await cur.fetchall()
        await cur.close()
        for p in players:
            await cx.execute("UPDATE users SET balance=balance+? WHERE id=?", (p["paid"], p["user_id"]))
        await cx.execute("UPDATE tournaments SET status='cancelled' WHERE id=?", (tid,))
    await c.answer("Bekor qilindi")
    for p in players:
        await tell(bot, p["user_id"], f"❌ <b>{esc(t['title'])}</b> turniri bekor qilindi."
                                      + (f" {fee_text(p['paid'])} balansingizga qaytarildi." if p["paid"] else ""))
    text, kb = await atn_view(tid)
    await safe_edit(c.message, text, kb)


# ============================================================ ♻️ QO'LDA TIKLASH (admin: .db fayl + izoh «/restore»)
@admin_r.message(F.document, F.caption.startswith("/restore"))
async def manual_restore(m: Message, bot: Bot):
    if not (m.document.file_name or "").endswith(".db"):
        return await m.answer("❗️ .db fayl yuboring.")
    tmp = Path(DB_PATH).with_suffix(".upload")
    try:
        await bot.download(m.document, destination=str(tmp))
        if not _sqlite_ok(tmp):
            return await m.answer("❌ Fayl yaroqsiz yoki buzuq baza.")
        async with db.lock:
            await db.conn.close()
            for ext in ("-wal", "-shm"):
                Path(DB_PATH + ext).unlink(missing_ok=True)
            os.replace(tmp, DB_PATH)
        await db.connect()
        BK["locked"] = False
        await m.answer("♻️ <b>Baza tiklandi.</b> Hamma ma'lumot yangilandi.", reply_markup=admin_kb())
    except Exception as e:
        log.exception("Qo'lda tiklash xatosi")
        await m.answer(f"❌ Tiklashda xato: {esc(e)}")
    finally:
        tmp.unlink(missing_ok=True)


# ============================================================ WEB APP (TELEGRAM MINI APP) — REST API
def verify_init_data(init_data: str) -> Optional[dict]:
    """Telegram WebApp initData imzosini tekshiradi (rasmiy algoritm). To'g'ri bo'lsa user dict qaytaradi."""
    if not init_data:
        return None
    try:
        pairs = urllib.parse.parse_qsl(init_data, strict_parsing=True)
    except ValueError:
        return None
    data = dict(pairs)
    recv_hash = data.pop("hash", None)
    if not recv_hash:
        return None
    check_string = "\n".join(f"{k}={v}" for k, v in sorted(data.items()))
    secret_key = hmac.new(b"WebAppData", BOT_TOKEN.encode(), hashlib.sha256).digest()
    calc_hash = hmac.new(secret_key, check_string.encode(), hashlib.sha256).hexdigest()
    if not hmac.compare_digest(calc_hash, recv_hash):
        return None
    try:
        auth_date = int(data.get("auth_date", "0"))
    except ValueError:
        return None
    if INIT_DATA_MAX_AGE and (time.time() - auth_date) > INIT_DATA_MAX_AGE:
        return None
    try:
        user = json.loads(data.get("user", "{}"))
    except json.JSONDecodeError:
        return None
    if not user.get("id"):
        return None
    return user


def api_ok(payload=None, **extra):
    body = {"ok": True}
    if payload is not None:
        body["data"] = payload
    body.update(extra)
    return web.json_response(body)


def api_err(message: str, status: int = 400):
    return web.json_response({"ok": False, "error": message}, status=status)


async def fsm_set_wait_check(bot: Bot, uid: int, kind: str, ref_id: int):
    """Web App orqali yaratilgan buyurtma/to'ldirish uchun botni 'chek kutish' holatiga o'tkazadi,
    shunda foydalanuvchi botga qaytib rasm yuborsa avtomatik ushlanadi."""
    key = f"{bot.id}:{uid}:{uid}:None:default"
    await db.execute(
        "INSERT INTO fsm(key,state,data) VALUES(?,?,?) "
        "ON CONFLICT(key) DO UPDATE SET state=excluded.state, data=excluded.data",
        (key, Pay.check.state, json.dumps({"kind": kind, "ref_id": ref_id})))


@web.middleware
async def api_auth_middleware(request: web.Request, handler):
    if not request.path.startswith("/api/"):
        return await handler(request)
    if request.path in ("/api/health",):
        return await handler(request)
    init_data = request.headers.get("X-Init-Data") or request.query.get("initData", "")
    user = verify_init_data(init_data)
    if not user:
        return api_err("Avtorizatsiya muvaffaqiyatsiz (initData noto'g'ri).", 401)
    tg_id = int(user["id"])
    name = " ".join(filter(None, [user.get("first_name"), user.get("last_name")])) or user.get("username") or str(tg_id)

    class _U:
        pass
    fake = _U()
    fake.id, fake.full_name, fake.username, fake.is_bot = tg_id, name, user.get("username"), False
    await register_user(fake)
    urow = await get_user(tg_id)
    if urow and urow["banned"]:
        return api_err("Siz bloklangansiz.", 403)
    request["tg_user"] = fake
    request["uid"] = tg_id
    return await handler(request)


routes = web.RouteTableDef()


@routes.get("/api/health")
async def h_health(request):
    return api_ok({"status": "ok", "service": "Pro UC Bot API"})


@routes.get("/api/profile")
async def h_profile(request):
    u = await get_user(request["uid"])
    thr = await Si("vip_threshold")
    refs = await db.fetchone("SELECT COUNT(*) n FROM users WHERE referrer_id=?", (u["id"],))
    return api_ok({
        "id": u["id"], "full_name": u["full_name"], "balance": u["balance"], "coins": u["coins"],
        "total_uc": u["total_uc"], "is_vip": bool(u["is_vip"]), "vip_threshold": thr,
        "vip_remaining": max(thr - u["total_uc"], 0), "referrals": refs["n"],
        "is_admin": await is_admin(u["id"]), "joined_at": u["joined_at"],
    })


@routes.get("/api/settings")
async def h_settings(request):
    return api_ok({
        "support": await S("support"), "coin_value": await Si("coin_value"),
        "coin_buy_price": await Si("coin_buy_price"), "coin_max_percent": await Si("coin_max_percent"),
        "topup_min": await Si("topup_min"), "cashback_per_100uc": await Si("cashback_per_100uc"),
        "vip_threshold": await Si("vip_threshold"),
        "shop_open": await shop_is_open(), "work_hours": await S("work_hours"),
    })


@routes.get("/api/packages")
async def h_packages(request):
    rows = await db.fetchall("SELECT id,name,uc,price FROM packages WHERE active=1 ORDER BY uc")
    fl = await flash_pct_now()
    return api_ok([{**dict(r), "price": flash_price(r["price"], fl)} for r in rows])


@routes.get("/api/orders")
async def h_orders(request):
    rows = await db.fetchall("SELECT * FROM orders WHERE user_id=? ORDER BY id DESC LIMIT 30", (request["uid"],))
    return api_ok([{
        "id": r["id"], "code": r["code"], "pkg_name": r["pkg_name"], "uc": r["uc"], "final": r["final"],
        "status": r["status"], "status_text": STATUS_TEXT[r["status"]], "pay_method": r["pay_method"],
        "pubg_id": r["pubg_id"], "created_at": r["created_at"],
    } for r in rows])


@routes.get("/api/leaderboard")
async def h_leaderboard(request):
    rows = await db.fetchall("SELECT full_name,total_uc,is_vip FROM users WHERE total_uc>0 "
                             "ORDER BY total_uc DESC LIMIT 10")
    return api_ok([dict(r) for r in rows])


@routes.post("/api/coupon/check")
async def h_coupon_check(request):
    body = await request.json()
    pkg = await db.fetchone("SELECT * FROM packages WHERE id=? AND active=1", (body.get("pkg_id"),))
    if not pkg:
        return api_err("Paket topilmadi.")
    pr = await calc_price(db.conn, request["uid"], pkg, (body.get("code") or "").strip(), False)
    if pr["coupon_error"]:
        return api_err(pr["coupon_error"])
    return api_ok({"discount": pr["coupon_disc"], "code": pr["coupon_code"]})


@routes.post("/api/quote")
async def h_quote(request):
    """Berilgan paket/kupon/tanga tanlovi bo'yicha yakuniy narxni hisoblab beradi (checkout sahifasi uchun)."""
    body = await request.json()
    pkg = await db.fetchone("SELECT * FROM packages WHERE id=? AND active=1", (body.get("pkg_id"),))
    if not pkg:
        return api_err("Paket topilmadi.")
    pr = await calc_price(db.conn, request["uid"], pkg, body.get("code"), bool(body.get("use_coins")))
    u = await get_user(request["uid"])
    return api_ok({
        "pkg_name": pkg["name"], "uc": pkg["uc"], "price": pr["price"], "coupon_disc": pr["coupon_disc"],
        "coupon_error": pr["coupon_error"], "coin_used": pr["coin_used"], "coin_disc": pr["coin_disc"],
        "final": pr["final"], "balance": u["balance"], "coins": u["coins"],
    })


@routes.post("/api/order/create")
async def h_order_create(request: web.Request):
    body = await request.json()
    pkg_id = body.get("pkg_id")
    pubg_id = str(body.get("pubg_id", "")).strip()
    method = body.get("method")
    if method not in ("card", "balance"):
        return api_err("To'lov usuli noto'g'ri.")
    if not (pubg_id.isdigit() and 7 <= len(pubg_id) <= 12):
        return api_err("PUBG ID noto'g'ri (7-12 raqam bo'lishi kerak).")
    if method == "card" and not await S("card_number"):
        return api_err("Hozircha karta orqali to'lov qabul qilinmayapti. Admin bilan bog'laning.")
    oid, res = await create_order(request["uid"], pkg_id, pubg_id, body.get("code"),
                                  bool(body.get("use_coins")), method)
    if oid is None:
        return api_err(res)
    o = await db.fetchone("SELECT * FROM orders WHERE id=?", (oid,))
    bot: Bot = request.app["bot"]
    if method == "balance":
        await notify_admins(bot, "order", oid)
        return api_ok({"order_id": oid, "code": o["code"], "status": "checking", "final": o["final"]})
    await fsm_set_wait_check(bot, request["uid"], "order", oid)
    return api_ok({
        "order_id": oid, "code": o["code"], "status": "awaiting_check", "final": o["final"],
        "card_number": await S("card_number"), "card_owner": await S("card_owner"),
        "instruction": "Karta orqali to'lang va to'lov izohiga kodni yozing, so'ng chekni botga (chatga) rasm "
                       "sifatida yuboring.",
    })


@routes.post("/api/order/cancel")
async def h_order_cancel(request):
    body = await request.json()
    o = await db.fetchone("SELECT user_id FROM orders WHERE id=?", (body.get("order_id"),))
    if not o or o["user_id"] != request["uid"]:
        return api_err("Buyurtma topilmadi.", 404)
    res = await cancel_order(int(body["order_id"]), ("awaiting_check",))
    if not res:
        return api_err("Bu buyurtmani bekor qilib bo'lmaydi.")
    return api_ok({"cancelled": True})


@routes.post("/api/topup/create")
async def h_topup_create(request):
    body = await request.json()
    n = body.get("amount")
    mn = await Si("topup_min")
    if not isinstance(n, int) or n < mn or n > 100_000_000:
        return api_err(f"Summa {fmt(mn)} dan 100 000 000 gacha bo'lishi kerak.")
    if not await S("card_number"):
        return api_err("Hozircha to'lov qabul qilinmayapti.")
    code = await gen_code()
    tid, _ = await db.execute("INSERT INTO topups(code,user_id,amount,status) VALUES(?,?,?,'awaiting_check')",
                              (code, request["uid"], n))
    bot: Bot = request.app["bot"]
    await fsm_set_wait_check(bot, request["uid"], "topup", tid)
    return api_ok({
        "topup_id": tid, "code": code, "amount": n,
        "card_number": await S("card_number"), "card_owner": await S("card_owner"),
        "instruction": "To'lov izohiga kodni yozing, so'ng chekni botga (chatga) rasm sifatida yuboring.",
    })


@routes.post("/api/coin/exchange")
async def h_coin_exchange(request):
    body = await request.json()
    n = body.get("amount")
    if not isinstance(n, int) or n <= 0:
        return api_err("Musbat son kiriting.")
    val = n * await Si("coin_value")
    _, rc = await db.execute("UPDATE users SET coins=coins-?, balance=balance+? WHERE id=? AND coins>=?",
                             (n, val, request["uid"], n))
    if not rc:
        return api_err("Tangangiz yetarli emas.")
    return api_ok({"exchanged": n, "credited": val})


@routes.post("/api/coin/buy")
async def h_coin_buy(request):
    body = await request.json()
    n = body.get("amount")
    if not isinstance(n, int) or n <= 0:
        return api_err("Musbat son kiriting.")
    cost = n * await Si("coin_buy_price")
    _, rc = await db.execute("UPDATE users SET balance=balance-?, coins=coins+? WHERE id=? AND balance>=?",
                             (cost, n, request["uid"], cost))
    if not rc:
        return api_err("Balans yetarli emas.")
    return api_ok({"bought": n, "cost": cost})


@routes.post("/api/daily")
async def h_daily(request):
    today = now_tz().strftime("%Y-%m-%d")
    bonus = await Si("daily_bonus")
    _, rc = await db.execute("UPDATE users SET coins=coins+?, last_daily=? WHERE id=? "
                             "AND (last_daily IS NULL OR last_daily<>?)",
                             (bonus, today, request["uid"], today))
    if not rc:
        return api_err("Bugungi bonusni allaqachon olgansiz.")
    return api_ok({"bonus": bonus})


def build_web_app(bot: Bot) -> web.Application:
    app = web.Application(middlewares=[api_auth_middleware])
    app["bot"] = bot
    app.add_routes(routes)

    async def health(_):
        return web.json_response({"status": "ok", "bot": "Pro UC Bot"})
    app.router.add_get("/", health)

    if WEBAPP_DIR.exists():
        idx = WEBAPP_DIR / "index.html"

        async def app_index(_):
            if idx.exists():
                return web.FileResponse(idx)
            return web.Response(text="Web App topilmadi (webapp/index.html yo'q)", status=404)

        # index.html — asosiy Mini App sahifasi (/app va /app/ ikkalasi ham)
        app.router.add_get("/app", app_index)
        app.router.add_get("/app/", app_index)
        # webapp papkasidagi qo'shimcha fayllar (rasm, ikonka va h.k.) uchun /app/assets/*
        assets_dir = WEBAPP_DIR / "assets"
        assets_dir.mkdir(exist_ok=True)
        app.router.add_static("/app/assets/", path=str(assets_dir), show_index=False, name="webapp_assets")
    return app


# ============================================================ FON VAZIFALAR
async def expire_pending(bot: Bot):
    mod = f"-{ORDER_EXPIRE_MIN} minutes"
    for r in await db.fetchall("SELECT id,user_id,code FROM orders WHERE status='awaiting_check' "
                               "AND created_at < datetime('now','+5 hours',?)", (mod,)):
        if await cancel_order(r["id"], ("awaiting_check",)):
            await tell(bot, r["user_id"], f"⌛️ Buyurtma <b>{r['code']}</b> vaqt tugagani uchun bekor qilindi "
                                          f"(tanga/kupon qaytarildi).")
    await db.execute("UPDATE topups SET status='cancelled' WHERE status='awaiting_check' "
                     "AND created_at < datetime('now','+5 hours',?)", (mod,))


async def auto_lottery(bot: Bot):
    now = now_tz()
    if now.day != 1:
        return
    prev = (now.replace(day=1) - timedelta(days=1)).strftime("%Y-%m")
    if not await db.fetchone("SELECT 1 FROM lottery_log WHERE month=?", (prev,)):
        log.info("Avto VIP o'yini: %s", await run_lottery(bot, prev))


async def maintenance(bot: Bot):
    tick = 0
    while True:
        jobs = [expire_pending, tournament_reminders, flash_tick, price_post_tick, daily_report]
        if tick % 60 == 0:
            jobs.append(auto_lottery)
        for job in jobs:   # bitta ish xato bersa, qolganlari to'xtamaydi
            try:
                await job(bot)
            except Exception as e:
                log.exception("maintenance xatosi: %s", job.__name__)
                await report_error(bot, f"maintenance/{job.__name__}", e)
        if tick % 60 == 0:
            try:
                await db.execute("DELETE FROM audit WHERE created_at < datetime('now','+5 hours','-90 days')")
            except Exception:
                log.exception("audit tozalash xatosi")
        tick += 1
        await asyncio.sleep(60)


# ============================================================ ISHGA TUSHIRISH
async def main():
    if not BOT_TOKEN:
        raise SystemExit("❌ BOT_TOKEN topilmadi (.env faylini tekshiring).")
    Path(DB_PATH).parent.mkdir(parents=True, exist_ok=True)
    moved = False
    try:
        moved = migrate_old_db(DB_PATH)
        if moved:
            log.info("♻️ Eski baza (%s) doimiy joyga ko'chirildi: %s", OLD_DB_PATH, DB_PATH)
    except Exception:
        log.exception("Eski bazani ko'chirishda xato (eski fayl joyida qoldi)")
    bot = Bot(BOT_TOKEN, default=DefaultBotProperties(parse_mode=ParseMode.HTML))
    restore = await tg_restore(bot)   # faqat BACKUP_CHAT_ID berilgan bo'lsa ishlaydi
    fresh = not (Path(DB_PATH).exists() and Path(DB_PATH).stat().st_size > 0)
    if restore == "error":
        BK["locked"] = True           # yaxshi zaxira ustiga bo'sh baza yozilib ketmasin
    await db.connect()
    if restore == "restored" and BK["restored_msg"]:
        await db.execute("UPDATE settings SET value=? WHERE key='backup_msg_id'", (str(BK["restored_msg"]),))
    for aid in ([SUPER_ADMIN_ID] if SUPER_ADMIN_ID else []):
        if restore == "restored":
            await tell(bot, aid, "♻️ <b>Baza Telegram zaxirasidan tiklandi.</b> Barcha ma'lumotlar joyida.")
        elif restore == "error":
            await tell(bot, aid, "⚠️ Zaxirani tiklab bo'lmadi, yangi baza ochildi. <b>Avto-zaxira o'chirildi</b> — eski "
                                 "zaxira ustiga yozilmasligi uchun. Pin qilingan .db faylni qo'lda tiklang: "
                                 "faylni botga «/restore» izohi bilan yuboring.")
    if not BACKUP_CHAT and SUPER_ADMIN_ID:   # eski versiyalar chatga yuborgan zaxira faylini (pin) tozalaymiz
        try:
            old_id = int(await S("backup_msg_id") or 0)
            if old_id:
                try:
                    await bot.unpin_chat_message(SUPER_ADMIN_ID, message_id=old_id)
                except Exception:
                    pass
                try:
                    await bot.delete_message(SUPER_ADMIN_ID, old_id)
                except Exception:
                    pass
                await db.execute("UPDATE settings SET value='0' WHERE key='backup_msg_id'")
        except Exception:
            log.exception("Eski zaxira xabarini tozalashda xato")
    if SUPER_ADMIN_ID:
        cnt = await db.fetchone("SELECT (SELECT COUNT(*) FROM users) u, (SELECT COUNT(*) FROM orders) o")
        size = Path(DB_PATH).stat().st_size // 1024 if Path(DB_PATH).exists() else 0
        msg = (f"🟢 <b>Bot ishga tushdi</b> ({now_tz():%d.%m %H:%M})\n💾 Baza: <code>{esc(DB_PATH)}</code> • {size} KB\n"
               f"👥 {cnt['u']} foydalanuvchi • 🧾 {cnt['o']} buyurtma")
        if moved:
            msg += "\n♻️ Eski baza doimiy joyga ko'chirildi."
        if fresh and restore not in ("restored",) and not moved:
            msg += ("\n\n⚠️ <b>YANGI BO'SH BAZA yaratildi.</b> Agar bot avval ishlagan bo'lsa — ma'lumot hosting diskida "
                    "saqlanmagan (disk doimiy emas). Hostingda <b>doimiy disk (Volume)</b> ulang yoki eski .db faylni "
                    "botga <code>/restore</code> izohi bilan yuboring.")
        await tell(bot, SUPER_ADMIN_ID, msg)
    dp = Dispatcher(storage=SQLiteStorage(db))
    admin_r.message.middleware(RoleMiddleware())
    admin_r.callback_query.middleware(RoleMiddleware())
    dp.errors.register(on_error)
    dp.message.filter(F.chat.type == "private")
    dp.message.outer_middleware(StateInterruptMiddleware())
    dp.message.outer_middleware(GateMiddleware())
    dp.callback_query.outer_middleware(GateMiddleware())
    dp.include_router(common_r)
    dp.include_router(admin_r)
    dp.include_router(user_r)

    maint = asyncio.create_task(maintenance(bot))
    bkp = asyncio.create_task(backup_loop(bot))
    runner = None
    try:
        app = build_web_app(bot)  # 🎮 Web App (Mini App) + REST API — har doim ishlaydi
        if WEBHOOK_URL:
            from aiogram.webhook.aiohttp_server import SimpleRequestHandler, setup_application
            SimpleRequestHandler(dp, bot, secret_token=WEBHOOK_SECRET or None).register(app, path=WEBHOOK_PATH)
            setup_application(app, dp, bot=bot)
            await bot.set_webhook(WEBHOOK_URL + WEBHOOK_PATH, secret_token=WEBHOOK_SECRET or None,
                                  allowed_updates=dp.resolve_used_update_types())
            runner = web.AppRunner(app)
            await runner.setup()
            await web.TCPSite(runner, WEB_HOST, WEB_PORT).start()
            log.info("Webhook + Web App server: %s:%s (Mini App: /app, API: /api/*)", WEB_HOST, WEB_PORT)
            stop = asyncio.Event()
            try:
                import signal
                for sg in (signal.SIGINT, signal.SIGTERM):
                    asyncio.get_running_loop().add_signal_handler(sg, stop.set)
            except (NotImplementedError, RuntimeError):
                pass
            await stop.wait()
        else:
            await bot.delete_webhook(drop_pending_updates=False)
            runner = web.AppRunner(app)
            await runner.setup()
            await web.TCPSite(runner, WEB_HOST, WEB_PORT).start()
            log.info("Polling rejimi + Web App server: %s:%s (Mini App: /app, API: /api/*)", WEB_HOST, WEB_PORT)
            await dp.start_polling(bot, allowed_updates=dp.resolve_used_update_types())
    finally:
        maint.cancel()
        bkp.cancel()
        try:   # o'chishdan oldin oxirgi zaxira
            await asyncio.wait_for(tg_backup(bot, force=True), timeout=12)
        except Exception:
            log.warning("Yakuniy zaxira olinmadi")
        if runner:
            await runner.cleanup()
        await db.close()
        await bot.session.close()


if __name__ == "__main__":
    try:
        asyncio.run(main())
    except (KeyboardInterrupt, SystemExit):
        pass
